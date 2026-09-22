"""
Put a reviewed language workbook back into NurtureHUB.

Reads the file produced by scripts.export_translation_sheet and:
  * applies "Hindi - correction" over the Hindi in use;
  * writes the Marathi column into a new mr/ set of translation files;
  * registers Marathi in the app (i18n/index.ts) the first time.

Rows are matched on the ID column, so reviewers may sort or filter freely.
Nothing is written unless --apply is given; without it this is a report.

Checks before anything is written:
  * every {{placeholder}} and <tag> in the English must appear in the
    translation — a dropped {{name}} shows the reader a blank where their
    name should be, or crashes the line;
  * the English in the sheet must still match the English in the app. If a
    sentence was reworded after the sheet went out, the row is reported and
    skipped rather than pairing new English with old Hindi.
  * server messages (IDs starting "server:") are reported, not applied —
    those sentences live in Python and are English-only today.

Usage (from backend/):
  venv-win/Scripts/python.exe -m scripts.import_translation_sheet --file ../reviewed.xlsx
  venv-win/Scripts/python.exe -m scripts.import_translation_sheet --file ../reviewed.xlsx --apply
"""
from __future__ import annotations

import argparse
import json
import re
from collections import defaultdict
from pathlib import Path
from typing import Dict, List, Tuple

from openpyxl import load_workbook

REPO = Path(__file__).resolve().parents[2]
LOCALES = REPO / "frontend" / "src" / "i18n" / "locales"
I18N_INDEX = REPO / "frontend" / "src" / "i18n" / "index.ts"

COL_ID, COL_EN, COL_HI, COL_FIX, COL_MR = 1, 3, 4, 5, 6


def _flatten(d: dict, prefix: str = "") -> Dict[str, str]:
    out: Dict[str, str] = {}
    for k, v in d.items():
        key = f"{prefix}.{k}" if prefix else k
        out.update(_flatten(v, key)) if isinstance(v, dict) else out.update({key: v})
    return out


def _nest(flat: Dict[str, str]) -> dict:
    root: dict = {}
    for key, value in flat.items():
        node = root
        parts = key.split(".")
        for p in parts[:-1]:
            node = node.setdefault(p, {})
        node[parts[-1]] = value
    return root


def _tokens(text: str) -> List[str]:
    return sorted(re.findall(r"\{\{[^}]+\}\}|<[^>]+>", text or ""))


def read_sheet(path: Path) -> List[dict]:
    wb = load_workbook(path, read_only=True, data_only=True)
    rows: List[dict] = []
    for name in wb.sheetnames:
        if name.strip().lower().startswith("read me"):
            continue
        ws = wb[name]
        for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            if not row or not row[COL_ID - 1]:
                continue
            # Never strip: a few strings carry a deliberate leading or trailing
            # space (" - Test Phase", "Quiz: ") and losing it would join words.
            get = lambda c: (str(row[c - 1]) if len(row) >= c and row[c - 1] is not None else "")  # noqa: E731
            rows.append({
                "sheet": name, "line": i, "id": get(COL_ID), "en": get(COL_EN),
                "hi": get(COL_HI), "fix": get(COL_FIX), "mr": get(COL_MR),
            })
    return rows


def main() -> None:
    ap = argparse.ArgumentParser(description="Apply a reviewed language workbook.")
    ap.add_argument("--file", required=True, help="the reviewed .xlsx")
    ap.add_argument("--apply", action="store_true", help="write the files (default: report only)")
    ap.add_argument("--force", action="store_true", help="apply rows that failed a check (not advised)")
    args = ap.parse_args()

    rows = read_sheet(Path(args.file))
    current: Dict[str, Dict[str, str]] = {}
    for p in sorted((LOCALES / "en").glob("*.json")):
        current[p.stem] = _flatten(json.loads(p.read_text(encoding="utf-8")))

    hi_updates: Dict[str, Dict[str, str]] = defaultdict(dict)
    mr_updates: Dict[str, Dict[str, str]] = defaultdict(dict)
    problems: List[str] = []
    server_rows: List[dict] = []
    unknown = stale = 0

    for r in rows:
        if r["id"].startswith("server:"):
            if r["fix"] or r["mr"]:
                server_rows.append(r)
            continue
        if ":" not in r["id"]:
            problems.append(f"{r['sheet']} line {r['line']}: ID '{r['id']}' is not in the expected form")
            continue
        ns, key = r["id"].split(":", 1)
        english = current.get(ns, {}).get(key)
        if english is None:
            unknown += 1
            problems.append(f"{r['sheet']} line {r['line']}: {r['id']} is no longer in the app - skipped")
            continue
        if r["en"].strip() and r["en"] != english:
            stale += 1
            problems.append(f"{r['sheet']} line {r['line']}: {r['id']} - the English changed since the sheet "
                            f"went out, so the translation may not fit. Skipped.")
            if not args.force:
                continue
        for label, value, bucket in (("Hindi correction", r["fix"], hi_updates), ("Marathi", r["mr"], mr_updates)):
            if not value.strip():
                continue
            # Trim what the reviewer may have added, then restore the spacing
            # the English carries on purpose.
            value = value.strip()
            if english[:1] == " ":
                value = " " + value
            if english[-1:] == " ":
                value = value + " "
            missing = [t for t in _tokens(english) if t not in value]
            if missing and not args.force:
                problems.append(f"{r['sheet']} line {r['line']}: {r['id']} - {label} is missing "
                                f"{' '.join(missing)} - skipped")
                continue
            bucket[ns][key] = value

    print(f"Read {len(rows)} rows from {args.file}")
    print(f"  Hindi corrections to apply : {sum(len(v) for v in hi_updates.values())}")
    print(f"  Marathi translations       : {sum(len(v) for v in mr_updates.values())}")
    if server_rows:
        print(f"  Server messages (not applied automatically): {len(server_rows)}")
        for r in server_rows[:10]:
            print(f"      {r['id']}")
        print("      These live in Python and the server sends them in English today;")
        print("      translating them needs a code change. Keep this sheet for that work.")
    if problems:
        print(f"\n  {len(problems)} row(s) need attention:")
        for p in problems[:25]:
            print(f"      {p}")
        if len(problems) > 25:
            print(f"      ... and {len(problems) - 25} more")

    if not args.apply:
        print("\nNothing written (add --apply to write the files).")
        return

    written: List[str] = []
    for ns, updates in hi_updates.items():
        path = LOCALES / "hi" / f"{ns}.json"
        flat = _flatten(json.loads(path.read_text(encoding="utf-8"))) if path.exists() else {}
        flat.update(updates)
        path.write_text(json.dumps(_nest(flat), ensure_ascii=False, indent=2) + "\n", encoding="utf-8", newline="\n")
        written.append(f"hi/{ns}.json (+{len(updates)})")

    if mr_updates:
        (LOCALES / "mr").mkdir(exist_ok=True)
        for ns, updates in mr_updates.items():
            path = LOCALES / "mr" / f"{ns}.json"
            flat = _flatten(json.loads(path.read_text(encoding="utf-8"))) if path.exists() else {}
            flat.update(updates)
            path.write_text(json.dumps(_nest(flat), ensure_ascii=False, indent=2) + "\n", encoding="utf-8", newline="\n")
            written.append(f"mr/{ns}.json (+{len(updates)})")
        # Every namespace needs a file, even an empty one: the app imports them
        # all by name. Anything not translated falls back to English.
        for p in sorted((LOCALES / "en").glob("*.json")):
            target = LOCALES / "mr" / p.name
            if not target.exists():
                target.write_text("{}\n", encoding="utf-8", newline="\n")
                written.append(f"mr/{p.name} (empty, falls back to English)")
        if register_marathi():
            written.append("i18n/index.ts (Marathi added to the language switcher)")

    print("\nWritten:")
    for w in written:
        print(f"  {w}")
    print("\nNext: cd frontend && npx tsc -b && npm run build, check the app, then commit.")


def register_marathi() -> bool:
    """Add Marathi to the app's language list. Returns True if it changed."""
    src = I18N_INDEX.read_text(encoding="utf-8")
    if "locales/mr/" in src:
        return False
    namespaces = re.search(r"export const NAMESPACES = \[\n(.*?)\] as const;", src, re.S)
    names = re.findall(r"'([a-zA-Z]+)'", namespaces.group(1)) if namespaces else []
    imports = "\n".join(
        f"import mr{n[0].upper() + n[1:]} from './locales/mr/{n}.json';" for n in names)
    src = src.replace("\n/**\n * Languages the UI actually ships",
                      f"\n{imports}\n\n/**\n * Languages the UI actually ships", 1)
    src = src.replace("  { code: 'hi', label: 'Hindi', native: 'हिन्दी' },",
                      "  { code: 'hi', label: 'Hindi', native: 'हिन्दी' },\n"
                      "  { code: 'mr', label: 'Marathi', native: 'मराठी' },", 1)
    block = ",\n".join(f"    {n}: mr{n[0].upper() + n[1:]}" for n in names)
    src = re.sub(r"(\n  hi: \{\n(?:.*?\n)*?  \},\n)", r"\1  mr: {\n" + block + ",\n  },\n", src, count=1)
    I18N_INDEX.write_text(src, encoding="utf-8", newline="\n")
    return True


if __name__ == "__main__":
    main()
