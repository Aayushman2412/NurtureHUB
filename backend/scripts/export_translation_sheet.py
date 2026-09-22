"""
Export every sentence NurtureHUB shows people into one Excel workbook for
language review: English, the Hindi we use today, and an empty Marathi column.

What goes in:
  * every phrase in the app's translation files (frontend/src/i18n/locales),
    split into "Learner app" (what health workers see) and "Admin panel";
  * the messages and the verification email the server sends, which live in
    Python code rather than the translation files (BACKEND_STRINGS below,
    kept honest by tests/test_translation_sheet.py).

The reviewed workbook goes back in through scripts.import_translation_sheet,
which matches rows on the ID column — so reviewers may re-order or filter
rows, but must not edit IDs.

Usage (from backend/):
  venv-win/Scripts/python.exe -m scripts.export_translation_sheet
  venv-win/Scripts/python.exe -m scripts.export_translation_sheet --out ../NurtureHUB_language_review.xlsx
"""
from __future__ import annotations

import argparse
import json
import re
from pathlib import Path
from typing import Dict, List, Tuple

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

REPO = Path(__file__).resolve().parents[2]
LOCALES = REPO / "frontend" / "src" / "i18n" / "locales"

# namespace -> (sheet, where it appears)
AREAS: Dict[str, Tuple[str, str]] = {
    "app": ("Learner app", "App menus and shell"),
    "auth": ("Learner app", "Sign in and sign up"),
    "common": ("Learner app", "Common buttons and words"),
    "landing": ("Learner app", "Public home page"),
    "dashboard": ("Learner app", "Learner dashboard"),
    "learner": ("Learner app", "Learner profile and registration"),
    "tutorials": ("Learner app", "Videos (learner side)"),
    "tests": ("Learner app", "Tests (learner side)"),
    "mother": ("Learner app", "Mother and child records"),
    "assessments": ("Learner app", "Assessment forms"),
    "growth": ("Learner app", "Growth monitoring"),
    "offline": ("Learner app", "Offline and sync messages"),
    "validation": ("Learner app", "Form error messages"),
    "admin": ("Admin panel", "Admin - general"),
    "adminTests": ("Admin panel", "Admin - test manager"),
    "adminTutorials": ("Admin panel", "Admin - videos manager"),
    "adminFormBuilder": ("Admin panel", "Admin - form builder"),
    "adminLiveMonitor": ("Admin panel", "Admin - live monitoring"),
    "adminResults": ("Admin panel", "Admin - results"),
    "resultsInsights": ("Admin panel", "Admin - results insights"),
    "pipelines": ("Admin panel", "Admin - data pipelines"),
}

# Sentences the SERVER sends. They are not in the translation files, so they
# are listed here by hand: (id, where, english, note). Placeholders use the
# Python {name} form the code uses. tests/test_translation_sheet.py checks
# each of these is still present in the file named in `where`.
BACKEND_STRINGS: List[Tuple[str, str, str, str]] = [
    ("server:email.otp.subject", "Verification email - subject (app/utils.py)",
     "NurtureHUB verification code: {otp}", "{otp} is the 6-digit code."),
    ("server:email.otp.text", "Verification email - plain text (app/utils.py)",
     "Your verification code for NurtureHUB is {otp}. It will expire in {minutes} minutes.", ""),
    ("server:email.otp.heading", "Verification email - heading (app/utils.py)",
     "NurtureHUB Verification Code", ""),
    ("server:email.otp.greeting", "Verification email (app/utils.py)", "Hello,", ""),
    ("server:email.otp.intro", "Verification email (app/utils.py)",
     "Thank you for using NurtureHUB. Please use the following One-Time Password (OTP) to verify your account:", ""),
    ("server:email.otp.expiry", "Verification email (app/utils.py)",
     "This code will expire in {minutes} minutes. If you did not request this, you can safely ignore this email.", ""),
    ("server:email.otp.footer", "Verification email (app/utils.py)",
     "NurtureHUB - Assessment & Training Platform", ""),
    ("server:notification.tutorialCompleted.title", "Notification after finishing a video (app/routers/tutorials.py)",
     "Tutorial Completed", ""),
    ("server:notification.tutorialCompleted.body", "Notification after finishing a video (app/routers/tutorials.py)",
     "You have successfully completed the tutorial: '{tutorial_title}'.", "{tutorial_title} is the video's name."),
    ("server:notification.testDone.title", "Notification after a test (app/routers/tests.py)",
     "Test Attempt Completed: {status}", "{status} is 'Passed' or 'Failed' (next two rows)."),
    ("server:notification.testDone.passed", "Notification after a test (app/routers/tests.py)", "Passed", ""),
    ("server:notification.testDone.failed", "Notification after a test (app/routers/tests.py)", "Failed", ""),
    ("server:notification.testDone.body", "Notification after a test (app/routers/tests.py)",
     "You completed the test '{test_title}'{how} with a score of {score}% ({correct}/{total} correct).",
     "{how} is empty, or the next row's words."),
    ("server:notification.testDone.auto", "Notification after a test (app/routers/tests.py)",
     " (submitted automatically when time ran out)", "Keep the leading space."),
    ("server:notification.resultsPending.title", "Notification when everything is done (app/flow.py)",
     "Results Pending", ""),
    ("server:notification.resultsPending.body", "Notification when everything is done (app/flow.py)",
     "Congratulations! You have completed all tutorials and tests. "
     "Please wait for your results - you will be notified once they are announced.", ""),
    ("server:notification.faceToFace.title", "Notification on face-to-face selection (app/routers/admin.py)",
     "Selected for Face-to-Face Training", ""),
    ("server:notification.faceToFace.body", "Notification on face-to-face selection (app/routers/admin.py)",
     "Congratulations! You have been selected for the face-to-face training. "
     "Please await further instructions.", ""),
    ("server:notification.protein.title", "Notification on an unusual protein total (app/routers/forms.py)",
     "Unusually high protein - {subject_name}", "{subject_name} is the mother's or child's name."),
    ("server:notification.protein.learner", "Notification on an unusual protein total (app/routers/forms.py)",
     "The recorded 24-hour protein total is {grams}g, above the {limit}g review threshold. "
     "Please re-check the portion sizes with the mother.", ""),
    ("server:notification.protein.admin", "Notification to admins on an unusual protein total (app/routers/forms.py)",
     "{learner_name} recorded a 24-hour protein total of {grams}g for {subject_name}, "
     "above the {limit}g review threshold. Worth confirming the portion sizes.", ""),
    ("server:notification.formSummary.body", "Notification after an assessment form (app/routers/forms.py)",
     "{green} step(s) as per LAP, {red} need attention. Open the assessment plan to see the recommended actions.", ""),
]

# ── Styling ────────────────────────────────────────────────────────────────

CORAL = "E85D4C"
HEADERS = [
    ("ID - do not change", 34, "EFEAE1"),
    ("Where it appears", 26, "EFEAE1"),
    ("English (in use now)", 48, "FFFFFF"),
    ("Hindi (in use now)", 48, "FFFFFF"),
    ("Hindi - correction (only if wrong)", 38, "EAF6EE"),
    ("Marathi (please fill in)", 48, "FFF8E1"),
    ("Words that must stay exactly", 22, "FFFFFF"),
    ("Reviewer notes", 28, "FFFFFF"),
]
DEVANAGARI_COLS = {4, 5, 6}          # 1-based: Hindi, correction, Marathi
THIN = Side(style="thin", color="D9D2C7")


def _flatten(d: dict, prefix: str = "") -> Dict[str, str]:
    out: Dict[str, str] = {}
    for k, v in d.items():
        key = f"{prefix}.{k}" if prefix else k
        if isinstance(v, dict):
            out.update(_flatten(v, key))
        else:
            out[key] = v
    return out


def _placeholders(text: str) -> str:
    """Everything a translator must copy across untouched."""
    found = re.findall(r"\{\{[^}]+\}\}|\{[a-zA-Z_][^}]*\}|<[^>]+>", text or "")
    seen, out = set(), []
    for f in found:
        if f not in seen:
            seen.add(f)
            out.append(f)
    return "  ".join(out)


def collect_rows() -> Dict[str, List[list]]:
    """{sheet name: rows}, each row matching HEADERS."""
    sheets: Dict[str, List[list]] = {"Learner app": [], "Admin panel": [], "Messages & email": []}
    for path in sorted((LOCALES / "en").glob("*.json")):
        ns = path.stem
        sheet, where = AREAS.get(ns, ("Admin panel", ns))
        en = _flatten(json.loads(path.read_text(encoding="utf-8")))
        hi_path = LOCALES / "hi" / path.name
        hi = _flatten(json.loads(hi_path.read_text(encoding="utf-8"))) if hi_path.exists() else {}
        for key, text in en.items():
            note = "" if str(hi.get(key, "")).strip() else "Hindi missing - please add"
            sheets[sheet].append([f"{ns}:{key}", where, text, hi.get(key, ""), "", "", _placeholders(text), note])
    for sid, where, text, note in BACKEND_STRINGS:
        sheets["Messages & email"].append([sid, where, text, "", "", "", _placeholders(text), note])
    return sheets


def _write_sheet(wb: Workbook, title: str, rows: List[list]) -> None:
    ws = wb.create_sheet(title)
    ws.append([h[0] for h in HEADERS])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF", size=11)
        cell.fill = PatternFill("solid", fgColor=CORAL)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 34
    for i, (_, width, _) in enumerate(HEADERS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width

    for row in rows:
        ws.append(row)
    for r in range(2, ws.max_row + 1):
        for c in range(1, len(HEADERS) + 1):
            cell = ws.cell(row=r, column=c)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
            fill = HEADERS[c - 1][2]
            if fill != "FFFFFF":
                cell.fill = PatternFill("solid", fgColor=fill)
            if c in DEVANAGARI_COLS:
                cell.font = Font(name="Nirmala UI", size=11)
            if c in (5, 6, 8):
                # Text format, so a translation that starts with + or = is not
                # taken for a formula when it is typed in.
                cell.number_format = "@"
            if c == 1:
                cell.font = Font(size=9, color="6B6357")
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{ws.max_row}"


def _write_readme(wb: Workbook, counts: Dict[str, int]) -> None:
    ws = wb.create_sheet("Read me first", 0)
    ws.column_dimensions["A"].width = 4
    ws.column_dimensions["B"].width = 110
    lines: List[Tuple[str, str]] = [
        ("title", "NurtureHUB - language review"),
        ("", ""),
        ("h", "What this file is"),
        ("p", "Every sentence NurtureHUB shows on screen, plus the messages and the verification "
              "email it sends. Each row is one sentence: the English in use now, the Hindi in use "
              "now, and an empty Marathi column."),
        ("", ""),
        ("h", "What we would like you to do"),
        ("n", "1. Read the English and the Hindi side by side. If the Hindi is fine, leave the "
              "correction column empty."),
        ("n", "2. If the Hindi is wrong or unclear, write the better Hindi in the column "
              "'Hindi - correction'. Please do not edit the 'Hindi (in use now)' column."),
        ("n", "3. Write the Marathi in the yellow column 'Marathi (please fill in)'."),
        ("n", "4. Use the last column for any note you want to leave us."),
        ("", ""),
        ("h", "Two things to keep exactly as they are"),
        ("n", "1. The ID column. We use it to put your words back into the app. Please do not "
              "change, delete or re-type it."),
        ("n", "2. Anything listed in 'Words that must stay exactly' - for example {{name}}, "
              "{{count}}, <b>. The app replaces these with real values (a name, a number). "
              "Copy them into the translation as they are; you may move them within the "
              "sentence so the sentence reads naturally."),
        ("", ""),
        ("h", "Good to know"),
        ("p", "Sentences are short on purpose: they are buttons, labels and messages, not "
              "paragraphs. Where a sentence ends with ... or a colon, please keep it."),
        ("p", "Some rows come in pairs ending in _one and _other: the first is used for one "
              "item, the second for two or more. Both need a translation."),
        ("p", "You may sort or filter rows while you work - we match on the ID, not the order."),
        ("", ""),
        ("h", "The sheets in this file"),
    ]
    for name, n in counts.items():
        lines.append(("n", f"{name}: {n} sentences"))
    lines += [
        ("", ""),
        ("h", "हिंदी में संक्षेप में"),
        ("p", "इस फ़ाइल में NurtureHUB की हर पंक्ति है - अभी इस्तेमाल हो रही अंग्रेज़ी और हिंदी, "
              "और मराठी के लिए खाली कॉलम। हिंदी ठीक हो तो कुछ न करें; सुधार चाहिए तो "
              "'Hindi - correction' कॉलम में लिखें। मराठी पीले कॉलम में लिखें। "
              "ID कॉलम और {{ }} या < > में लिखी चीज़ें बिलकुल वैसी ही रहने दें।"),
        ("", ""),
        ("p", "Thank you - once we get this back, the corrections and the Marathi go straight "
              "into NurtureHUB."),
    ]
    r = 1
    for kind, text in lines:
        cell = ws.cell(row=r, column=2, value=text)
        if kind == "title":
            cell.font = Font(bold=True, size=18, color="D14432")
            ws.row_dimensions[r].height = 30
        elif kind == "h":
            cell.font = Font(bold=True, size=12, color="26221C")
        else:
            cell.font = Font(size=11, name="Nirmala UI")
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            ws.row_dimensions[r].height = max(16, 15 * (len(text) // 95 + 1))
        r += 1


def main() -> None:
    ap = argparse.ArgumentParser(description="Export every NurtureHUB sentence for language review.")
    ap.add_argument("--out", default=str(REPO / "NurtureHUB_language_review.xlsx"))
    args = ap.parse_args()

    sheets = collect_rows()
    wb = Workbook()
    wb.remove(wb.active)
    counts = {name: len(rows) for name, rows in sheets.items()}
    _write_readme(wb, counts)
    for name, rows in sheets.items():
        _write_sheet(wb, name, rows)
    out = Path(args.out)
    wb.save(out)

    total = sum(counts.values())
    missing_hi = sum(1 for rows in sheets.values() for r in rows if r[1] and not str(r[3]).strip() and not r[0].startswith("server:"))
    print(f"Wrote {out}")
    for name, n in counts.items():
        print(f"  {name:<20}{n:>6} sentences")
    print(f"  {'TOTAL':<20}{total:>6} sentences ({missing_hi} without Hindi today)")


if __name__ == "__main__":
    main()
