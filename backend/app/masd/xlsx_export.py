"""The MASD report as an Excel workbook.

Sheet 1 is the learner-level MASD file the analysts' scripts used to build
from a raw export — one row per learner with their training batch and
follow-up, adoptions by type against the targets in force, activities by
subtype against what was expected, activity on their own cases, and
nil-activity days. The other sheets hold the tables behind the dashboard, so
anyone can re-cut them. Pregnancies needing follow-up appear as counts per
learner: the workbook carries no mother or child identifiers.
"""
from __future__ import annotations

from io import BytesIO
from typing import Any, Dict, List, Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.masd import rules as R
from app.masd.insights import asked as _asked

HEAD_FILL = PatternFill("solid", fgColor="E85D4C")
HEAD_FONT = Font(bold=True, color="FFFFFF")
THIN = Side(style="thin", color="E4DDD3")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
GOOD, OK, LOW = (PatternFill("solid", fgColor=c) for c in ("DDF1E3", "FBEBC8", "F8D7D3"))


def _tint(v: Optional[float]) -> Optional[PatternFill]:
    if v is None:
        return None
    return GOOD if v >= 80 else OK if v >= 60 else LOW


def _sheet(wb: Workbook, title: str, header: List[str], rows: List[List[Any]], widths: Optional[List[int]] = None,
           pct_cols: tuple = (), note: Optional[str] = None):
    ws = wb.create_sheet(title)
    start = 1
    if note:
        ws.cell(row=1, column=1, value=note).font = Font(italic=True, color="6B6357")
        start = 3
    for c, h in enumerate(header, start=1):
        cell = ws.cell(row=start, column=c, value=h)
        cell.fill, cell.font, cell.border = HEAD_FILL, HEAD_FONT, BORDER
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    for r, row in enumerate(rows, start=start + 1):
        for c, v in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.border = BORDER
            if c in pct_cols and isinstance(v, (int, float)):
                fill = _tint(v)
                if fill:
                    cell.fill = fill
                cell.number_format = "0.0"
    for i, w in enumerate(widths or [], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = ws.cell(row=start + 1, column=2)
    ws.auto_filter.ref = f"A{start}:{get_column_letter(len(header))}{start + len(rows)}"
    ws.row_dimensions[start].height = 32
    return ws


def build_workbook(report: Dict[str, Any]) -> BytesIO:
    wb = Workbook()
    wb.remove(wb.active)
    p, cal = report["project"], report["calendar"]
    tn = cal["targets"]["now"]
    note = (f"NurtureHUB MASD · {p['name']} · as of {report['as_of']} · targets in force: ANC {tn['anc']}, "
            f"PNC <5M {tn['pnc_lt5']}, PNC ≥5M {tn['pnc_ge5']}, Staff Nurse {tn['nurse']} · each tranche's follow-up = "
            f"days since it opened (or the learner's batch ended, if later) − its buffer, rounded down to "
            f"{cal['step_days']}; expected = per tranche, the adoptions asked for × the table at its follow-up")
    steps = [u["step"] for u in report.get("uptake") or []]
    types = list(R.ADOPTION_TYPES)
    tl = {t: R.ADOPTION_TYPE_LABELS[t] for t in types}

    header = (["Learner", "Email", "Block", "Designation", "Role group", "Department", "F2F", "MT/FL",
               "Batch", "Training ended", "Follow-up (days)", "Follow-up counted (days)"]
              + [f"{tl[t]} target" for t in types] + [f"{tl[t]} adopted" for t in types]
              + ["Not typed", "Total adoptions", "Target", "% target fulfilled",
                 "Antenatal care", "Protein count", "Growth monitoring", "Breastfeeding", "Complementary feeding",
                 "Total activities", "Expected activities", "% of expected activity"]
              + [f"% expected activity · {tl[t]}" for t in types]
              + ["Own cases: expected", "Own cases: done", "% own-case activity", "Own-case band",
                 "Avg activities / adoption", "Last activity", "Nil-activity days",
                 "Follow-up counted, by tranche (days)"]
              + [h for n in steps for h in (f"Tranche {n}: days to first adoption",
                                            f"Tranche {n}: days to take all")])
    band_label = dict(R.OWN_BANDS)
    rows = []
    for l in sorted(report["learners"], key=lambda x: (x["block"], x["name"] or "")):
        a, act, own = l["adoptions"], l["activities"], l["own"]
        rows.append([l["name"], l["email"], l["block"], l["role"], l["role_group"], l["department"],
                     "Yes" if l["f2f"] else "No",
                     {"master_trainer": "Master Trainer", "facilitator": "Facilitator"}.get(l["trainer_role"] or "", ""),
                     l["batch"] or "", l["training_end"], l["fu_raw"], l["fu_days"]]
                    + [l["by_type"][t]["target"] for t in types] + [l["by_type"][t]["adopted"] for t in types]
                    + [a["unknown"], a["total"], l["target"], l["fulfilment_pct"],
                       act["anc"], act["protein"], act["gm"], act["bf"], act["cf"], act["total"], l["ideal"]["total"],
                       l["intensity_pct"]]
                    + [l["by_type"][t]["activity_pct"] for t in types]
                    + [own["expected"]["total"], own["actual"]["total"], own["pct"], band_label.get(own["band"] or "", ""),
                       l["avg_per_adoption"], l["last_activity"], l["nil_days"],
                       " · ".join(str(fu) for fu in l.get("tranche_fu") or [])]
                    + [v for n in steps for v in _uptake_cells(l, n)])
    pct_learner = (22, 30, 31, 32, 33, 36)
    _sheet(wb, "MASD learners", header, rows,
           widths=[24, 30, 14, 20, 12, 10, 6, 14, 12, 12, 10, 10] + [8] * 8 + [7, 10] + [10] * 8 + [11] * 3
           + [10, 10, 10, 12, 10, 12, 10, 14] + [12] * (2 * len(steps)),
           pct_cols=pct_learner, note=note)

    def group_rows(groups):
        return [[g["label"], g["learners"], g["adoptions"], g["target"], g["fulfilment_pct"], g["activities"],
                 g["ideal"], g["intensity_pct"], g["own_pct"], g["nil_days_avg"]]
                + [g["subtype_pct"][k] for k in R.ACTIVITY_KEYS] for g in groups]

    g_head = ["Group", "Learners", "Adoptions", "Target", "% target", "Activities", "Expected", "% of expected",
              "% own-case activity", "Avg nil-activity days"] + [f"% {R.ACTIVITY_LABELS[k]}" for k in R.ACTIVITY_KEYS]
    _sheet(wb, "By block", g_head, group_rows(report["blocks"]), widths=[18] + [11] * 14,
           pct_cols=(5, 8, 9, 11, 12, 13, 14, 15))
    _sheet(wb, "By cadre", g_head, group_rows(report["roles"] + report["departments"]), widths=[18] + [11] * 14,
           pct_cols=(5, 8, 9, 11, 12, 13, 14, 15))

    # Adoption type: target met vs expected activity done, per block and department.
    t_head = ["Group", "Learners"]
    for t in types:
        t_head += [f"{tl[t]} target", f"{tl[t]} adopted", f"{tl[t]} % target", f"{tl[t]} expected",
                   f"{tl[t]} done", f"{tl[t]} % activity"]
    t_rows = []
    for g in report["blocks"] + report["departments"] + report["roles"]:
        row = [g["label"], g["learners"]]
        for t in types:
            v = g["by_type"][t]
            row += [v["target"], v["adopted"], v["adoption_pct"], v["expected"], v["actual"], v["activity_pct"]]
        t_rows.append(row)
    _sheet(wb, "By adoption type", t_head, t_rows, widths=[18, 9] + [10] * 18,
           pct_cols=tuple(c for i in range(3) for c in (5 + 6 * i, 8 + 6 * i)))

    b_head = ["Group", "Learners banded"] + [band_label[k] for k in R.OWN_BAND_KEYS] + ["% own-case activity"]
    _sheet(wb, "Own-case bands", b_head,
           [[g["label"], g["own_banded"]] + [g["own_bands"][k] for k in R.OWN_BAND_KEYS] + [g["own_pct"]]
            for g in report["departments"] + report["blocks"] + report["roles"]],
           widths=[18, 10] + [11] * 6 + [12], pct_cols=(9,),
           note="Learners with at least one adoption, by the share of their own cases' expected activity they have done")

    _sheet(wb, "Batches", ["Batch", "Training ended", "Learners", "Follow-up (days)", "Counted as (days)",
                           "Counted, by tranche (days)"]
           + [f"Expected {R.ACTIVITY_LABELS[k]}" for k in R.ACTIVITY_KEYS] + ["Expected total", "Staff Nurse total"],
           [[b["name"] or "Project training date", b["end_date"], b["learners"], b["fu_raw"], b["fu_days"],
             " · ".join(str(fu) for fu in b["tranche_fu"])]
            + [b["expected"]["community"]["forms"][k] for k in R.ACTIVITY_KEYS]
            + [b["expected"]["community"]["total"], b["expected"]["nurse"]["total"]] for b in cal["batches"]],
           widths=[22, 13, 9, 11, 11, 14] + [12] * 7,
           note="What one learner of each batch is expected to have done by the report date, tranche by tranche")

    uptake = {u["step"]: u for u in report.get("uptake") or []}
    keys = ["before", "d7", "d15", "d30", "later", "not_yet"]
    key_label = {"before": "before it opened", "d7": "within 7 d", "d15": "8–15 d", "d30": "16–30 d",
                 "later": "after 30 d", "not_yet": "not yet"}
    tr_rows = []
    for tr in cal.get("tranches") or []:
        u = uptake.get(tr["step"]) or {}
        tr_rows.append([f"Tranche {tr['step']}", tr["from"], _asked(tr["added"]), tr["buffer"], tr["clock_from"],
                        tr["fu_raw"], tr["fu_days"]]
                       + [tr["community"]["forms"][k] for k in R.ACTIVITY_KEYS]
                       + [tr["community"]["total"], tr["nurse"]["total"],
                          u.get("learners"), u.get("started"), u.get("started_pct"), u.get("median_days_to_start"),
                          u.get("completed"), u.get("completed_pct"), u.get("median_days_to_complete")]
                       + [(u.get("started_by") or {}).get(k) for k in keys])
    _sheet(wb, "Tranches", ["Tranche", "Opened", "Asked for", "Buffer (days)", "Follow-up counts from",
                            "Follow-up (days)", "Counted as (days)"]
           + [f"Expected {R.ACTIVITY_LABELS[k]}" for k in R.ACTIVITY_KEYS]
           + ["Expected per learner", "Staff Nurse", "Learners asked", "Started", "% started", "Median days to start",
              "Took all of it", "% took all", "Median days to take all"]
           + [f"First adoption {key_label[k]}" for k in keys],
           tr_rows, widths=[11, 12, 34, 9, 13, 10, 10] + [11] * 5 + [11] * 9 + [12] * 6, pct_cols=(17, 20),
           note="Each tranche of adoptions with its own follow-up (from the project's training date), what one learner "
                "is expected to have done for it, and how quickly learners took it up once it opened")

    table = report["rules"]["expected"]["table"]
    _sheet(wb, "Expected forms", ["Adoption type", "Form"] + [f"{d} d" for d in table["durations"]],
           [[R.ADOPTION_TYPE_LABELS.get(t, "Staff Nurse <5M"), R.ACTIVITY_LABELS[k]] + table["rows"][t][k]
            for t, k in R.EXPECTED_ROWS],
           widths=[16, 20] + [6] * len(table["durations"]),
           note="Cumulative forms one adoption is expected to generate after N days of follow-up")

    fl = report["flags"]
    _sheet(wb, "Pregnancy flags", ["Learner", "Block", "Pregnancies flagged"],
           [[x["name"], x["block"], x["n"]] for x in fl["by_learner"]], widths=[26, 16, 12],
           note=(f"{fl['summary']['open_pregnancies']} pregnant women under follow-up · {fl['summary']['edd_passed']} past "
                 f"the due date · {fl['summary']['anc_behind']} behind on antenatal checks · {fl['summary']['no_lmp']} "
                 f"without LMP (mothers are listed on the dashboard, not in this file)"))

    cmp = report.get("comparison")
    if cmp:
        _sheet(wb, "Then vs now", ["Block", f"% target {cmp['then']['as_of']}", f"% target {cmp['now']['as_of']}",
                                   "Change", f"% expected {cmp['then']['as_of']}", f"% expected {cmp['now']['as_of']}",
                                   "Change"],
               [[b["label"], b["fulfilment_then"], b["fulfilment_now"], b["fulfilment_change"],
                 b["intensity_then"], b["intensity_now"], b["intensity_change"]] for b in cmp["blocks"]],
               widths=[18] + [14] * 6)

    o = report["outcomes"]
    ex = o["exclusions"]
    _sheet(wb, "Exclusions", ["Reason", "MT+FL", "Other", "Total"],
           [[x["label"], x["mtfl"], x["other"], x["total"]] for x in ex["reasons"]]
           + [["Included (final analysis)", ex["included"]["mtfl"], ex["included"]["other"], ex["included"]["total"]],
              ["Total", ex["total"]["mtfl"], ex["total"]["other"], ex["total"]["total"]]],
           widths=[52, 10, 10, 10])

    prev_rows = []
    pv = o["prevalence"]
    for label, block in (("Overall", pv["overall"]), ("MT+FL", pv["mtfl"]), ("Other", pv["other"])):
        prev_rows += _prev(label, "All ages", block)
    for band in (R.BAND_LT6, R.BAND_6_11):
        for label in ("all", "mtfl", "other"):
            prev_rows += _prev({"all": "All", "mtfl": "MT+FL", "other": "Other"}[label], R.BAND_LABELS[band],
                               pv["bands"][band][label])
        for yn in ("yes", "no"):
            prev_rows += _prev(f"Visit rule met: {yn}", R.BAND_LABELS[band], o["compliance"][band][yn])
    _sheet(wb, "Malnutrition", ["Group", "Age band", "Indicator", "n", "BV %", "AV %", "LV %", "Abs change (pts)",
                                "Rel change %"], prev_rows, widths=[22, 14, 14, 8, 9, 9, 9, 14, 13])

    _sheet(wb, "Needs attention", ["Learner", "Block", "Cadre", "Adoptions", "% target", "% of expected", "Nil-activity days",
                                   "Reasons"],
           [[a["name"], a["block"], a["role_group"], a["adoptions"], a["fulfilment_pct"], a["intensity_pct"],
             a["nil_days"], ", ".join(a["reasons"])] for a in report["attention"]],
           widths=[24, 14, 12, 10, 10, 11, 14, 40], pct_cols=(6,))
    _sheet(wb, "Data corrections", ["Learner", "Block", "Cases", "Reasons"],
           [[f["name"], f["block"], f["cases"], "; ".join(f"{R.EXCLUSION_LABELS[k]} ({n})" for k, n in f["reasons"].items())]
            for f in o["data_fixes"]], widths=[24, 14, 8, 90])
    _sheet(wb, "Findings", ["Section", "Tone", "Finding"],
           [[section, f["tone"], f["text"]] for section, items in report["insights"].items() for f in items],
           widths=[16, 8, 140])

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def _uptake_cells(learner: Dict[str, Any], step: int) -> List[Any]:
    """Days from tranche `step` opening to the learner's first adoption for it
    and to the one that completed it; negative = made before it opened."""
    u = next((x for x in learner.get("uptake") or [] if x["step"] == step), None)
    return [u["started_days"], u["completed_days"]] if u else [None, None]


def _prev(group: str, band: str, block: Dict[str, Any]) -> List[List[Any]]:
    if not block or not block.get("n"):
        return []
    return [[group, band, R.INDICATOR_LABELS[i], block["n"], block[i]["bv"], block[i]["av"], block[i]["lv"],
             block[i]["abs_change"], block[i]["rel_change"]] for i in R.INDICATORS]
