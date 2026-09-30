"""The MASD report as a PowerPoint deck.

Mirrors the analysts' district deck section by section — definitions, the
expected-forms reference and what each learner is expected to have done by
now, then-vs-now progress, adoption against % of expected activity by block,
cadre and adoption type, activity within learners' own cases, activity
subtypes, the pregnancies needing follow-up, the outcome cohort and
malnutrition against the NFHS benchmarks, key takeaways — with every chart a
native PowerPoint chart (editable, re-colourable) and the same "→" findings
the dashboard shows.
NurtureHUB adds three slides the scripts never had: the weekly activity trend,
the learners who need a supervisor's call, and the cases waiting on a data
correction.

`build_deck(report)` takes the dict from app.masd.report.build and returns an
in-memory .pptx.
"""
from __future__ import annotations

from datetime import date
from io import BytesIO
from typing import Any, Callable, Dict, List, Optional, Sequence

from pptx import Presentation
from pptx.chart.data import BubbleChartData, CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

from app.masd import rules as R

# ── Look ───────────────────────────────────────────────────────────────────

FONT = "Calibri"
INK = "26221C"
MUTED = "6B6357"
FAINT = "9C9189"
CORAL = "E85D4C"
CORAL_DARK = "B8412F"
SAND = "F6F1EA"
LINE = "E4DDD3"
WHITE = "FFFFFF"
THEN = "A8A29E"      # the earlier date in a then-vs-now pair
TEAL = "2A7F8F"
GOOD, OK, LOW = "2F9E56", "E0A11B", "D6453D"
SUBTYPE_COLORS = {"anc": "7C5CBF", "protein": "E0A11B", "gm": "2A7F8F", "bf": "E85D4C", "cf": "5C9E3C"}
TONE_COLOR = {"good": GOOD, "watch": "C2410C", "info": MUTED}
SLIDE_W, SLIDE_H = Inches(13.333), Inches(7.5)
MARGIN = Inches(0.55)


def _rgb(hex_: str) -> RGBColor:
    return RGBColor.from_string(hex_)


def _e(*values) -> tuple:
    """Whole EMUs. Layout maths divides lengths (half a slide, a third) and
    gets floats; a float offset or extent written into a table or chart frame
    makes PowerPoint refuse to open the whole file."""
    return tuple(Emu(int(round(v))) for v in values)


def _pct(v: Optional[float]) -> str:
    return "—" if v is None else f"{v:.1f}%"


def _num(v) -> str:
    if v is None:
        return "—"
    if isinstance(v, float) and not v.is_integer():
        return f"{v:,.1f}"
    return f"{int(round(v)):,}"


def _one(v) -> str:
    """A rate or an average, always to one decimal (97.0, not 97)."""
    return "—" if v is None else f"{v:,.1f}"


def _tint(v: Optional[float]) -> Optional[str]:
    """Pale cell fill by how good a percentage is (≥80 green, ≥60 amber, else red)."""
    if v is None:
        return None
    return "DDF1E3" if v >= 80 else "FBEBC8" if v >= 60 else "F8D7D3"


def _prev_tint(v: Optional[float]) -> Optional[str]:
    """Pale fill for a malnutrition prevalence (lower is better)."""
    if v is None:
        return None
    return "DDF1E3" if v < 10 else "FBEBC8" if v < 20 else "F8D7D3"


# ── Deck primitives ────────────────────────────────────────────────────────


class Deck:
    def __init__(self, report: Dict[str, Any]):
        self.r = report
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = SLIDE_W, SLIDE_H
        self.blank = self.prs.slide_layouts[6]
        self.n = 0
        p = report["project"]
        self.footer = f"NurtureHUB · MASD · {p['name']} · as of {_d(report['as_of'])}"

    # text
    def text(self, slide, x, y, w, h, runs, size=14, color=INK, bold=False, align=PP_ALIGN.LEFT,
             anchor=MSO_ANCHOR.TOP, font=FONT):
        box = slide.shapes.add_textbox(*_e(x, y, w, h))
        tf = box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = Emu(0)
        tf.margin_top = tf.margin_bottom = Emu(0)
        tf.vertical_anchor = anchor
        paragraphs = runs if isinstance(runs, list) else [runs]
        for i, para in enumerate(paragraphs):
            pgh = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            pgh.alignment = align
            parts = para if isinstance(para, tuple) and isinstance(para[0], tuple) else (para,)
            for part in parts:
                if isinstance(part, tuple):
                    txt, opts = part
                else:
                    txt, opts = part, {}
                run = pgh.add_run()
                run.text = txt
                f = run.font
                f.name = font
                f.size = Pt(opts.get("size", size))
                f.bold = opts.get("bold", bold)
                f.color.rgb = _rgb(opts.get("color", color))
        return box

    def slide(self, title: str, subtitle: Optional[str] = None):
        s = self.prs.slides.add_slide(self.blank)
        self.n += 1
        self.text(s, MARGIN, Inches(0.38), SLIDE_W - 2 * MARGIN, Inches(0.6), title, size=26, bold=True)
        if subtitle:
            self.text(s, MARGIN, Inches(0.98), SLIDE_W - 2 * MARGIN, Inches(0.4), subtitle, size=13, color=MUTED)
        self.text(s, MARGIN, SLIDE_H - Inches(0.42), Inches(9), Inches(0.25), self.footer, size=9, color=FAINT)
        self.text(s, SLIDE_W - MARGIN - Inches(1), SLIDE_H - Inches(0.42), Inches(1), Inches(0.25),
                  str(self.n), size=9, color=FAINT, align=PP_ALIGN.RIGHT)
        return s

    def findings(self, slide, items: List[Dict[str, str]], x, y, w, h, size=13, limit=4):
        items = [f for f in items if f.get("text")][:limit]
        if not items:
            return
        paras = [((("→ ", {"color": TONE_COLOR.get(f["tone"], MUTED), "bold": True}),
                   (f["text"], {}))) for f in items]
        box = self.text(slide, x, y, w, h, paras, size=size, color=INK)
        for pgh in box.text_frame.paragraphs:
            pgh.space_after = Pt(6)

    def card(self, slide, x, y, w, h, value: str, label: str, note: str = "", tone: str = INK):
        x, y, w, h = _e(x, y, w, h)
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
        shape.adjustments[0] = 0.08
        shape.fill.solid()
        shape.fill.fore_color.rgb = _rgb(SAND)
        shape.line.fill.background()
        shape.shadow.inherit = False
        self.text(slide, x + Inches(0.2), y + Inches(0.18), w - Inches(0.4), Inches(0.7), value,
                  size=30, bold=True, color=tone)
        self.text(slide, x + Inches(0.2), y + Inches(0.9), w - Inches(0.4), Inches(0.35), label,
                  size=13, bold=True, color=INK)
        if note:
            self.text(slide, x + Inches(0.2), y + Inches(1.22), w - Inches(0.4), h - Inches(1.3), note,
                      size=11, color=MUTED)

    def table(self, slide, x, y, w, rows: List[List[Any]], col_w: Optional[Sequence[float]] = None,
              fills: Optional[Callable[[int, int, Any], Optional[str]]] = None, size: Optional[int] = None,
              bold_rows: Sequence[int] = (), header_rows: int = 1, left_cols: Sequence[int] = (0,)):
        x, y, w = _e(x, y, w)
        n_rows, n_cols = len(rows), max(len(r) for r in rows)
        size = size or (12 if n_rows <= 8 else 11 if n_rows <= 12 else 10 if n_rows <= 16 else 9)
        row_h = Pt(size * 2.05)
        shape = slide.shapes.add_table(n_rows, n_cols, x, y, w, Emu(int(row_h * n_rows)))
        tbl = shape.table
        tbl.first_row = False
        tbl.horz_banding = False
        total = sum(col_w) if col_w else n_cols
        for i in range(n_cols):
            frac = (col_w[i] if col_w else 1) / total
            tbl.columns[i].width = Emu(int(w * frac))
        for r in range(n_rows):
            tbl.rows[r].height = Emu(int(row_h))
            for c in range(n_cols):
                value = rows[r][c] if c < len(rows[r]) else ""
                cell = tbl.cell(r, c)
                cell.margin_left = cell.margin_right = Inches(0.06)
                cell.margin_top = cell.margin_bottom = Inches(0.02)
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                is_head = r < header_rows
                fill = CORAL if is_head else (fills(r, c, value) if fills else None) or WHITE
                cell.fill.solid()
                cell.fill.fore_color.rgb = _rgb(fill)
                tf = cell.text_frame
                tf.word_wrap = True
                pgh = tf.paragraphs[0]
                pgh.alignment = PP_ALIGN.LEFT if c in left_cols else PP_ALIGN.CENTER
                run = pgh.add_run()
                run.text = "" if value is None else str(value)
                run.font.name = FONT
                run.font.size = Pt(size)
                run.font.bold = is_head or r in bold_rows
                run.font.color.rgb = _rgb(WHITE if is_head else INK)
        return shape

    def chart(self, slide, kind, x, y, w, h, categories: List[str], series: List[tuple],
              colors: Sequence[str], title: Optional[str] = None, max_value: Optional[float] = None,
              number_format: str = '0.0', legend: bool = True, labels: bool = True, label_size=9):
        data = CategoryChartData()
        data.categories = categories
        for name, values in series:
            data.add_series(name, [None if v is None else float(v) for v in values], number_format=number_format)
        chart = slide.shapes.add_chart(kind, *_e(x, y, w, h), data).chart
        chart.font.name = FONT
        chart.font.size = Pt(10)
        chart.font.color.rgb = _rgb(MUTED)
        chart.has_title = bool(title)
        if title:
            chart.chart_title.text_frame.text = title
            tp = chart.chart_title.text_frame.paragraphs[0]
            tp.runs[0].font.size = Pt(12)
            tp.runs[0].font.bold = True
            tp.runs[0].font.color.rgb = _rgb(INK)
        chart.has_legend = legend and len(series) > 1
        if chart.has_legend:
            chart.legend.position = XL_LEGEND_POSITION.BOTTOM
            chart.legend.include_in_layout = False
            chart.legend.font.size = Pt(10)
        plot = chart.plots[0]
        if kind in (XL_CHART_TYPE.COLUMN_CLUSTERED, XL_CHART_TYPE.BAR_CLUSTERED):
            plot.gap_width = 70
            plot.overlap = -8
        for s, color in zip(plot.series, colors):
            if kind in (XL_CHART_TYPE.LINE_MARKERS, XL_CHART_TYPE.LINE):
                s.format.line.color.rgb = _rgb(color)
                s.format.line.width = Pt(2.25)
                s.smooth = False
                s.marker.format.fill.solid()
                s.marker.format.fill.fore_color.rgb = _rgb(color)
                s.marker.format.line.color.rgb = _rgb(color)
            else:
                s.format.fill.solid()
                s.format.fill.fore_color.rgb = _rgb(color)
        if labels:
            plot.has_data_labels = True
            dl = plot.data_labels
            dl.number_format = number_format
            dl.number_format_is_linked = False
            dl.font.size = Pt(label_size)
            dl.font.color.rgb = _rgb(INK)
            if kind in (XL_CHART_TYPE.COLUMN_CLUSTERED, XL_CHART_TYPE.BAR_CLUSTERED):
                dl.position = XL_LABEL_POSITION.OUTSIDE_END
        va = chart.value_axis
        va.has_major_gridlines = True
        va.major_gridlines.format.line.color.rgb = _rgb(LINE)
        va.format.line.fill.background()
        va.tick_labels.font.size = Pt(10)
        if max_value:
            va.maximum_scale = max_value
        va.minimum_scale = 0
        ca = chart.category_axis
        ca.tick_labels.font.size = Pt(10)
        ca.format.line.color.rgb = _rgb(LINE)
        return chart

    def bubble(self, slide, x, y, w, h, points: List[Dict[str, Any]], x_title: str, y_title: str):
        data = BubbleChartData()
        for p in points:
            s = data.add_series(p["label"])
            s.add_data_point(p["x"], p["y"], max(p["size"], 1))
        chart = slide.shapes.add_chart(XL_CHART_TYPE.BUBBLE, *_e(x, y, w, h), data).chart
        chart.font.name = FONT
        chart.font.size = Pt(10)
        chart.has_legend = False
        plot = chart.plots[0]
        plot.bubble_scale = 32
        palette = [CORAL, TEAL, "7C5CBF", "E0A11B", "5C9E3C", "C2410C", "3B6EA8", "8A6E4B", "B83280", "4B8F8C"]
        for i, s in enumerate(plot.series):
            s.format.fill.solid()
            s.format.fill.fore_color.rgb = _rgb(palette[i % len(palette)])
        plot.has_data_labels = True
        dl = plot.data_labels
        dl.show_series_name = True
        dl.show_value = False
        dl.font.size = Pt(10)
        dl.font.color.rgb = _rgb(INK)
        for axis, title in ((chart.category_axis, x_title), (chart.value_axis, y_title)):
            axis.has_title = True
            axis.axis_title.text_frame.text = title
            axis.axis_title.text_frame.paragraphs[0].runs[0].font.size = Pt(11)
            axis.axis_title.text_frame.paragraphs[0].runs[0].font.color.rgb = _rgb(MUTED)
            axis.tick_labels.font.size = Pt(10)
            axis.has_major_gridlines = True
            axis.major_gridlines.format.line.color.rgb = _rgb(LINE)
        xs = [p["x"] for p in points]
        ys = [p["y"] for p in points]
        if xs:
            chart.category_axis.minimum_scale = max(0, (min(xs) // 10) * 10 - 10)
            chart.category_axis.maximum_scale = (max(xs) // 10) * 10 + 20
        if ys:
            chart.value_axis.minimum_scale = max(0, (min(ys) // 10) * 10 - 10)
            chart.value_axis.maximum_scale = (max(ys) // 10) * 10 + 20
        return chart

    def save(self) -> BytesIO:
        buffer = BytesIO()
        self.prs.save(buffer)
        buffer.seek(0)
        return buffer


def _d(iso: Optional[str]) -> str:
    if not iso:
        return "—"
    return date.fromisoformat(iso).strftime("%d %b %Y")


# ── Slides ─────────────────────────────────────────────────────────────────


def _title(d: Deck):
    r = d.r
    s = d.prs.slides.add_slide(d.blank)
    d.n += 1
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
    bg.fill.solid()
    bg.fill.fore_color.rgb = _rgb(INK)
    bg.line.fill.background()
    p = r["project"]
    d.text(s, Inches(0.9), Inches(2.0), Inches(11.5), Inches(0.5), "NurtureHUB · Member Activity Summary",
           size=16, color="F2B8AE", bold=True)
    d.text(s, Inches(0.9), Inches(2.6), Inches(11.5), Inches(1.2), f"{p['name']}: MASD & Outcome Analysis",
           size=40, color=WHITE, bold=True)
    d.text(s, Inches(0.9), Inches(3.75), Inches(11.5), Inches(0.5), f"As of {_d(r['as_of'])}",
           size=22, color="E8E1D8")
    d.text(s, Inches(0.9), Inches(4.5), Inches(11.5), Inches(0.8),
           "Project-level report | Expected activity, adoption fulfilment & malnutrition outcomes vs NFHS",
           size=15, color="C9C0B5")
    d.text(s, Inches(0.9), Inches(6.5), Inches(11.5), Inches(0.4),
           f"Generated by NurtureHUB from its own records on {date.today():%d %b %Y}", size=11, color="9C9189")


def _definitions(d: Deck):
    r = d.r
    s = d.slide("Abbreviations & definitions")
    cal = r["calendar"]
    tn = cal["targets"]["now"]
    batches = [b for b in cal["batches"] if b["id"] is not None]
    training = (f"{len(batches)} batches, ended {_d(batches[0]['end_date'])} – {_d(batches[-1]['end_date'])}"
                if batches else f"F2F training {_d(cal['training_date'])}" + (" (inferred)" if cal["inferred"] else ""))
    rows = [
        ["Term", "Meaning"],
        ["F2F training", "Face-to-face training — the learners this report covers"],
        ["MASD", "Member Activity Summary Details"],
        ["Adoption / MCD", "A mother (and her child) taken on by a learner — Mother-Child Dyad"],
        ["ANC · PNC <5M · PNC ≥5M", "Adopted while pregnant · after birth, child under 5 months · 5 months or older"],
        ["BV · AV · LV", "Birth details · Adoption visit (baseline) · Last visit (endline)"],
        ["MT+FL", "Master Trainers + Facilitators"],
        ["Follow-up (FU)", f"Days since the learner's F2F batch ended, less a {cal['buffer_days']}-day buffer, rounded "
                           f"down to a multiple of {cal['step_days']}"],
        ["Adoption target fulfilment", f"Adoptions ÷ the targets in force (now ANC {tn['anc']} · <5M {tn['pnc_lt5']} · "
                                       f"≥5M {tn['pnc_ge5']}; Staff Nurse {tn['nurse']} PNC <5M)"],
        ["% of expected activity", "Activities done ÷ activities expected: the targets × the forms each adoption "
                                   "type should generate at the learner's follow-up"],
        ["Own-case activity", "Activities done ÷ what the learner's own cases were due by their actual follow-up"],
        ["Nil-activity days", "Days since the learner's last recorded activity"],
        ["Programme calendar", training],
    ]
    d.table(s, MARGIN, Inches(1.2), SLIDE_W - 2 * MARGIN, rows, col_w=[2.6, 8.9], size=11, left_cols=(0, 1))


_REF_DURATIONS = (15, 30, 45, 60, 75, 90, 105, 120, 150, 180, 225, 270)


def _reference(d: Deck):
    exp = d.r["rules"]["expected"]
    table = exp["table"]
    s = d.slide("Reference: forms expected per adoption, by follow-up",
                "Cumulative count of each form one adoption should have generated after N days of follow-up"
                + ("" if not exp["table_is_default"] else " (default LAP table — edit on the dashboard)"))
    cols = [i for i, dur in enumerate(table["durations"]) if dur in _REF_DURATIONS]
    rows = [["Adoption type", "Form"] + [f"{table['durations'][i]} d" for i in cols]]
    names = {R.ANC: "ANC", R.PNC_LT5: "PNC <5M", R.PNC_GE5: "PNC ≥5M", R.NURSE: "Staff Nurse <5M"}
    for atype, key in R.EXPECTED_ROWS:
        rows.append([names[atype], R.ACTIVITY_LABELS[key]] + [table["rows"][atype][key][i] for i in cols])
    d.table(s, MARGIN, Inches(1.35), SLIDE_W - 2 * MARGIN, rows, col_w=[2.0, 2.2] + [0.75] * len(cols), size=11,
            left_cols=(0, 1))
    d.findings(s, [{"tone": "info", "text": "A learner's expected activity = the adoption targets in force × these "
                                            "counts at the learner's follow-up. It rises as the follow-up lengthens, "
                                            "so every date in this deck is read against what was due by then."}],
               MARGIN, Inches(5.9), SLIDE_W - 2 * MARGIN, Inches(0.9))


def _expected_now(d: Deck):
    r = d.r
    cal = r["calendar"]
    tn = cal["targets"]["now"]
    batches = cal["batches"]
    if not batches:
        return
    s = d.slide("What each learner is expected to have done by now",
                f"Targets in force on {_d(r['as_of'])}: ANC {tn['anc']} · PNC <5M {tn['pnc_lt5']} · PNC ≥5M "
                f"{tn['pnc_ge5']} · Staff Nurse {tn['nurse']} PNC <5M")
    rows = [["Batch", "Training ended", "Learners", "Follow-up (days)", "Counted as"]
            + [R.ACTIVITY_LABELS[k] for k in R.ACTIVITY_KEYS] + ["Total", "Staff Nurse"]]
    for b in batches:
        com = b["expected"]["community"]
        rows.append([b["name"] or "Project training date", _d(b["end_date"]), b["learners"],
                     _num(b["fu_raw"]), f"{b['fu_days']} d"]
                    + [com["forms"][k] for k in R.ACTIVITY_KEYS] + [com["total"], b["expected"]["nurse"]["total"]])
    d.table(s, MARGIN, Inches(1.35), SLIDE_W - 2 * MARGIN, rows,
            col_w=[1.9, 1.5, 1, 1.2, 1.1] + [1.05] * 5 + [0.9, 1.1], size=11, bold_rows=(), left_cols=(0,))
    ex = batches[0]
    com = ex["expected"]["community"]
    parts = []
    for t, v in com["by_type"].items():
        forms = " + ".join(f"{n} {R.ACTIVITY_LABELS[k].lower()}" for k, n in v["forms"].items() if n)
        parts.append(f"{v['adoptions']} {R.ADOPTION_TYPE_LABELS[t]} → {forms or 'nothing yet'}")
    d.findings(s, [{"tone": "info", "text": f"Worked example — {ex['name'] or 'the project'} at {ex['fu_days']} days: "
                                            + "; ".join(parts) + f" = {com['total']} activities per learner."}],
               MARGIN, Inches(1.55) + Pt(11 * 2.05) * len(rows), SLIDE_W - 2 * MARGIN, Inches(1.4), size=12)


def _by_type(d: Deck):
    r = d.r
    s_ = r["summary"]
    types = list(R.ADOPTION_TYPES)
    s = d.slide("Adoption vs activity, by adoption type",
                "% of each type's adoption target met, and % of the activity those adoptions were expected to generate")
    top = max([100.0] + [s_["by_type"][t][k] or 0 for t in types for k in ("adoption_pct", "activity_pct")])
    d.chart(s, XL_CHART_TYPE.COLUMN_CLUSTERED, MARGIN, Inches(1.35), Inches(6.2), Inches(4.2),
            [R.ADOPTION_TYPE_LABELS[t] for t in types],
            [("% adoption target met", [s_["by_type"][t]["adoption_pct"] for t in types]),
             ("% expected activity done", [s_["by_type"][t]["activity_pct"] for t in types])],
            [TEAL, CORAL], max_value=(int(top // 20) + 1) * 20)
    d.findings(s, r["insights"]["by_type"], MARGIN + Inches(6.5), Inches(1.5), Inches(5.7), Inches(4.5), size=12, limit=4)

    groups = [g for g in r["blocks"] if g["learners"]]
    if not groups:
        return
    s = d.slide("Adoption type by block (data)", "Adoption target met / expected activity done — best and worst "
                                                 "block per type on activity marked ▲ ▼")
    head = ["Block", "Learners"]
    for t in types:
        head += [f"{R.ADOPTION_TYPE_LABELS[t]} adopt.", f"{R.ADOPTION_TYPE_LABELS[t]} activity"]
    rows = [head]
    marks = {}
    for t in types:
        ranked = [g for g in groups if g["by_type"][t]["activity_pct"] is not None]
        if len(ranked) >= 2:
            marks[(max(ranked, key=lambda g: g["by_type"][t]["activity_pct"])["key"], t)] = " ▲"
            marks[(min(ranked, key=lambda g: g["by_type"][t]["activity_pct"])["key"], t)] = " ▼"
    for g in groups + [dict(r["summary"], key="total", label="Total")]:
        row = [g["label"], g["learners"]]
        for t in types:
            v = g["by_type"][t]
            row += [_pct(v["adoption_pct"]), _pct(v["activity_pct"]) + marks.get((g["key"], t), "")]
        rows.append(row)

    def fill(r_, c, v):
        if c >= 3 and c % 2 == 1 and isinstance(v, str) and "%" in v:
            return _tint(float(v.split("%")[0]))
        return None
    d.table(s, MARGIN, Inches(1.35), SLIDE_W - 2 * MARGIN, rows, col_w=[2.2, 1] + [1.35] * 6, fills=fill,
            bold_rows=[len(rows) - 1])


def _own_bands(d: Deck):
    r = d.r
    depts = [g for g in r["departments"] if g["own_banded"]]
    if not depts:
        return
    labels = dict(R.OWN_BANDS)
    s = d.slide("Within the cases they adopted: how active are learners?",
                "Share of learners by the % of their own cases' expected activity they have done (by department)")
    cats = [f"{g['label'].replace(' subtotal', '')} (n={g['own_banded']})" for g in depts]
    series = [(labels[k], [round(100.0 * g["own_bands"][k] / g["own_banded"], 1) for g in depts]) for k in R.OWN_BAND_KEYS]
    colors = ["B8412F", "D6453D", "E0A11B", "C9B458", "7FB069", "2F9E56"]
    d.chart(s, XL_CHART_TYPE.COLUMN_CLUSTERED, MARGIN, Inches(1.35), Inches(8), Inches(4.4), cats, series, colors,
            max_value=100, label_size=8)
    d.findings(s, r["insights"]["own_cases"], MARGIN + Inches(8.3), Inches(1.5), Inches(3.9), Inches(4.5), size=12, limit=3)
    rows = [["Department", "Learners"] + [labels[k] for k in R.OWN_BAND_KEYS] + ["Own-case activity"]]
    for g in depts:
        rows.append([g["label"].replace(" subtotal", ""), g["own_banded"]]
                    + [f"{100.0 * g['own_bands'][k] / g['own_banded']:.0f}%" for k in R.OWN_BAND_KEYS] + [_pct(g["own_pct"])])
    d.table(s, MARGIN, Inches(5.9), SLIDE_W - 2 * MARGIN, rows, col_w=[2.2, 1] + [1.15] * 6 + [1.5], size=10,
            bold_rows=[len(rows) - 1])


def _flags(d: Deck):
    f = d.r["flags"]
    sm = f["summary"]
    if not sm["open_pregnancies"]:
        return
    s = d.slide("Pregnancies that need following up",
                "Due date (LMP + 280 days) passed with no birth entered, fortnightly antenatal checks missed, or no LMP")
    cards = [(f"{sm['open_pregnancies']:,}", "Pregnant women under follow-up", "", INK),
             (f"{sm['edd_passed']:,}", "Past the due date", "no birth recorded yet", LOW if sm["edd_passed"] else INK),
             (f"{sm['anc_behind']:,}", "Antenatal checks behind", "2+ fortnightly checks missed", OK if sm["anc_behind"] else INK),
             (f"{sm['no_lmp']:,}", "No LMP recorded", "due date cannot be tracked", INK)]
    w = (SLIDE_W - 2 * MARGIN - Inches(0.3) * 3) / 4
    for i, (v, label, note, color) in enumerate(cards):
        d.card(s, MARGIN + i * (w + Inches(0.3)), Inches(1.35), w, Inches(1.6), v, label, note, color)
    rows = [["Learner", "Block", "Pregnancies flagged"]] + [[x["name"], x["block"], x["n"]] for x in f["by_learner"][:10]]
    if len(rows) > 1:
        d.table(s, MARGIN, Inches(3.25), Inches(6.4), rows, col_w=[3, 2, 1.4], size=11)
    d.findings(s, d.r["insights"]["flags"], MARGIN + Inches(6.8), Inches(3.3), Inches(5.4), Inches(3.5), size=12, limit=4)


def _headline(d: Deck):
    r = d.r
    s = d.slide("At a glance", f"{r['summary']['learners']} F2F-trained learners · as of {_d(r['as_of'])}")
    sm = r["summary"]
    ex = r["outcomes"]["exclusions"]
    tone = lambda v: GOOD if (v or 0) >= 80 else OK if (v or 0) >= 60 else LOW  # noqa: E731
    cards = [
        (_pct(sm["fulfilment_pct"]), "Adoption target met", f"{sm['adoptions']:,} of {sm['target']:,} adoptions",
         GOOD if (sm["fulfilment_pct"] or 0) >= 100 else OK),
        (_pct(sm["intensity_pct"]), "Expected activity done", f"{sm['activities']:,} of {sm['ideal']:,} expected by now",
         tone(sm["intensity_pct"])),
        ("—" if sm["nil_days_avg"] is None else f"{sm['nil_days_avg']:.1f}", "Avg nil-activity days",
         "days since the last activity", INK),
        (f"{ex['included']['total']:,}", "Dyads in outcome analysis", f"{_pct(ex['inclusion_pct'])} of eligible cases",
         INK),
    ]
    w = (SLIDE_W - 2 * MARGIN - Inches(0.3) * 3) / 4
    for i, (v, label, note, color) in enumerate(cards):
        d.card(s, MARGIN + i * (w + Inches(0.3)), Inches(1.5), w, Inches(1.75), v, label, note, color)
    d.findings(s, r["insights"]["overview"], MARGIN, Inches(3.6), SLIDE_W - 2 * MARGIN, Inches(3.2), limit=5)


def _progress(d: Deck):
    cmp = d.r.get("comparison")
    if not cmp:
        return
    t, n = cmp["then"], cmp["now"]
    s = d.slide(f"Progress: {_d(t['as_of'])} vs {_d(n['as_of'])}")
    rows = [
        ["Metric", _d(t["as_of"]), _d(n["as_of"])],
        ["Follow-up counted", f"{t['fu_days'] or 0} days", f"{n['fu_days'] or 0} days"],
        ["Learners (F2F)", t["learners"], n["learners"]],
        ["Adoption target basis", f"{t['target_per_learner']} adoptions / learner", f"{n['target_per_learner']} adoptions / learner"],
        ["Adoptions completed", f"{t['adoptions']:,} ({_pct(t['fulfilment_pct'])} of {t['target']:,})",
         f"{n['adoptions']:,} ({_pct(n['fulfilment_pct'])} of {n['target']:,})"],
        ["Expected activity done", _pct(t["intensity_pct"]), _pct(n["intensity_pct"])],
        ["Activity on own cases", _pct(t["own_pct"]), _pct(n["own_pct"])],
        ["Avg nil-activity days", _one(t["nil_days_avg"]), _one(n["nil_days_avg"])],
        ["Outcome cohort (final n)", f"{t['cohort']} ({t['cohort_lt6']} <6m + {t['cohort_6_11']} 6–11m)",
         f"{n['cohort']} ({n['cohort_lt6']} <6m + {n['cohort_6_11']} 6–11m)"],
    ]
    d.table(s, MARGIN, Inches(1.3), SLIDE_W - 2 * MARGIN, rows, col_w=[3.2, 4.1, 4.1], size=13)
    d.findings(s, d.r["insights"]["progress"], MARGIN, Inches(5.0), SLIDE_W - 2 * MARGIN, Inches(1.8), limit=3)


def _progress_blocks(d: Deck):
    cmp = d.r.get("comparison")
    if not cmp or not cmp["blocks"]:
        return
    t, n = cmp["then"], cmp["now"]
    blocks = cmp["blocks"]
    for metric, title, key_then, key_now, key_change in (
        ("adoption", "Adoptions: then vs now — by block (% of target met)", "fulfilment_then", "fulfilment_now",
         "fulfilment_change"),
        ("intensity", "Expected activity done: then vs now — by block", "intensity_then", "intensity_now",
         "intensity_change"),
    ):
        s = d.slide(title)
        rows = [["Block", f"{_d(t['as_of'])} (%)", f"{_d(n['as_of'])} (%)", "Change (pts)"]]
        for b in blocks:
            ch = b[key_change]
            rows.append([b["label"], _one(b[key_then]), _one(b[key_now]), "—" if ch is None else f"{ch:+.1f}"])
        d.table(s, MARGIN, Inches(1.3), Inches(5.3), rows, col_w=[2.2, 1.3, 1.3, 1.2],
                fills=lambda r_, c, v: (("DDF1E3" if str(v).startswith("+") else "F8D7D3") if c == 3 and v != "—" else None))
        d.chart(s, XL_CHART_TYPE.COLUMN_CLUSTERED, MARGIN + Inches(5.6), Inches(1.2), Inches(6.7), Inches(3.9),
                [b["label"] for b in blocks],
                [(_d(t["as_of"]), [b[key_then] for b in blocks]), (_d(n["as_of"]), [b[key_now] for b in blocks])],
                [THEN, CORAL], label_size=8)
        items = [f for f in d.r["insights"]["progress"]
                 if (metric == "adoption" and "fulfilment" in f["text"].lower())
                 or (metric == "intensity" and "expected activity" in f["text"].lower())]
        d.findings(s, items or d.r["insights"]["progress"][:1], MARGIN + Inches(5.6), Inches(5.3), Inches(6.7),
                   Inches(1.5), size=12, limit=2)


def _progress_nurses_uw(d: Deck):
    cmp = d.r.get("comparison")
    if not cmp:
        return
    t, n = cmp["then"], cmp["now"]
    if t["nurses"] and n["nurses"]:
        s = d.slide("Staff Nurse / Nursing Staff: then vs now")
        tn, nn = t["nurses"], n["nurses"]

        def ch(a, b):
            return "—" if a is None or b is None else f"{b - a:+.1f} pts"

        rows = [["Metric", _d(t["as_of"]), _d(n["as_of"]), "Change"],
                ["Learners", tn["learners"], nn["learners"], "—" if tn["learners"] == nn["learners"] else nn["learners"] - tn["learners"]],
                ["Adoption fulfilment", _pct(tn["fulfilment_pct"]), _pct(nn["fulfilment_pct"]), ch(tn["fulfilment_pct"], nn["fulfilment_pct"])],
                ["Activity intensity vs ideal", _pct(tn["intensity_pct"]), _pct(nn["intensity_pct"]), ch(tn["intensity_pct"], nn["intensity_pct"])],
                ["Avg nil-activity days", _one(tn["nil_days_avg"]), _one(nn["nil_days_avg"]),
                 ch(tn["nil_days_avg"], nn["nil_days_avg"]).replace(" pts", " days")]]
        d.table(s, MARGIN, Inches(1.3), SLIDE_W - 2 * MARGIN, rows, col_w=[3.5, 2.5, 2.5, 2.5], size=14)
        nurse_f = [f for f in d.r["insights"]["roles"] if "Staff Nurse" in f["text"] or "Nursing" in f["text"]]
        d.findings(s, nurse_f, MARGIN, Inches(4.2), SLIDE_W - 2 * MARGIN, Inches(2.2))
    # Underweight AV → LV by band, interim vs final.
    s = d.slide("Underweight reduction: interim vs final", "Share of children below −2 SD weight-for-age, adoption visit → last visit")
    rows = [["Age band", "Report", "n", "AV (baseline)", "LV (endline)", "Abs change"]]
    cats, av, lv = [], [], []
    for band in (R.BAND_LT6, R.BAND_6_11):
        for label, snap in (("Interim", t), ("Final", n)):
            u = snap["underweight"][band]
            rows.append([R.BAND_LABELS[band], f"{label} ({_d(snap['as_of'])})", u["n"], _pct(u["av"]), _pct(u["lv"]),
                         "—" if u["abs_change"] is None else f"{u['abs_change']:+.1f} pts"])
            cats.append(f"{R.BAND_LABELS[band]} {label.lower()}")
            av.append(u["av"])
            lv.append(u["lv"])
    d.table(s, MARGIN, Inches(1.4), Inches(6.6), rows, col_w=[1.4, 2.2, 0.7, 1.2, 1.2, 1.2], size=12)
    d.chart(s, XL_CHART_TYPE.COLUMN_CLUSTERED, MARGIN + Inches(6.9), Inches(1.3), Inches(5.3), Inches(4.2), cats,
            [("AV (baseline)", av), ("LV (endline)", lv)], [THEN, CORAL])


def _distribution(d: Deck):
    r = d.r
    groups = [g for g in R.ROLE_GROUPS if any(g in v for v in r["distribution"].values())]
    if not groups:
        return
    s = d.slide("Distribution of learners across blocks & role groups",
                "Row-wise shares: each cell is the share of that block's F2F learners")
    rows = [["Block"] + groups + ["Total"]]
    total_n = sum(sum(v.values()) for v in r["distribution"].values())
    col_tot = {g: 0 for g in groups}
    for block in sorted(r["distribution"]):
        counts = r["distribution"][block]
        n = sum(counts.values())
        row = [block]
        for g in groups:
            c = counts.get(g, 0)
            col_tot[g] += c
            row.append(f"{c} ({100 * c / n:.0f}%)" if c else "")
        row.append(f"{n} ({100 * n / total_n:.0f}%)")
        rows.append(row)
    rows.append(["Total"] + [f"{col_tot[g]} ({100 * col_tot[g] / total_n:.0f}%)" for g in groups] + [f"{total_n} (100%)"])
    d.table(s, MARGIN, Inches(1.4), SLIDE_W - 2 * MARGIN, rows, col_w=[2] + [1.2] * len(groups) + [1.4],
            bold_rows=[len(rows) - 1])


def _adoption_activity(d: Deck, key: str):
    r = d.r
    groups = [g for g in (r["roles"] if key == "roles" else r["blocks"]) if g["learners"]]
    if not groups:
        return
    word = "role group" if key == "roles" else "block"
    s = d.slide(f"Adoption & activity fulfilment by {word}",
                "% of the adoption target met and % of the expected activities done — 100% is the target line")
    top = max([100.0] + [g["fulfilment_pct"] or 0 for g in groups] + [g["intensity_pct"] or 0 for g in groups])
    d.chart(s, XL_CHART_TYPE.COLUMN_CLUSTERED, MARGIN, Inches(1.35), SLIDE_W - 2 * MARGIN, Inches(4.0),
            [f"{g['label']} (n={g['learners']})" for g in groups],
            [("% adoption target achieved", [g["fulfilment_pct"] for g in groups]),
             ("% activity achieved", [g["intensity_pct"] for g in groups])],
            [TEAL, CORAL], max_value=(int(top // 20) + 1) * 20, label_size=9)
    d.findings(s, r["insights"][key], MARGIN, Inches(5.5), SLIDE_W - 2 * MARGIN, Inches(1.4), size=12, limit=2)

    # The relationship, as bubbles.
    pts = [{"label": g["label"], "x": g["fulfilment_pct"] or 0, "y": g["intensity_pct"] or 0, "size": g["learners"]}
           for g in groups if g["fulfilment_pct"] is not None and g["intensity_pct"] is not None]
    if len(pts) >= 2:
        s = d.slide(f"The relationship: adoption target met vs expected activity done ({word}s)",
                    "Bubble size = number of learners · right = more cases adopted · up = more follow-up done")
        d.bubble(s, MARGIN, Inches(1.35), Inches(8.3), Inches(5.3), pts,
                 "Adoption target met (%) → more cases adopted", "Expected activity done (%) → more follow-up")
        d.findings(s, r["insights"][key][1:] or r["insights"][key], MARGIN + Inches(8.6), Inches(1.5), Inches(3.6),
                   Inches(5), size=12, limit=3)

    # The data.
    s = d.slide(f"Adoption & activity with nil-activity days by {word} (data)")
    rows = [["Group", "Learners", "% adoption target", "% activity", "Avg nil-activity days"]]
    for g in groups:
        rows.append([g["label"], g["learners"], _pct(g["fulfilment_pct"]), _pct(g["intensity_pct"]), _one(g["nil_days_avg"])])
    tot = r["summary"]
    rows.append(["Total", tot["learners"], _pct(tot["fulfilment_pct"]), _pct(tot["intensity_pct"]), _one(tot["nil_days_avg"])])
    d.table(s, MARGIN, Inches(1.35), SLIDE_W - 2 * MARGIN, rows, col_w=[3, 1.4, 2, 2, 2.2],
            fills=lambda r_, c, v: _tint(float(v.rstrip("%"))) if c == 3 and isinstance(v, str) and v.endswith("%") else None,
            bold_rows=[len(rows) - 1])


def _cadre_targets(d: Deck):
    r = d.r
    roles = {g["key"]: g for g in r["roles"] if g["learners"]}
    if not roles:
        return
    s = d.slide("Cadre level: target vs actual adoptions and activities")
    rows = [["Cadre", "Learners", "Adoptions target", "Actual", "% fulfilled", "Activities expected", "Actual", "% fulfilled"]]
    bold = []

    def line(label, g):
        rows.append([label, g["learners"], _num(g["target"]), _num(g["adoptions"]), _pct(g["fulfilment_pct"]),
                     _num(g["ideal"]), _num(g["activities"]), _pct(g["intensity_pct"])])

    for dept, members in ((R.WCD, ("AWW", "AWSup")), (R.HFW, ("ASHA", "ASHA Sup", "ANM", "CHO", "Nursing Staff", "Other"))):
        for key in members:
            if key in roles:
                line(key, roles[key])
        sub = _sum_groups([roles[k] for k in members if k in roles])
        if sub["learners"]:
            line(f"{dept} subtotal", sub)
            bold.append(len(rows) - 1)
    line("Grand total", r["departments"][-1])
    bold.append(len(rows) - 1)
    d.table(s, MARGIN, Inches(1.35), SLIDE_W - 2 * MARGIN, rows, col_w=[2.4, 1, 1.4, 1, 1.2, 1.6, 1, 1.2], bold_rows=bold,
            fills=lambda r_, c, v: _tint(float(v.rstrip("%"))) if c in (4, 7) and isinstance(v, str) and v.endswith("%") else None)


def _sum_groups(groups: List[Dict[str, Any]]) -> Dict[str, Any]:
    tot = {k: sum(g[k] for g in groups) for k in ("learners", "target", "adoptions", "ideal", "activities")}
    tot["fulfilment_pct"] = round(100.0 * tot["adoptions"] / tot["target"], 1) if tot["target"] else None
    tot["intensity_pct"] = round(100.0 * tot["activities"] / tot["ideal"], 1) if tot["ideal"] else None
    return tot


def _subtypes(d: Deck, key: str):
    r = d.r
    groups = [g for g in (r["roles"] if key == "roles" else r["blocks"]) if g["learners"]]
    if not groups:
        return
    word = "role group" if key == "roles" else "block"
    s = d.slide(f"Activity subtypes by {word}", "% of the expected count done, per activity type")
    series = [(R.ACTIVITY_LABELS[k], [g["subtype_pct"][k] for g in groups]) for k in R.ACTIVITY_KEYS]
    top = max([100.0] + [v for _, vals in series for v in vals if v is not None])
    d.chart(s, XL_CHART_TYPE.COLUMN_CLUSTERED, MARGIN, Inches(1.35), SLIDE_W - 2 * MARGIN, Inches(4.1),
            [f"{g['label']} (n={g['learners']})" for g in groups], series,
            [SUBTYPE_COLORS[k] for k in R.ACTIVITY_KEYS], max_value=(int(top // 20) + 1) * 20, labels=len(groups) <= 6,
            label_size=7)
    d.findings(s, r["insights"]["subtypes"], MARGIN, Inches(5.6), SLIDE_W - 2 * MARGIN, Inches(1.3), size=12, limit=2)

    s = d.slide(f"Activity subtypes by {word} (data)", "Cells tinted green ≥80%, amber 60–79%, red below 60%")
    rows = [["Group", "Learners"] + [R.ACTIVITY_LABELS[k] for k in R.ACTIVITY_KEYS]]
    for g in groups:
        rows.append([g["label"], g["learners"]] + [_one(g["subtype_pct"][k]) for k in R.ACTIVITY_KEYS])
    d.table(s, MARGIN, Inches(1.35), SLIDE_W - 2 * MARGIN, rows, col_w=[2.4, 1] + [1.6] * 5,
            fills=lambda r_, c, v: _tint(float(v)) if c >= 2 and isinstance(v, str) and v not in ("—", "") else None)


def _weekly(d: Deck):
    weeks = d.r["weekly"]
    if len(weeks) < 2:
        return
    s = d.slide("Activity week by week", "All activities filed by F2F learners each week since training (new in NurtureHUB)")
    cats = [date.fromisoformat(w["week"]).strftime("%d %b") for w in weeks]
    as_of = date.fromisoformat(d.r["as_of"])
    part = (as_of - date.fromisoformat(weeks[-1]["week"])).days < 6
    if part:
        cats[-1] += " (part)"
    d.chart(s, XL_CHART_TYPE.LINE_MARKERS, MARGIN, Inches(1.35), SLIDE_W - 2 * MARGIN, Inches(4.3), cats,
            [("Activities", [w["total"] for w in weeks]), ("New adoptions", [w["adoptions"] for w in weeks])],
            [CORAL, TEAL], number_format="0", labels=False)
    peak = max(weeks, key=lambda w: w["total"])
    last = weeks[-1]
    d.findings(s, [{"tone": "info", "text": f"Busiest week: {_d(peak['week'])} with {peak['total']:,} activities. "
                                            f"The latest {'(part) ' if part else ''}week ({_d(last['week'])}) "
                                            f"had {last['total']:,}."}],
               MARGIN, Inches(5.85), SLIDE_W - 2 * MARGIN, Inches(1))


def _attention(d: Deck):
    rows_in = d.r["attention"]
    if not rows_in:
        return
    s = d.slide("Learners who need a supervisor's call",
                f"{len(rows_in)} F2F learners — no activity, 14+ silent days, under 40% of expected activity or under half the adoptions")
    reason = {"no_activity": "No activity", "inactive": "Silent 14+ days", "low_intensity": "Low expected activity",
              "few_adoptions": "Few adoptions"}
    rows = [["Learner", "Block", "Cadre", "Adoptions", "% activity", "Nil days", "Why"]]
    for a in rows_in[:14]:
        rows.append([a["name"], a["block"], a["role_group"], a["adoptions"], _pct(a["intensity_pct"]),
                     _num(a["nil_days"]), ", ".join(reason[x] for x in a["reasons"])])
    d.table(s, MARGIN, Inches(1.35), SLIDE_W - 2 * MARGIN, rows, col_w=[2.6, 1.6, 1.3, 1, 1.1, 0.9, 3.3],
            left_cols=(0, 6))
    if len(rows_in) > 14:
        d.text(s, MARGIN, SLIDE_H - Inches(0.8), Inches(10), Inches(0.3),
               f"…and {len(rows_in) - 14} more — the full list is in the MASD dashboard and the Excel download.",
               size=11, color=MUTED)


def _feedback(d: Deck):
    ins = d.r["insights"]
    good = [f for f in ins["progress"] + ins["blocks"] + ins["roles"] + ins["subtypes"] if f["tone"] == "good"]
    watch = [f for f in ins["progress"] + ins["blocks"] + ins["roles"] + ins["subtypes"] if f["tone"] == "watch"]
    s = d.slide("Feedback: what's going well, what needs attention")
    half = (SLIDE_W - 2 * MARGIN - Inches(0.4)) / 2
    for i, (label, items, color) in enumerate((("Going well", good, GOOD), ("Needs attention", watch, "C2410C"))):
        x = MARGIN + i * (half + Inches(0.4))
        d.text(s, x, Inches(1.35), half, Inches(0.4), label, size=18, bold=True, color=color)
        d.findings(s, items or [{"tone": "info", "text": "Nothing stands out on this side."}], x, Inches(1.9), half,
                   Inches(4.9), size=13, limit=5)


def _divider(d: Deck, title: str, sub: str):
    s = d.prs.slides.add_slide(d.blank)
    d.n += 1
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
    bg.fill.solid()
    bg.fill.fore_color.rgb = _rgb(INK)
    bg.line.fill.background()
    d.text(s, Inches(0.9), Inches(3.0), Inches(11.5), Inches(1), title, size=40, bold=True, color=WHITE)
    d.text(s, Inches(0.9), Inches(4.0), Inches(11.5), Inches(0.6), sub, size=18, color="C9C0B5")


def _funnel(d: Deck):
    o = d.r["outcomes"]
    s = d.slide("Pre-analysis exclusion: learner & case funnel",
                "From every learner with a registered case to the F2F-eligible cases the outcome analysis uses")
    names = {"start": "START: every learner with a registered case", "no_f2f": "Not selected for F2F training",
             "nursing_staff": "Staff Nurse role group (analysed separately)"}
    rows = [["Step", "Learners removed", "Learners after", "Cases removed", "Cases after"]]
    for st in o["funnel"]:
        rows.append([names.get(st["key"], st["key"]), st["learners_removed"] or "", st["learners_after"],
                     st["cases_removed"] or "", st["cases_after"]])
    last = o["funnel"][-1]
    rows.append(["FINAL: eligible for the outcome analysis", "", last["learners_after"], "", last["cases_after"]])
    d.table(s, MARGIN, Inches(1.4), SLIDE_W - 2 * MARGIN, rows, col_w=[5, 1.6, 1.6, 1.6, 1.6], bold_rows=[len(rows) - 1], size=13)
    d.findings(s, [{"tone": "info", "text": "Staff Nurses' cases are facility-based and analysed separately; learners not "
                                            "selected for F2F training are outside the programme cohort."}],
               MARGIN, Inches(4.6), SLIDE_W - 2 * MARGIN, Inches(1))


def _exclusions(d: Deck):
    o = d.r["outcomes"]
    ex = o["exclusions"]
    s = d.slide("Primary exclusion reasons: MT+FL vs other learners", "Mother-child dyad level; the first failed check is the reason")
    tm, to = ex["total"]["mtfl"] or 1, ex["total"]["other"] or 1
    tt = ex["total"]["total"] or 1
    rows = [["Exclusion reason", "MT+FL", "Other", "Total"]]
    for x in ex["reasons"]:
        if not x["total"]:
            continue
        rows.append([x["label"], f"{x['mtfl']} ({100 * x['mtfl'] / tm:.1f}%)", f"{x['other']} ({100 * x['other'] / to:.1f}%)",
                     f"{x['total']} ({100 * x['total'] / tt:.1f}%)"])
    inc = ex["included"]
    rows.append(["Included (final analysis)", f"{inc['mtfl']} ({_pct(ex['inclusion_pct_mtfl'])})",
                 f"{inc['other']} ({_pct(ex['inclusion_pct_other'])})", f"{inc['total']} ({_pct(ex['inclusion_pct'])})"])
    rows.append(["Total", ex["total"]["mtfl"], ex["total"]["other"], ex["total"]["total"]])
    d.table(s, MARGIN, Inches(1.35), Inches(8.4), rows, col_w=[4.6, 1.3, 1.3, 1.3], bold_rows=[len(rows) - 2, len(rows) - 1])
    d.findings(s, d.r["insights"]["outcomes"], MARGIN + Inches(8.7), Inches(1.4), Inches(3.5), Inches(5), size=12)


def _retention(d: Deck):
    o = d.r["outcomes"]["retention"]
    if len(o["learners"]) <= 1:
        return
    s = d.slide("Retention by cadre: learners and adoptions", "How many of each cadre's learners and cases reach the final analysis")
    half = (SLIDE_W - 2 * MARGIN - Inches(0.4)) / 2
    for i, (key, head) in enumerate((("learners", "F2F learners"), ("adoptions", "Total adoptions"))):
        rows = [["Role group", head, "In final analysis", "% retained"]]
        for x in o[key]:
            rows.append([x["role_group"], x["total"], x["final"], _pct(x["pct"])])
        d.table(s, MARGIN + i * (half + Inches(0.4)), Inches(1.4), half, rows, col_w=[1.8, 1.3, 1.5, 1.2],
                bold_rows=[len(rows) - 1],
                fills=lambda r_, c, v: _tint(float(v.rstrip("%"))) if c == 3 and isinstance(v, str) and v.endswith("%") else None)


def _demographics(d: Deck):
    demo = d.r["outcomes"]["demographics"]
    if not demo["n"]:
        return
    s = d.slide(f"Mother demographic profile (n={demo['n']:,})", "Mothers of the analysed dyads")
    panels = [("Mother's age", "mother_age"), ("Ration card", "ration_card"), ("Social category", "social_category"),
              ("Delivery place", "delivery_place"), ("Delivery method", "delivery_method")]
    w = (SLIDE_W - 2 * MARGIN - Inches(0.3) * 2) / 3
    h = Inches(2.7)
    for i, (title, key) in enumerate(panels):
        items = [x for x in demo[key] if x["n"]][:6]
        if not items:
            continue
        col, row = i % 3, i // 3
        d.chart(s, XL_CHART_TYPE.BAR_CLUSTERED, MARGIN + col * (w + Inches(0.3)), Inches(1.3) + row * (h + Inches(0.15)),
                w, h, [f"{x['label']} (n={x['n']})" for x in items][::-1], [(title, [x["pct"] for x in items][::-1])],
                [CORAL], title=title, legend=False, label_size=9)
    d.text(s, MARGIN + 2 * (w + Inches(0.3)), Inches(1.3) + h + Inches(0.4), w, Inches(1.8),
           "Family type is not collected in NurtureHUB's mother registration, so it is not shown.", size=11, color=MUTED)


def _nfhs(d: Deck):
    b = d.r["outcomes"].get("benchmarks")
    if not b:
        return
    trend = b.get("district_trend")
    if trend and trend.get("rounds"):
        s = d.slide(f"NFHS rounds: {trend.get('label') or 'district trend'}", "Children under five")
        rounds = trend["rounds"]
        rows = [["Indicator"] + rounds]
        for ind in R.INDICATORS:
            rows.append([R.INDICATOR_LABELS[ind]] + [_pct(v) for v in trend[ind]])
        d.table(s, MARGIN, Inches(1.4), Inches(6.2), rows, col_w=[2] + [1.2] * len(rounds), size=14)
        d.chart(s, XL_CHART_TYPE.COLUMN_CLUSTERED, MARGIN + Inches(6.5), Inches(1.3), Inches(5.7), Inches(4.3),
                [R.INDICATOR_LABELS[i] for i in R.INDICATORS],
                [(rd, [trend[i][k] for i in R.INDICATORS]) for k, rd in enumerate(rounds)],
                [THEN, FAINT, CORAL, TEAL][: len(rounds)])
    bands = b.get("age_bands") or {}
    if bands:
        s = d.slide(f"{b.get('label') or 'NFHS'}: malnutrition by age group", "The benchmark the project's cohort is read against")
        rows = [["Indicator", R.BAND_LABELS[R.BAND_LT6], R.BAND_LABELS[R.BAND_6_11]]]
        for ind in R.INDICATORS:
            rows.append([R.INDICATOR_LABELS[ind], _pct((bands.get(R.BAND_LT6) or {}).get(ind)),
                         _pct((bands.get(R.BAND_6_11) or {}).get(ind))])
        d.table(s, MARGIN, Inches(1.4), Inches(5.5), rows, col_w=[2, 1.6, 1.6], size=14)
        d.chart(s, XL_CHART_TYPE.COLUMN_CLUSTERED, MARGIN + Inches(5.9), Inches(1.3), Inches(6.3), Inches(4.3),
                [R.INDICATOR_LABELS[i] for i in R.INDICATORS],
                [(R.BAND_LABELS[bd], [(bands.get(bd) or {}).get(i) for i in R.INDICATORS]) for bd in (R.BAND_LT6, R.BAND_6_11)],
                [TEAL, CORAL])


def _prev_rows(p: Dict[str, Any], bench: Optional[Dict[str, Any]]):
    bands = (bench or {}).get("age_bands") or {}
    head = ["Indicator"]
    if bands:
        head += [f"<6m {bench.get('label') or 'NFHS'}", f"6–11m {bench.get('label') or 'NFHS'}"]
    head += ["BV", "AV (baseline)", "LV (endline)", "Abs diff (AV→LV)", "Rel diff (AV→LV)"]
    rows = [head]
    for ind in R.INDICATORS:
        x = p[ind]
        row = [R.INDICATOR_LABELS[ind]]
        if bands:
            row += [_pct((bands.get(R.BAND_LT6) or {}).get(ind)), _pct((bands.get(R.BAND_6_11) or {}).get(ind))]
        row += [_pct(x["bv"]), _pct(x["av"]), _pct(x["lv"]),
                "—" if x["abs_change"] is None else f"{x['abs_change']:+.1f} pts",
                "—" if x["rel_change"] is None else f"{x['rel_change']:+.1f}%"]
        rows.append(row)
    return rows


def _prevalence_slide(d: Deck, title: str, p: Dict[str, Any], findings_key: str):
    if not p.get("n"):
        return
    bench = d.r["outcomes"].get("benchmarks")
    s = d.slide(f"{title} (n={p['n']:,})", "Share below −2 SD: stunting (length-for-age), underweight (weight-for-age), wasting (weight-for-length)")
    rows = _prev_rows(p, bench)
    ncol = len(rows[0])
    d.table(s, MARGIN, Inches(1.4), SLIDE_W - 2 * MARGIN, rows, col_w=[1.8] + [1.3] * (ncol - 1), size=13,
            fills=lambda r_, c, v: _prev_tint(float(v.rstrip("%"))) if c >= 1 and isinstance(v, str) and v.endswith("%")
            and "pts" not in v and not v.startswith(("+", "-")) else None)
    d.chart(s, XL_CHART_TYPE.COLUMN_CLUSTERED, MARGIN, Inches(3.3), Inches(7.2), Inches(3.4),
            [R.INDICATOR_LABELS[i] for i in R.INDICATORS],
            [("Baseline (AV)", [p[i]["av"] for i in R.INDICATORS]), ("Last visit (LV)", [p[i]["lv"] for i in R.INDICATORS])],
            [THEN, CORAL])
    d.findings(s, d.r["insights"].get(findings_key, []), MARGIN + Inches(7.5), Inches(3.5), Inches(4.7), Inches(3), size=12)


def _compliance(d: Deck):
    comp = d.r["outcomes"]["compliance"]
    for band, key in ((R.BAND_LT6, "compliance_lt6"), (R.BAND_6_11, "compliance_6_11")):
        c = comp[band]
        rule = c["rule"]
        n = (c["yes"].get("n") or 0) + (c["no"].get("n") or 0)
        if not n:
            continue
        who = "ANC + <5 months" if band == R.BAND_LT6 else "≥5 months"
        s = d.slide(f"Minimum PNC visit compliance: {who} (n={n:,})",
                    f"Compliant = {rule['min_visits']}+ growth checks with ≥{rule['min_follow_up_days']} days of follow-up")
        rows = [["Indicator", "Group", "n", "BV", "AV (baseline)", "LV (endline)"]]
        for ind in R.INDICATORS:
            for label, g in (("Yes", c["yes"]), ("No", c["no"])):
                rows.append([R.INDICATOR_LABELS[ind] if label == "Yes" else "", label, g.get("n", 0),
                             _pct(g[ind]["bv"]) if g.get("n") else "—", _pct(g[ind]["av"]) if g.get("n") else "—",
                             _pct(g[ind]["lv"]) if g.get("n") else "—"])
        d.table(s, MARGIN, Inches(1.4), Inches(6.2), rows, col_w=[1.6, 0.8, 0.8, 1, 1.3, 1.3], size=12)
        cats = [R.INDICATOR_LABELS[i] for i in R.INDICATORS]
        get = lambda g, p: [g[i][p] if g.get("n") else None for i in R.INDICATORS]  # noqa: E731
        d.chart(s, XL_CHART_TYPE.COLUMN_CLUSTERED, MARGIN + Inches(6.5), Inches(1.3), Inches(5.7), Inches(3.8), cats,
                [("Yes – AV", get(c["yes"], "av")), ("Yes – LV", get(c["yes"], "lv")),
                 ("No – AV", get(c["no"], "av")), ("No – LV", get(c["no"], "lv"))],
                ["9FC5CC", TEAL, "F2B8AE", CORAL], label_size=8)
        d.findings(s, d.r["insights"][key], MARGIN + Inches(6.5), Inches(5.3), Inches(5.7), Inches(1.4), size=12)


def _bands(d: Deck):
    bands = d.r["outcomes"]["prevalence"]["bands"]
    for band in (R.BAND_LT6, R.BAND_6_11):
        m, o = bands[band]["mtfl"], bands[band]["other"]
        if not m.get("n") and not o.get("n"):
            continue
        s = d.slide(f"Age {R.BAND_LABELS[band]}: MT+FL vs other learners (AV vs LV)")
        rows = [["Indicator", "MT+FL AV", "MT+FL LV", "MT+FL rel diff", "Other AV", "Other LV", "Other rel diff"]]
        for ind in R.INDICATORS:
            def cell(g, p):
                return f"{_pct(g[ind][p])} (n={g['n']})" if g.get("n") else "—"

            def rel(g):
                return "—" if not g.get("n") or g[ind]["rel_change"] is None else f"{g[ind]['rel_change']:+.1f}%"
            rows.append([R.INDICATOR_LABELS[ind], cell(m, "av"), cell(m, "lv"), rel(m), cell(o, "av"), cell(o, "lv"), rel(o)])
        d.table(s, MARGIN, Inches(1.35), SLIDE_W - 2 * MARGIN, rows, col_w=[1.7, 1.6, 1.6, 1.4, 1.6, 1.6, 1.4], size=12)
        get = lambda g, p: [g[i][p] if g.get("n") else None for i in R.INDICATORS]  # noqa: E731
        d.chart(s, XL_CHART_TYPE.COLUMN_CLUSTERED, MARGIN, Inches(3.2), Inches(7.4), Inches(3.5),
                [R.INDICATOR_LABELS[i] for i in R.INDICATORS],
                [("MT+FL AV", get(m, "av")), ("MT+FL LV", get(m, "lv")), ("Other AV", get(o, "av")), ("Other LV", get(o, "lv"))],
                ["9FC5CC", TEAL, "F2B8AE", CORAL], label_size=8)
        items = [f for f in d.r["insights"]["bands"] if R.BAND_LABELS[band] in f["text"]]
        d.findings(s, items, MARGIN + Inches(7.7), Inches(3.4), Inches(4.5), Inches(3.2), size=12)

    bench = (d.r["outcomes"].get("benchmarks") or {})
    bb = bench.get("age_bands") or {}
    s = d.slide(f"Age bands vs {bench.get('label') or 'NFHS'} benchmark (overall)" if bb else "Age bands: overall (AV vs LV)")
    y = Inches(1.35)
    for band in (R.BAND_LT6, R.BAND_6_11):
        a = bands[band]["all"]
        d.text(s, MARGIN, y, Inches(10), Inches(0.35), f"Age {R.BAND_LABELS[band]} (n={a.get('n', 0):,})", size=14, bold=True)
        rows = [["Indicator"] + (["NFHS"] if bb else []) + ["AV (baseline)", "LV (endline)", "Abs diff (AV→LV)", "Rel diff (AV→LV)"]]
        for ind in R.INDICATORS:
            x = a[ind]
            rows.append([R.INDICATOR_LABELS[ind]] + ([_pct((bb.get(band) or {}).get(ind))] if bb else [])
                        + [_pct(x["av"]), _pct(x["lv"]),
                           "—" if x["abs_change"] is None else f"{x['abs_change']:+.1f} pts",
                           "—" if x["rel_change"] is None else f"{x['rel_change']:+.1f}%"])
        d.table(s, MARGIN, y + Inches(0.4), SLIDE_W - 2 * MARGIN, rows, size=12)
        y += Inches(2.6)


def _data_fixes(d: Deck):
    fixes = d.r["outcomes"]["data_fixes"]
    if not fixes:
        return
    s = d.slide("Cases waiting on a data correction",
                "Excluded for a data-entry problem, not for missing follow-up — correcting them returns the case to the analysis")
    rows = [["Learner", "Block", "Cases", "What to check"]]
    for f in fixes[:14]:
        rows.append([f["name"], f["block"], f["cases"],
                     "; ".join(f"{R.EXCLUSION_LABELS[k]} ({n})" for k, n in f["reasons"].items())])
    d.table(s, MARGIN, Inches(1.4), SLIDE_W - 2 * MARGIN, rows, col_w=[2.4, 1.6, 0.8, 7], left_cols=(0, 3))


def _takeaways(d: Deck):
    s = d.slide("Key takeaways")
    d.findings(s, d.r["insights"]["takeaways"], MARGIN, Inches(1.4), SLIDE_W - 2 * MARGIN, Inches(5.4), size=16, limit=6)


def build_deck(report: Dict[str, Any]) -> BytesIO:
    d = Deck(report)
    _title(d)
    _definitions(d)
    _reference(d)
    _expected_now(d)
    _headline(d)
    _progress(d)
    _progress_blocks(d)
    _progress_nurses_uw(d)
    _distribution(d)
    _adoption_activity(d, "roles")
    _adoption_activity(d, "blocks")
    _cadre_targets(d)
    _by_type(d)
    _own_bands(d)
    _subtypes(d, "roles")
    _subtypes(d, "blocks")
    _weekly(d)
    _attention(d)
    _flags(d)
    _feedback(d)
    _divider(d, "Outcomes-centric analysis", "Mother-child dyads followed from birth to the last visit")
    _funnel(d)
    _exclusions(d)
    _retention(d)
    _demographics(d)
    _nfhs(d)
    p = report["outcomes"]["prevalence"]
    _prevalence_slide(d, "MT+FL learners' cases: malnutrition indicators", p["mtfl"], "mtfl")
    _prevalence_slide(d, "Other learners' cases", p["other"], "other")
    _compliance(d)
    _bands(d)
    _prevalence_slide(d, "Overall: all analysed cases", p["overall"], "overall")
    _data_fixes(d)
    _takeaways(d)
    return d.save()
