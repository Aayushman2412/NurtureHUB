"""Admin MASD dashboard: /api/admin/masd/*.

The analysts' district report — adoption target fulfilment, activity
intensity, and malnutrition outcomes against NFHS — computed from the app's
own data (app/masd/). The report is aggregate: it names learners, never a
mother or child, but it is derived from their records, so the routes sit
behind the PHI emergency switch and every read and download is audited.
"""
from __future__ import annotations

import io
import json
import re
from datetime import date
from typing import Dict, List, Literal, Optional

from fastapi import APIRouter, Depends, File, HTTPException, Query, Request, UploadFile
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field, model_validator
from sqlalchemy import func
from sqlalchemy.orm import Session

from app import models, projects
from app.config import settings
from app.database import get_db
from app.dependencies import get_admin_email, get_current_admin, require_phi_access
from app.masd import data as masd_data
from app.masd import report as masd_report
from app.masd import rules as R
from app.rate_limit import limiter, principal_key
from app.security import audit, phi

router = APIRouter(
    prefix="/api/admin/masd",
    tags=["admin-masd"],
    dependencies=[Depends(get_current_admin), Depends(require_phi_access)],
)


def _project(db: Session, slug: str) -> models.ProgramDistrict:
    project = db.query(models.ProgramDistrict).filter(models.ProgramDistrict.slug == slug).first()
    if project is None:
        raise HTTPException(status_code=404, detail="Project not found")
    return project


def _dates(as_of: Optional[date], compare: Optional[date]) -> tuple:
    as_of = as_of or date.today()
    if compare is not None and compare >= as_of:
        raise HTTPException(status_code=422, detail="The comparison date must be before the report date")
    return as_of, compare


def _log_read(report: dict, project: models.ProgramDistrict) -> None:
    phi.log_list(
        resource_type="masd_report",
        count=report["outcomes"]["exclusions"]["total"]["total"],
        subject_type=phi.CHILD,
        detail={"scope": "aggregate_programme_report", "project": project.slug,
                "as_of": report["as_of"], "fields": ["z_scores", "dates", "demographic_categories"]},
    )
    flagged = report["flags"]["items"]
    if flagged:
        # The follow-up flags name pregnancies by record ID, with their dates.
        phi.log_list(
            resource_type="masd_pregnancy_flags",
            count=len(flagged),
            subject_type=phi.MOTHER,
            detail={"project": project.slug, "as_of": report["as_of"],
                    "fields": ["mother_uid", "adoption_date", "lmp", "edd", "antenatal_check_dates"]},
        )


@router.get("/projects")
def masd_projects(db: Session = Depends(get_db)):
    """Projects for the picker, with how many F2F learners each has and
    whether its programme calendar has been set."""
    saved = {s.program_district_id: s for s in db.query(models.MasdProjectSettings).all()}
    out = []
    for p in db.query(models.ProgramDistrict).order_by(models.ProgramDistrict.name).all():
        ids = projects.member_project_ids(db, p)
        f2f = (
            db.query(func.count(models.FaceToFaceSelection.id))
            .join(models.User, models.User.id == models.FaceToFaceSelection.user_id)
            .filter(models.User.program_district_id.in_(ids))
            .scalar()
        )
        s = saved.get(p.id)
        out.append({
            "slug": p.slug, "name": p.name, "level": getattr(p, "level", None),
            "f2f_learners": f2f,
            "calendar_saved": bool(s and (s.training_date or s.tranche2_start)),
        })
    return out


@router.get("/report")
def masd_report_json(
    project: str = Query(...),
    as_of: Optional[date] = Query(None),
    compare: Optional[date] = Query(None),
    db: Session = Depends(get_db),
):
    p = _project(db, project)
    as_of, compare = _dates(as_of, compare)
    report = masd_report.build(db, p, as_of, compare)
    _log_read(report, p)
    return report


class CumulativeIn(BaseModel):
    as_of: Optional[date] = None
    # None = every F2F learner; otherwise the learners the dashboard's filters keep.
    learner_ids: Optional[List[int]] = Field(None, max_length=20000)
    kind: Optional[Literal["anc", "pnc_lt5", "pnc_ge5", "sn"]] = None
    per_case: bool = True


@router.post("/cumulative")
def masd_cumulative(payload: CumulativeIn, project: str = Query(...), db: Session = Depends(get_db)):
    """Expected vs actual adoptions and activities, day by day, for the
    filtered learners and one kind of adoption (review of 12 Oct 2026).
    Aggregate counts only."""
    from app.masd import engine as masd_engine

    p = _project(db, project)
    as_of, _ = _dates(payload.as_of, None)
    out = masd_engine.cumulative_for(masd_data.load(db, p), as_of, learner_ids=payload.learner_ids,
                                     kind=payload.kind, per_case=payload.per_case)
    phi.log_list(
        resource_type="masd_cumulative",
        count=out["learners"],
        subject_type=phi.CHILD,
        detail={"scope": "aggregate_daily_counts", "project": p.slug, "as_of": as_of.isoformat(),
                "kind": payload.kind, "filtered": payload.learner_ids is not None},
    )
    return out


def _download(db: Session, p: models.ProgramDistrict, as_of: date, compare: Optional[date],
              fmt: str, admin_email: str):
    report = masd_report.build(db, p, as_of, compare)
    if fmt == "pptx":
        from app.masd.pptx_export import build_deck
        buffer, media = build_deck(report), "application/vnd.openxmlformats-officedocument.presentationml.presentation"
    else:
        from app.masd.xlsx_export import build_workbook
        buffer, media = build_workbook(report), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    # Synchronous and committed before the bytes go out: this is the moment the
    # figures leave the platform's custody (SECURITY.md).
    phi.log_export(
        resource_type="masd_report",
        record_count=report["outcomes"]["exclusions"]["total"]["total"],
        fmt=fmt,
        subject_type=phi.CHILD,
        db=db,
        detail={"project": p.slug, "as_of": report["as_of"], "by": admin_email,
                "contents": "aggregate programme report; learner names, no mother or child identifiers",
                "custody_note": "the exported file is in the custody of whoever holds it"},
    )
    db.commit()
    name = f"MASD_{p.name.replace(' ', '_')}_{as_of:%d%m%y}.{fmt}"
    return StreamingResponse(buffer, media_type=media,
                             headers={"Content-Disposition": f'attachment; filename="{name}"'})


@router.get("/report/pptx")
@limiter.limit(lambda: settings.RATE_LIMIT_EXPORT, key_func=principal_key)
def masd_report_pptx(
    request: Request,
    project: str = Query(...),
    as_of: Optional[date] = Query(None),
    compare: Optional[date] = Query(None),
    admin_email: str = Depends(get_admin_email),
    db: Session = Depends(get_db),
):
    """The report as a PowerPoint deck with native (editable) charts."""
    p = _project(db, project)
    as_of, compare = _dates(as_of, compare)
    return _download(db, p, as_of, compare, "pptx", admin_email)


@router.get("/report/xlsx")
@limiter.limit(lambda: settings.RATE_LIMIT_EXPORT, key_func=principal_key)
def masd_report_xlsx(
    request: Request,
    project: str = Query(...),
    as_of: Optional[date] = Query(None),
    compare: Optional[date] = Query(None),
    admin_email: str = Depends(get_admin_email),
    db: Session = Depends(get_db),
):
    """The learner-level MASD sheet plus the block, cadre and outcome tables."""
    p = _project(db, project)
    as_of, compare = _dates(as_of, compare)
    return _download(db, p, as_of, compare, "xlsx", admin_email)


# ── Follow-up lists as Excel (review of 7 Oct 2026) ─────────────────────────

_FLAG_SHEETS = {
    # kind: (report list, sheet title, [(column, key)])
    "anc": ("items", "Pregnancies to follow up", [
        ("Mother ID", "mother_uid"), ("Learner", "learner"), ("Learner ID", "learner_id"), ("Block", "block"),
        ("Adopted", "adopted"), ("LMP", "lmp"), ("Due date (as per LMP)", "edd"),
        ("Days past due date", "days_past_edd"), ("ANC checks done", "anc_done"),
        ("ANC checks due", "anc_expected"), ("Last ANC check", "last_anc"),
        ("Days since last contact", "days_since_contact"), ("Why", "reasons")]),
    "pnc": ("pnc", "PNC cases to follow up", [
        ("Mother ID", "mother_uid"), ("Adoption type", "type"), ("Learner", "learner"), ("Learner ID", "learner_id"),
        ("Batch monitor", "monitor"), ("Block", "block"), ("Adopted", "adopted"), ("Baby's date of birth", "dob"),
        ("Growth checks due", "visits_due"), ("Growth checks done", "visits_done"),
        ("Visits missed till date", "visits_missed"), ("Last visit", "last_visit"),
        ("Days since last visit", "days_since_visit"), ("Why", "reasons")]),
    "mismatch": ("type_mismatch", "Adoption type mismatches", [
        ("Mother ID", "mother_uid"), ("Learner", "learner"), ("Learner ID", "learner_id"),
        ("Batch monitor", "monitor"), ("Block", "block"), ("Adopted", "adopted"),
        ("Type chosen by the learner", "declared"), ("Type from the dates", "computed")]),
}
_TYPE_LABELS = {"anc": "ANC", "pnc_lt5": "PNC <5M", "pnc_ge5": "PNC ≥5M", "unknown": "Not known"}


@router.get("/flags/xlsx")
@limiter.limit(lambda: settings.RATE_LIMIT_EXPORT, key_func=principal_key)
def masd_flags_xlsx(
    request: Request,
    project: str = Query(...),
    kind: Literal["anc", "pnc", "mismatch"] = Query(...),
    as_of: Optional[date] = Query(None),
    types: Optional[str] = Query(None, description="comma-separated adoption types to keep"),
    blocks: Optional[str] = Query(None, description="comma-separated blocks to keep"),
    reasons: Optional[str] = Query(None, description="comma-separated reasons to keep"),
    admin_email: str = Depends(get_admin_email),
    db: Session = Depends(get_db),
):
    """One follow-up list, filtered as on screen. Mothers by record ID only —
    the learner knows who they are."""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    p = _project(db, project)
    as_of, _ = _dates(as_of, None)
    report = masd_report.build(db, p, as_of, None)
    key, title, cols = _FLAG_SHEETS[kind]
    keep_types = set(types.split(",")) if types else None
    keep_blocks = set(blocks.split(",")) if blocks else None
    keep_reasons = set(reasons.split(",")) if reasons else None
    rows = [i for i in report["flags"][key]
            if (keep_types is None or i.get("type", i.get("computed")) in keep_types)
            and (keep_blocks is None or i.get("block") in keep_blocks)
            and (keep_reasons is None or keep_reasons & set(i.get("reasons", [])))]
    labels = {**R.FLAG_REASONS, **R.PNC_FLAG_REASONS}

    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    ws.append([c for c, _ in cols])
    for c in ws[1]:
        c.font = Font(bold=True)
    for i in rows:
        out = []
        for _, k in cols:
            v = i.get(k)
            if k == "reasons":
                v = "; ".join(labels.get(r, r) for r in v or [])
            elif k in ("type", "declared", "computed"):
                v = _TYPE_LABELS.get(v, v)
            out.append(v)
        ws.append(out)
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = max(12, min(40, len(str(col[0].value)) + 4))
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    phi.log_export(
        resource_type="masd_followup_list",
        record_count=len(rows),
        fmt="xlsx",
        subject_type=phi.MOTHER,
        db=db,
        detail={"project": p.slug, "as_of": report["as_of"], "list": kind, "by": admin_email,
                "subject_ids": [i["mother_id"] for i in rows],
                "fields": [c for c, _ in cols],
                "custody_note": "the exported file is in the custody of whoever holds it"},
    )
    db.commit()
    name = f"MASD_{kind}_followup_{p.name.replace(' ', '_')}_{as_of:%d%m%y}.xlsx"
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f'attachment; filename="{name}"'})


# ── Programme calendar and benchmarks ──────────────────────────────────────


class BandValues(BaseModel):
    stunting: Optional[float] = Field(None, ge=0, le=100)
    underweight: Optional[float] = Field(None, ge=0, le=100)
    wasting: Optional[float] = Field(None, ge=0, le=100)
    severe_stunting: Optional[float] = Field(None, ge=0, le=100)
    severe_underweight: Optional[float] = Field(None, ge=0, le=100)
    sam: Optional[float] = Field(None, ge=0, le=100)
    # The survey's sample size for this band, to weight the bands together.
    n: Optional[int] = Field(None, ge=0, le=10_000_000)


class DistrictTrend(BaseModel):
    label: str = Field("", max_length=120)
    rounds: List[str] = Field(default_factory=list, max_length=4)
    stunting: List[Optional[float]] = Field(default_factory=list, max_length=4)
    underweight: List[Optional[float]] = Field(default_factory=list, max_length=4)
    wasting: List[Optional[float]] = Field(default_factory=list, max_length=4)

    @model_validator(mode="after")
    def _same_length(self):
        n = len(self.rounds)
        for name in ("stunting", "underweight", "wasting"):
            values = getattr(self, name)
            if len(values) != n:
                raise ValueError(f"{name} needs one value per survey round")
            if any(v is not None and not 0 <= v <= 100 for v in values):
                raise ValueError(f"{name} values must be percentages (0–100)")
        return self


class Benchmarks(BaseModel):
    label: str = Field("", max_length=120)
    age_bands: Dict[Literal["lt6", "m6_11"], BandValues] = Field(default_factory=dict)
    # District NFHS fact sheets publish under-5 figures only.
    under5: Optional[BandValues] = None
    district_trend: Optional[DistrictTrend] = None


class TargetStep(BaseModel):
    """Expected adoptions per learner from `from` onwards — one tranche. Its
    follow-up counts after `buffer` days (None = the default: a week for the
    first tranche, four days for later ones)."""
    from_: date = Field(..., alias="from")
    anc: int = Field(0, ge=0, le=50)
    pnc_lt5: int = Field(0, ge=0, le=50)
    pnc_ge5: int = Field(0, ge=0, le=50)
    nurse: int = Field(0, ge=0, le=50)
    buffer: Optional[int] = Field(None, ge=0, le=60)

    model_config = {"populate_by_name": True}

    def stored(self) -> dict:
        return {"from": self.from_.isoformat(), "anc": self.anc, "pnc_lt5": self.pnc_lt5,
                "pnc_ge5": self.pnc_ge5, "nurse": self.nurse, "buffer": self.buffer}


class TrainerIn(BaseModel):
    name: str = Field(..., min_length=1, max_length=120)
    phone: Optional[str] = Field(None, max_length=20)
    role: Optional[str] = Field(None, max_length=60)      # "Trainer", "Assistant trainer", ...

    @model_validator(mode="after")
    def _phone(self):
        if self.phone:
            digits = re.sub(r"\D", "", self.phone)
            if not 10 <= len(digits) <= 13:
                raise ValueError("A trainer's phone number should have 10 digits")
            self.phone = digits[-10:]
        self.name = self.name.strip()
        self.role = (self.role or "").strip() or None
        return self


class BatchIn(BaseModel):
    id: Optional[int] = None          # None = a new batch
    name: str = Field(..., min_length=1, max_length=80)
    start_date: Optional[date] = None
    end_date: date
    trainers: List[TrainerIn] = Field(default_factory=list, max_length=6)

    @model_validator(mode="after")
    def _dates(self):
        if self.start_date and self.start_date > self.end_date:
            raise ValueError(f"Batch {self.name}: training cannot start after it ends")
        return self


def _trainers(raw: Optional[str]) -> List[dict]:
    try:
        out = json.loads(raw) if raw else []
    except ValueError:
        return []
    return out if isinstance(out, list) else []


def _batch_out(b: models.MasdTrainingBatch, learners: int) -> dict:
    return {"id": b.id, "name": b.name, "start_date": b.start_date.isoformat() if b.start_date else None,
            "end_date": b.end_date.isoformat(), "trainers": _trainers(b.trainers), "learners": learners}


class SettingsIn(BaseModel):
    training_date: Optional[date] = None
    tranche2_start: Optional[date] = None
    benchmarks: Optional[Benchmarks] = None
    # None = leave as they are (older clients); [] = back to the default.
    targets: Optional[List[TargetStep]] = Field(None, max_length=12)
    batches: Optional[List[BatchIn]] = Field(None, max_length=30)

    @model_validator(mode="after")
    def _order(self):
        if self.training_date and self.tranche2_start and self.tranche2_start <= self.training_date:
            raise ValueError("Tranche 2 must start after the training date")
        if self.targets:
            days = [t.from_ for t in self.targets]
            if len(set(days)) != len(days):
                raise ValueError("Two target steps start on the same date")
        if self.batches:
            names = [b.name.strip().lower() for b in self.batches]
            if len(set(names)) != len(names):
                raise ValueError("Two batches have the same name")
        return self


def _own_batches(db: Session, p: models.ProgramDistrict) -> List[models.MasdTrainingBatch]:
    return (db.query(models.MasdTrainingBatch)
            .filter(models.MasdTrainingBatch.program_district_id == p.id)
            .order_by(models.MasdTrainingBatch.end_date, models.MasdTrainingBatch.name).all())


def _settings_out(db: Session, p: models.ProgramDistrict, s: Optional[models.MasdProjectSettings]) -> dict:
    training = s.training_date if s else None
    tranche2 = s.tranche2_start if s else None
    counts = dict(
        db.query(models.FaceToFaceSelection.batch_id, func.count(models.FaceToFaceSelection.id))
        .filter(models.FaceToFaceSelection.batch_id.isnot(None))
        .group_by(models.FaceToFaceSelection.batch_id).all()
    )
    return {
        "project": p.slug,
        "training_date": training.isoformat() if training else None,
        "tranche2_start": tranche2.isoformat() if tranche2 else None,
        "benchmarks": masd_data.benchmarks_for(p, s),
        "benchmarks_are_default": not (s and s.benchmarks_json) and p.slug in R.DEFAULT_BENCHMARKS,
        "targets": (s.targets_json if s and s.targets_json else R.default_targets(p.slug, training, tranche2)),
        "targets_are_default": not (s and s.targets_json),
        "batches": [_batch_out(b, counts.get(b.id, 0)) for b in _own_batches(db, p)],
        "updated_by": s.updated_by if s else None,
        "updated_at": s.updated_at.isoformat() if s and s.updated_at else None,
    }


@router.get("/settings")
def get_settings(project: str = Query(...), db: Session = Depends(get_db)):
    p = _project(db, project)
    return _settings_out(db, p, masd_data.settings_for(db, p))


@router.put("/settings")
def put_settings(
    payload: SettingsIn,
    project: str = Query(...),
    admin_email: str = Depends(get_admin_email),
    db: Session = Depends(get_db),
):
    p = _project(db, project)
    s = masd_data.settings_for(db, p)
    if s is None:
        s = models.MasdProjectSettings(program_district_id=p.id)
        db.add(s)
    s.training_date = payload.training_date
    s.tranche2_start = payload.tranche2_start
    s.benchmarks_json = payload.benchmarks.model_dump() if payload.benchmarks else None
    if payload.targets is not None:
        s.targets_json = sorted((t.stored() for t in payload.targets), key=lambda t: t["from"]) or None
    if payload.batches is not None:
        existing = {b.id: b for b in _own_batches(db, p)}
        keep = set()
        for b in payload.batches:
            row = existing.get(b.id) if b.id is not None else None
            if row is None:
                row = models.MasdTrainingBatch(program_district_id=p.id)
                db.add(row)
            row.name, row.end_date, row.updated_by = b.name.strip(), b.end_date, admin_email
            row.start_date = b.start_date
            row.trainers = json.dumps([t.model_dump() for t in b.trainers]) if b.trainers else None
            if b.id is not None:
                keep.add(b.id)
        for bid, row in existing.items():
            if bid not in keep:
                db.delete(row)        # its learners fall back to the project's training date
    s.updated_by = admin_email
    audit.record_sync(
        audit.Action.ADMIN_CONFIG_CHANGE,
        db=db,
        resource_type="masd_settings",
        resource_id=p.slug,
        is_phi=False,
        # Trainers' phone numbers stay out of the trail: only that one was set.
        detail={"project": p.slug, "by": admin_email,
                "changes": payload.model_dump(mode="json", by_alias=True,
                                              exclude={"batches": {"__all__": {"trainers"}}}),
                "trainers": [{"batch": b.name, "names": [t.name for t in b.trainers],
                              "phones_set": sum(1 for t in b.trainers if t.phone)}
                             for b in (payload.batches or []) if b.trainers]},
    )
    db.commit()
    db.refresh(s)
    return _settings_out(db, p, s)


# ── The expected-forms table (programme-wide) ──────────────────────────────


class ExpectedFormsIn(BaseModel):
    durations: List[int]
    rows: Dict[str, Dict[str, List[int]]]


@router.get("/expected-forms")
def get_expected_forms(db: Session = Depends(get_db)):
    table, is_default = masd_data.expected_forms_for(db)
    row = db.get(models.MasdRuleTable, masd_data.EXPECTED_FORMS_KEY)
    return {"table": table, "is_default": is_default, "default": R.default_expected_forms(),
            "rows": R.EXPECTED_ROWS,
            "updated_by": row.updated_by if row and not is_default else None,
            "updated_at": row.updated_at.isoformat() if row and row.updated_at and not is_default else None}


@router.put("/expected-forms")
def put_expected_forms(
    payload: ExpectedFormsIn,
    admin_email: str = Depends(get_admin_email),
    db: Session = Depends(get_db),
):
    """Replace the cumulative expected-forms table (e.g. with the analysts'
    revised LAP sheet). Every project's expected activity follows it."""
    table = payload.model_dump()
    problem = R.validate_expected_forms(table)
    if problem:
        raise HTTPException(status_code=422, detail=problem)
    row = db.get(models.MasdRuleTable, masd_data.EXPECTED_FORMS_KEY)
    if row is None:
        row = models.MasdRuleTable(key=masd_data.EXPECTED_FORMS_KEY, value_json=table)
        db.add(row)
    row.value_json, row.updated_by = table, admin_email
    audit.record_sync(
        audit.Action.ADMIN_CONFIG_CHANGE, db=db, resource_type="masd_rule_table",
        resource_id=masd_data.EXPECTED_FORMS_KEY, is_phi=False,
        detail={"table": masd_data.EXPECTED_FORMS_KEY, "by": admin_email},
    )
    db.commit()
    return get_expected_forms(db)


@router.delete("/expected-forms")
def reset_expected_forms(admin_email: str = Depends(get_admin_email), db: Session = Depends(get_db)):
    """Back to the default table."""
    row = db.get(models.MasdRuleTable, masd_data.EXPECTED_FORMS_KEY)
    if row is not None:
        db.delete(row)
        audit.record_sync(
            audit.Action.ADMIN_CONFIG_CHANGE, db=db, resource_type="masd_rule_table",
            resource_id=masd_data.EXPECTED_FORMS_KEY, is_phi=False,
            detail={"table": masd_data.EXPECTED_FORMS_KEY, "reset": True, "by": admin_email},
        )
        db.commit()
    return get_expected_forms(db)


# ── Which batch an F2F learner trained in ──────────────────────────────────


class BatchAssignIn(BaseModel):
    batch_id: Optional[int] = None


@router.put("/learners/{user_id}/batch")
def put_learner_batch(
    user_id: int,
    payload: BatchAssignIn,
    admin_email: str = Depends(get_admin_email),
    db: Session = Depends(get_db),
):
    """Put an F2F learner in a training batch (or none: the project's training
    date then starts their follow-up). The batch must belong to the learner's
    project, or to the state project it sits under."""
    selection = (db.query(models.FaceToFaceSelection)
                 .filter(models.FaceToFaceSelection.user_id == user_id).first())
    if selection is None:
        raise HTTPException(status_code=404, detail="This learner is not selected for face-to-face training")
    batch = None
    if payload.batch_id is not None:
        batch = db.get(models.MasdTrainingBatch, payload.batch_id)
        user = db.get(models.User, user_id)
        allowed = set()
        if user is not None and user.program_district_id is not None:
            allowed.add(user.program_district_id)
            pd = db.get(models.ProgramDistrict, user.program_district_id)
            if pd is not None and getattr(pd, "parent_id", None):
                allowed.add(pd.parent_id)
        if batch is None or batch.program_district_id not in allowed:
            raise HTTPException(status_code=404, detail="Batch not found for this learner's project")
    before = selection.batch_id
    selection.batch_id = payload.batch_id
    audit.record_sync(
        audit.Action.ADMIN_USER_CHANGE, db=db, resource_type="user", resource_id=user_id, is_phi=False,
        detail={"action": "f2f_batch", "from": before, "to": payload.batch_id, "by": admin_email},
    )
    db.commit()
    return {"user_id": user_id, "batch_id": selection.batch_id,
            "batch": batch.name if batch else None,
            "end_date": batch.end_date.isoformat() if batch else None}


@router.get("/batches")
def list_batches(project: str = Query(...), db: Session = Depends(get_db)):
    """A project's training batches (with those of the state it sits under),
    for the batch pickers on the Results and MASD pages."""
    p = _project(db, project)
    ids = [p.id] + ([p.parent_id] if getattr(p, "parent_id", None) else [])
    ids += projects.member_project_ids(db, p)
    rows = (db.query(models.MasdTrainingBatch)
            .filter(models.MasdTrainingBatch.program_district_id.in_(set(ids)))
            .order_by(models.MasdTrainingBatch.end_date, models.MasdTrainingBatch.name).all())
    return [{"id": b.id, "name": b.name, "end_date": b.end_date.isoformat(),
             "start_date": b.start_date.isoformat() if b.start_date else None} for b in rows]


# ── Batches from an Excel: user ID, name, batch (review of 6 Oct 2026) ──────

_BATCH_HEADERS = ("User ID", "Name", "Batch")


def _project_batches(db: Session, p: models.ProgramDistrict) -> List[models.MasdTrainingBatch]:
    ids = {p.id} | set(projects.member_project_ids(db, p))
    if getattr(p, "parent_id", None):
        ids.add(p.parent_id)
    return (db.query(models.MasdTrainingBatch)
            .filter(models.MasdTrainingBatch.program_district_id.in_(ids))
            .order_by(models.MasdTrainingBatch.end_date, models.MasdTrainingBatch.name).all())


def _batch_key(value) -> str:
    """"Batch 3", "batch-3", "B3", "3" and 3.0 all name batch 3; a name
    without a number is matched on its letters."""
    text = str(value).strip().lower()
    if re.fullmatch(r"\d+(\.0+)?", text):
        return str(int(float(text)))
    digits = re.findall(r"\d+", text)
    return str(int(digits[-1])) if digits else re.sub(r"[^a-z]", "", text)


@router.get("/batches/assign-template")
def batch_assign_template(project: str = Query(...), db: Session = Depends(get_db)):
    """An Excel to fill in: one row per learner, with the project's batches on
    a second sheet."""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    p = _project(db, project)
    wb = Workbook()
    ws = wb.active
    ws.title = "Learners"
    ws.append(list(_BATCH_HEADERS))
    for c in ws[1]:
        c.font = Font(bold=True)
    for letter, width in (("A", 12), ("B", 32), ("C", 14)):
        ws.column_dimensions[letter].width = width
    ref = wb.create_sheet("Batches")
    ref.append(["Batch", "Training from", "Training to"])
    for c in ref[1]:
        c.font = Font(bold=True)
    for b in _project_batches(db, p):
        ref.append([b.name, b.start_date.isoformat() if b.start_date else "", b.end_date.isoformat()])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="batch_assignment_{p.slug}.xlsx"'},
    )


@router.post("/batches/assign-upload")
async def batch_assign_upload(
    project: str = Query(...),
    apply: bool = Query(False, description="False = check the file and report; True = assign"),
    file: UploadFile = File(...),
    admin_email: str = Depends(get_admin_email),
    db: Session = Depends(get_db),
):
    """Put many learners in their training batches from an Excel of user ID,
    name and batch number. Checked first (apply=false) and reported row by
    row; nothing changes until it is applied. The user ID decides who; the
    name is only compared, so a typo in it is reported, not fatal."""
    from openpyxl import load_workbook

    p = _project(db, project)
    raw = await file.read()
    if len(raw) > 5 * 1024 * 1024:
        raise HTTPException(status_code=413, detail="The file is larger than 5 MB")
    try:
        ws = load_workbook(io.BytesIO(raw), read_only=True, data_only=True).worksheets[0]
        rows = list(ws.iter_rows(values_only=True))
    except Exception:  # noqa: BLE001 — any unreadable file gets the same answer
        raise HTTPException(status_code=422, detail="Could not read the file — upload an .xlsx")
    if not rows:
        raise HTTPException(status_code=422, detail="The file is empty")
    head = [str(c or "").strip().lower() for c in rows[0]]

    def col(*names: str) -> Optional[int]:
        for name in names:
            for i, h in enumerate(head):
                if h == name or (len(name) > 2 and name in h):
                    return i
        return None

    c_id, c_name, c_batch = col("user id", "userid", "user_id", "id"), col("name"), col("batch")
    if c_id is None or c_batch is None:
        raise HTTPException(status_code=422, detail="The first row must have the columns User ID, Name and Batch")

    by_key: Dict[str, List[models.MasdTrainingBatch]] = {}
    for b in _project_batches(db, p):
        by_key.setdefault(_batch_key(b.name), []).append(b)
    allowed = {p.id} | set(projects.member_project_ids(db, p))

    results, seen = [], set()
    for n, row in enumerate(rows[1:], start=2):
        if row is None or all(v in (None, "") for v in row):
            continue

        def cell(i: Optional[int]):
            return row[i] if i is not None and i < len(row) else None

        uid_raw, name, batch_raw = cell(c_id), cell(c_name), cell(c_batch)
        out = {"row": n, "user_id": None, "name": str(name or "").strip(), "batch": str(batch_raw or "").strip(),
               "status": "error", "message": ""}
        results.append(out)
        try:
            uid = int(float(str(uid_raw).strip()))
        except (TypeError, ValueError):
            out["message"] = "User ID is not a number"
            continue
        out["user_id"] = uid
        if uid in seen:
            out["message"] = "This user ID is on an earlier row too"
            continue
        seen.add(uid)
        user = db.get(models.User, uid)
        if user is None or user.is_admin or user.program_district_id not in allowed:
            out["message"] = "No learner with this user ID in this project"
            continue
        out["learner"] = user.full_name or ""
        match = by_key.get(_batch_key(batch_raw)) if batch_raw not in (None, "") else None
        if not match:
            out["message"] = "No batch with this name or number in this project"
            continue
        if len(match) > 1:
            out["message"] = "More than one batch has this number — use the batch's full name"
            continue
        batch = match[0]
        out["batch"] = batch.name
        sel = db.query(models.FaceToFaceSelection).filter(models.FaceToFaceSelection.user_id == uid).first()
        if sel is None:
            out["message"] = "This learner is not selected for face-to-face training"
            continue
        warn = []
        if name and user.full_name and str(name).strip().lower() != user.full_name.strip().lower():
            warn.append("Name differs from the learner's profile")
        if sel.batch_id == batch.id:
            out["status"], out["message"] = "unchanged", "; ".join(warn) or "Already in this batch"
            continue
        out["status"] = "assign"
        out["from"] = sel.batch.name if sel.batch else None
        out["message"] = "; ".join(warn)
        if apply:
            sel.batch_id = batch.id
            out["status"] = "assigned"

    summary = {k: sum(1 for r in results if r["status"] == k) for k in ("assign", "assigned", "unchanged", "error")}
    if apply and summary["assigned"]:
        audit.record_sync(
            audit.Action.ADMIN_USER_CHANGE, db=db, resource_type="masd_batch_assignment", resource_id=p.slug,
            is_phi=False, detail={"action": "f2f_batch_upload", "assigned": summary["assigned"],
                                  "rows": len(results), "by": admin_email},
        )
        db.commit()
    return {"applied": apply, "summary": summary, "rows": results}


class TrainerRoleIn(BaseModel):
    trainer_role: Optional[Literal["master_trainer", "facilitator"]] = None


@router.put("/learners/{user_id}/trainer-role")
def put_trainer_role(
    user_id: int,
    payload: TrainerRoleIn,
    admin_email: str = Depends(get_admin_email),
    db: Session = Depends(get_db),
):
    """Mark an F2F learner as a Master Trainer or Facilitator (or clear it).
    Only F2F-selected learners can hold the role — it is earned in F2F training."""
    selection = (
        db.query(models.FaceToFaceSelection)
        .filter(models.FaceToFaceSelection.user_id == user_id)
        .first()
    )
    if selection is None:
        raise HTTPException(status_code=404, detail="This learner is not selected for face-to-face training")
    before = selection.trainer_role
    selection.trainer_role = payload.trainer_role
    audit.record_sync(
        audit.Action.ADMIN_USER_CHANGE,
        db=db,
        resource_type="user",
        resource_id=user_id,
        is_phi=False,
        detail={"action": "trainer_role", "from": before, "to": payload.trainer_role, "by": admin_email},
    )
    db.commit()
    return {"user_id": user_id, "trainer_role": selection.trainer_role}


# ── Batch monitors (review of 5 Oct 2026) ────────────────────────────────────
# A batch monitor is a learner the programme trusts to watch over a few others:
# the dashboard tells them "these many of yours are doing well; these two or
# three are yours to chase". Tag a learner here, then give learners a monitor —
# one at a time or a whole list at once.

class BatchMonitorTagIn(BaseModel):
    is_batch_monitor: bool


class MonitorAssignIn(BaseModel):
    batch_monitor_id: Optional[int] = None


class MonitorLearnersIn(BaseModel):
    user_ids: List[int] = Field(default_factory=list, max_length=5000)


def _learner_or_404(db: Session, user_id: int) -> models.User:
    user = db.query(models.User).filter(models.User.id == user_id, models.User.is_admin.is_(False)).first()
    if user is None:
        raise HTTPException(status_code=404, detail="Learner not found")
    return user


def _monitor_change(db: Session, admin_email: str, user_id: int, detail: dict) -> None:
    audit.record_sync(
        audit.Action.ADMIN_USER_CHANGE, db=db, resource_type="user", resource_id=user_id,
        is_phi=False, detail={**detail, "by": admin_email},
    )


@router.put("/learners/{user_id}/batch-monitor")
def put_batch_monitor_tag(
    user_id: int,
    payload: BatchMonitorTagIn,
    admin_email: str = Depends(get_admin_email),
    db: Session = Depends(get_db),
):
    """Tag a learner as a batch monitor, or remove the tag. Removing it frees
    the learners they watched (they show as having no monitor again)."""
    user = _learner_or_404(db, user_id)
    before = bool(user.is_batch_monitor)
    user.is_batch_monitor = payload.is_batch_monitor
    released = 0
    if before and not payload.is_batch_monitor:
        released = (db.query(models.User).filter(models.User.batch_monitor_id == user_id)
                    .update({models.User.batch_monitor_id: None}, synchronize_session="fetch"))
    if payload.is_batch_monitor and user.batch_monitor_id is not None:
        user.batch_monitor_id = None  # a monitor is not watched by another monitor
    _monitor_change(db, admin_email, user_id, {"action": "batch_monitor_tag", "from": before,
                                               "to": payload.is_batch_monitor, "released": released})
    db.commit()
    return {"user_id": user_id, "is_batch_monitor": user.is_batch_monitor, "released": released}


@router.put("/learners/{user_id}/monitor")
def put_learner_monitor(
    user_id: int,
    payload: MonitorAssignIn,
    admin_email: str = Depends(get_admin_email),
    db: Session = Depends(get_db),
):
    """Give a learner a batch monitor (or none)."""
    user = _learner_or_404(db, user_id)
    if payload.batch_monitor_id is not None:
        if payload.batch_monitor_id == user_id:
            raise HTTPException(status_code=400, detail="A learner cannot monitor themselves")
        monitor = _learner_or_404(db, payload.batch_monitor_id)
        if not monitor.is_batch_monitor:
            raise HTTPException(status_code=400, detail="That learner is not tagged as a batch monitor")
        if user.is_batch_monitor:
            raise HTTPException(status_code=400, detail="A batch monitor is not watched by another monitor")
    before = user.batch_monitor_id
    user.batch_monitor_id = payload.batch_monitor_id
    _monitor_change(db, admin_email, user_id, {"action": "batch_monitor_assign", "from": before,
                                               "to": payload.batch_monitor_id})
    db.commit()
    return {"user_id": user_id, "batch_monitor_id": user.batch_monitor_id}


@router.put("/monitors/{monitor_id}/learners")
def put_monitor_learners(
    monitor_id: int,
    payload: MonitorLearnersIn,
    admin_email: str = Depends(get_admin_email),
    db: Session = Depends(get_db),
):
    """Give many learners the same batch monitor in one go (a whole training
    batch, a block). Other monitors and the monitor themself are skipped."""
    monitor = _learner_or_404(db, monitor_id)
    if not monitor.is_batch_monitor:
        raise HTTPException(status_code=400, detail="That learner is not tagged as a batch monitor")
    ids = [i for i in dict.fromkeys(payload.user_ids) if i != monitor_id]
    learners = db.query(models.User).filter(
        models.User.id.in_(ids), models.User.is_admin.is_(False),
        models.User.is_batch_monitor.is_(False)).all() if ids else []
    for u in learners:
        u.batch_monitor_id = monitor_id
    _monitor_change(db, admin_email, monitor_id, {"action": "batch_monitor_assign_many",
                                                  "count": len(learners), "requested": len(ids)})
    db.commit()
    return {"monitor_id": monitor_id, "assigned": len(learners), "skipped": len(ids) - len(learners)}
