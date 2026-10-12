"""
Every question of every form, one row each, with the rules that govern it.

The review sheet asked for on 5 Oct 2026: a single workbook a programme lead
reads question by question — answer type, limits, skip patterns, verdicts —
and flags the variables that change from one project to the next (the
location lists, the cadres) so a new project starts from a ready reckoner.

Both form dialects are walked:
  * flat forms (registration, check growth, antenatal): the ordered field
    list, with `showIf` conditions as the skip pattern;
  * flow forms (breastfeeding, complementary feeding, protein intake): the
    canvas graph in the order a learner meets it — following `next` from the
    start node — with section children, matrices (one row for the grid and
    one per food item), information blocks, per-answer branching, verdicts
    and coaching.

`catalog_rows()` is pure (schema in, rows out) so the admin download, the
script and the tests all share it. Rows are plain dicts keyed by COLUMNS.
"""
from __future__ import annotations

import re
from typing import Any, Dict, Iterable, List, Optional

# (key, heading, width) — the order of the sheet.
COLUMNS: List[tuple] = [
    ("form", "Form", 22),
    ("seq", "No.", 6),
    ("section", "Section / block", 22),
    ("variable", "Variable name", 24),
    ("question", "Question", 48),
    ("item", "Item", 14),
    ("answer_type", "Answer type", 16),
    ("required", "Required", 11),
    ("options", "Answer options", 48),
    ("limits", "Limits (blocks the entry)", 24),
    ("expected", "Expected range (flags, does not block)", 22),
    ("date_rules", "Date rules", 22),
    ("shown_if", "Shown only if (skip pattern)", 36),
    ("next", "Goes to next", 36),
    ("verdicts", "Verdicts", 40),
    ("coaching", "Coaching shown for an answer", 44),
    ("notes", "Help text / notes", 40),
    ("varies", "May change per project", 22),
    ("flag", "Reviewer flag", 14),
    ("comments", "Comments", 30),
]

_FLAT_TYPES = {
    "text": "Text", "textarea": "Long text", "number": "Number", "date": "Date",
    "dropdown": "Choose one (list)", "radio": "Choose one", "checkbox": "Choose any",
    "image": "Photo",
}
_FLOW_TYPES = {
    "single": "Choose one", "multi": "Choose any", "text": "Text", "date": "Date", "number": "Number",
}
_MATRIX_COL_TYPES = {
    "number": "Number", "text": "Text", "date": "Date", "datetime": "Date & time", "dropdown": "Choose one (list)",
}
_COMPUTED = {
    "gestational_age": "Worked out from the mother's LMP at the assessment date (read only)",
    "weight_gain": "This visit's weight minus the previous visit's (read only)",
}
_BUILTIN_VERDICTS = {"green": "As per LAP", "red": "Needs tutorial"}

# Variables a new project is likely to change: its places, facilities and
# cadres. Matched on the variable name and the question wording.
_VARIES = [
    (re.compile(r"\bration\b|ration_card|state-specific", re.I), "State-specific list"),
    (re.compile(r"hwc|phc|sub.?centre|facility|anganwadi centre|\bawc\b|health (and|&) wellness", re.I), "Facility list"),
    (re.compile(r"\b(state|district|taluk|tehsil|block|village|ward|hamlet)\b|(^|_)(state|district|taluk|block|village)(_|$)", re.I),
     "Location list"),
    (re.compile(r"designation|cadre|department", re.I), "Cadre / department list"),
]


def _num(v: Any) -> str:
    if v is None or v == "":
        return ""
    try:
        f = float(v)
    except (TypeError, ValueError):
        return str(v)
    return str(int(f)) if f.is_integer() else f"{f:g}"


def _range(lo: Any, hi: Any) -> str:
    lo, hi = _num(lo), _num(hi)
    if lo and hi:
        return f"{lo} – {hi}"
    if lo:
        return f"at least {lo}"
    if hi:
        return f"at most {hi}"
    return ""


def _decimals(d: Any) -> str:
    if d is None or d == "":
        return ""
    d = int(d)
    return "whole numbers" if d == 0 else f"up to {d} decimal place{'s' if d > 1 else ''}"


def _join(*parts: str, sep: str = "; ") -> str:
    return sep.join(p for p in parts if p)


def _varies(variable: str, question: str, notes: str, options: Optional[list], answer_type: str) -> str:
    # Counts and dates are not lists ("Years at current facility" is a number).
    if answer_type in ("Number", "Date", "Photo"):
        return ""
    # The variable name decides first; the wording only for a list the project
    # fills in itself ("…during the stay at the facility?" with Yes/No is not a
    # facility list); help text only for a state lookup.
    dynamic = options is not None and len(options) == 0
    for pattern, label in _VARIES:
        if pattern.search(variable) or (dynamic and pattern.search(question)):
            return label
    if _VARIES[0][0].search(notes or ""):
        return _VARIES[0][1]
    if answer_type.startswith("Choose") and dynamic:
        return "List filled per project"
    return ""


def _row(form: str, **kw: Any) -> Dict[str, Any]:
    row = {key: "" for key, _, _ in COLUMNS}
    row["form"] = form
    row.update({k: v for k, v in kw.items() if v is not None})
    return row


# ── Flat forms ────────────────────────────────────────────────────────────────

def _flat_rows(form: str, schema: dict) -> List[Dict[str, Any]]:
    fields = [f for f in (schema.get("fields") or []) if isinstance(f, dict)]
    by_id = {f.get("id"): f for f in fields}

    def option_labels(field: dict, values: Iterable[str]) -> str:
        labels = {o.get("value"): o.get("label") for o in (field.get("options") or []) if isinstance(o, dict)}
        return " or ".join(f"‘{labels.get(v, v)}’" for v in values)

    def condition(c: dict) -> str:
        kind = c.get("kind")
        if kind == "ageLtDays":
            return f"the child is younger than {c.get('days')} days"
        if kind == "ageGteDays":
            return f"the child is {c.get('days')} days or older"
        target = by_id.get(c.get("fieldId")) or {}
        name = target.get("label") or c.get("fieldId") or "?"
        return f"‘{name}’ is {option_labels(target, c.get('anyOf') or [])}"

    rows = []
    for i, f in enumerate(fields, start=1):
        ftype = f.get("type") or "text"
        answer_type = _FLAT_TYPES.get(ftype, ftype)
        opts = f.get("options")
        options = ""
        if isinstance(opts, list) and opts:
            options = "; ".join(
                (o.get("label") or o.get("value") or "") + (" (clears the others)" if o.get("exclusive") else "")
                for o in opts if isinstance(o, dict)
            )
        if f.get("allowOther"):
            options = _join(options, "Other (please specify)")
        notes = f.get("helpText") or ""
        computed = f.get("computed")
        read_only = bool(computed) or "read only" in notes.lower() or "read-only" in notes.lower()
        required = "Read only" if read_only else ("Yes" if f.get("required") else "No")
        if f.get("showIf") and required == "Yes":
            required = "Yes, when shown"
        limits = ""
        if ftype == "number":
            limits = _join(_range(f.get("min"), f.get("max")), _decimals(f.get("decimals")))
        date_rules = ""
        if ftype == "date":
            date_rules = _join(
                "not in the future" if f.get("noFuture") else "",
                "not before the child's date of birth" if f.get("notBeforeDob") else "",
            )
        rows.append(_row(
            form,
            seq=i,
            variable=f.get("id") or "",
            question=f.get("label") or "",
            item="Question" if not read_only else "Auto / read only",
            answer_type=answer_type,
            required=required,
            options=options or ("(list kept per project)" if ftype in ("dropdown", "radio", "checkbox") and isinstance(opts, list) and not opts else ""),
            limits=limits,
            expected=_range(f.get("flagMin"), f.get("flagMax")) if ftype == "number" else "",
            date_rules=date_rules,
            shown_if=" and ".join(condition(c) for c in (f.get("showIf") or []) if isinstance(c, dict)),
            notes=_join(notes, _COMPUTED.get(computed or "", "")),
            varies=_varies(f.get("id") or "", f.get("label") or "", notes, opts if isinstance(opts, list) else None, answer_type),
        ))
    return rows


# ── Flow forms ────────────────────────────────────────────────────────────────

def _flow_order(schema: dict) -> List[dict]:
    """Nodes in the order a learner meets them: the default `next` chain from
    the start node, then option branches, then anything unreachable."""
    nodes: Dict[str, dict] = {k: v for k, v in (schema.get("nodes") or {}).items() if isinstance(v, dict)}
    order: List[str] = []
    seen = set()

    def visit(node_id: Optional[str]) -> None:
        while node_id and node_id in nodes and node_id not in seen:
            seen.add(node_id)
            order.append(node_id)
            node = nodes[node_id]
            for opt in node.get("options") or []:
                branch = opt.get("next") if isinstance(opt, dict) else None
                if branch and branch not in seen:
                    pending.append(branch)
            node_id = node.get("next")

    pending: List[str] = []
    start = schema.get("startNodeId")
    visit(start if start in nodes else (next(iter(nodes), None)))
    while pending:
        visit(pending.pop(0))
    for node_id in nodes:
        visit(node_id)
    return [nodes[n] for n in order]


def _flow_rows(form: str, schema: dict) -> List[Dict[str, Any]]:
    nodes: Dict[str, dict] = {k: v for k, v in (schema.get("nodes") or {}).items() if isinstance(v, dict)}
    verdict_names = dict(_BUILTIN_VERDICTS)
    for v in schema.get("verdicts") or []:
        if isinstance(v, dict) and v.get("id"):
            verdict_names[v["id"]] = v.get("label") or v["id"]

    def title_of(node_id: Optional[str]) -> str:
        n = nodes.get(node_id or "")
        return f"‘{n.get('title') or node_id}’" if n else ""

    def option_label(node_id: str, option_id: str) -> str:
        n = nodes.get(node_id) or {}
        for o in n.get("options") or []:
            if isinstance(o, dict) and o.get("id") == option_id:
                return o.get("label") or option_id
        return option_id

    def visible_if(rule: Any) -> str:
        if not isinstance(rule, dict) or not rule.get("nodeId"):
            return ""
        answers = " or ".join(f"‘{option_label(rule['nodeId'], a)}’" for a in rule.get("anyOf") or [])
        return f"{title_of(rule['nodeId'])} is answered {answers}"

    def options_text(opts: list) -> str:
        return "; ".join(
            (o.get("label") or "") + (" (clears the others)" if o.get("exclusive") else "")
            for o in opts if isinstance(o, dict)
        )

    def verdicts_text(opts: list) -> str:
        marked = [(o.get("label") or "", verdict_names.get(o.get("verdict"), o.get("verdict")))
                  for o in opts if isinstance(o, dict) and o.get("verdict")]
        return "; ".join(f"{label} → {v}" for label, v in marked)

    def coaching_text(opts: list) -> str:
        out = []
        for o in opts:
            if not isinstance(o, dict):
                continue
            a = o.get("action") or {}
            kind = a.get("type") or "none"
            if kind == "none":
                continue
            what = a.get("message") or ""
            if kind in ("youtube", "video"):
                clip = _range(a.get("startSeconds"), a.get("endSeconds"))
                what = _join(what, f"video{' (' + clip + ' s)' if clip else ''}", sep=" — ")
            out.append(f"If ‘{o.get('label') or ''}’: {what}")
        return " | ".join(out)

    def branching(node: dict) -> str:
        parts = []
        for o in node.get("options") or []:
            if isinstance(o, dict) and o.get("next"):
                parts.append(f"if ‘{o.get('label') or ''}’ → {title_of(o['next'])}")
        default = title_of(node.get("next"))
        if parts:
            parts.append(f"otherwise → {default or 'end of form'}")
            return "; ".join(parts)
        return default or "end of form"

    def numeric(q: dict) -> tuple:
        n = q.get("numeric") or {}
        return _decimals(n.get("decimals")), _range(n.get("flagMin"), n.get("flagMax"))

    rows: List[Dict[str, Any]] = []
    seq = 0

    def question_row(q: dict, section: str = "", shown: str = "", nxt: str = "") -> None:
        nonlocal seq
        seq += 1
        opts = [o for o in (q.get("options") or []) if isinstance(o, dict)]
        qtype = q.get("questionType") or "single"
        decimals, expected = numeric(q) if qtype == "number" else ("", "")
        rows.append(_row(
            form, seq=seq, section=section, variable=q.get("id") or "", question=q.get("title") or "",
            item="Question", answer_type=_FLOW_TYPES.get(qtype, qtype),
            required="Yes" if q.get("required") else "No",
            options=options_text(opts), limits=decimals, expected=expected,
            shown_if=shown, next=nxt, verdicts=verdicts_text(opts), coaching=coaching_text(opts),
            notes=q.get("helpText") or "",
            varies=_varies(q.get("id") or "", q.get("title") or "", q.get("helpText") or "", None, ""),
        ))

    for node in _flow_order(schema):
        kind = node.get("kind")
        shown = visible_if(node.get("visibleIf"))
        if kind == "question":
            question_row(node, shown=shown, nxt=branching(node))
        elif kind == "section":
            title = node.get("title") or ""
            children = [c for c in node.get("children") or [] if isinstance(c, dict)]
            for i, child in enumerate(children):
                question_row(child, section=title, shown=shown,
                             nxt=branching(node) if i == len(children) - 1 else "")
        elif kind == "info":
            seq += 1
            rows.append(_row(
                form, seq=seq, variable=node.get("id") or "", question=node.get("title") or "",
                item="Information shown", shown_if=shown, next=branching(node), notes=node.get("body") or "",
            ))
        elif kind == "matrix":
            seq += 1
            title = node.get("title") or ""
            cols = [c for c in node.get("columns") or [] if isinstance(c, dict)]
            col_desc = []
            for c in cols:
                ctype = _MATRIX_COL_TYPES.get(c.get("type") or "", c.get("type") or "")
                bits = [ctype]
                if c.get("options"):
                    bits.append(", ".join(str(o.get("label")) for o in c["options"] if isinstance(o, dict)))
                n = c.get("numeric") or {}
                if n:
                    bits.append(_join(_decimals(n.get("decimals")),
                                      ("expected " + _range(n.get("flagMin"), n.get("flagMax"))) if _range(n.get("flagMin"), n.get("flagMax")) else "", sep=", "))
                if c.get("zeroesRow"):
                    bits.append("0 here blanks the rest of the row")
                if c.get("learnerHidden"):
                    bits.append("hidden from the learner")
                col_desc.append(f"{c.get('label')}: " + "; ".join(b for b in bits if b))
            rows.append(_row(
                form, seq=seq, section=title, variable=node.get("id") or "", question=title,
                item="Grid", answer_type="Grid (rows × columns)",
                required="Yes" if node.get("required") else "No",
                options=" | ".join(col_desc), shown_if=shown, next=branching(node),
                notes=_join(node.get("helpText") or "", f"unit: {node['unitLabel']}" if node.get("unitLabel") else ""),
            ))
            for r in node.get("rows") or []:
                if not isinstance(r, dict):
                    continue
                facts = _join(
                    f"unit: {r['unit']}" if r.get("unit") else "",
                    f"protein per serving: {_num(r['proteinPerServing'])} g" if r.get("proteinPerServing") is not None else "",
                    "high-quality protein" if r.get("highQuality") else "",
                    r.get("helpText") or "",
                )
                rows.append(_row(
                    form, seq="", section=title, variable=r.get("id") or "", question=r.get("label") or "",
                    item="Grid row", shown_if=visible_if(r.get("visibleIf")), notes=facts,
                ))
    return rows


def catalog_rows(form_title: str, schema: dict) -> List[Dict[str, Any]]:
    """One row per question (and per grid row / information block) of a form."""
    if not isinstance(schema, dict):
        return []
    if "nodes" in schema:
        return _flow_rows(form_title, schema)
    return _flat_rows(form_title, schema)


# ── Workbook ─────────────────────────────────────────────────────────────────

README = [
    ("What this is", "Every question NurtureHUB asks, form by form, in the order a learner meets it, with the rules that govern each one."),
    ("Why", "Review each question and flag the variables that change from project to project (places, facilities, cadres) — "
            "the flagged list becomes the ready reckoner for setting up a new project."),
    ("Limits (blocks the entry)", "A value outside this range cannot be saved."),
    ("Expected range", "A value outside this range is saved but flagged red for follow-up."),
    ("Shown only if", "The skip pattern: the question appears only when this holds; otherwise it is skipped."),
    ("Goes to next", "For the step-by-step assessments: where the form goes after this question, including jumps for particular answers."),
    ("Verdicts", "As per LAP (green) or Needs tutorial (red): how an answer is scored in the learner's plan."),
    ("May change per project", "Filled in automatically where the question is a place, facility or cadre list — a starting point for your own flags."),
    ("Reviewer flag / Comments", "Empty columns for your review."),
    ("Registration forms", "The learner, mother and child registration pages follow the EP HST tools instrument; the Form Builder keeps the same questions as a version history."),
]


def build_workbook(forms: List[dict], project_name: str, generated_on: str):
    """forms: [{"title", "form_key", "version", "default_version", "pinned", "schema"}] in sheet order."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    head_fill = PatternFill("solid", fgColor="1F2937")
    head_font = Font(bold=True, color="FFFFFF")
    band = PatternFill("solid", fgColor="F3F4F6")
    thin = Side(style="thin", color="D1D5DB")
    wrap = Alignment(wrap_text=True, vertical="top")

    ws = wb.active
    ws.title = "Read me"
    ws.append(["NurtureHUB — every form question", f"{project_name} · {generated_on}"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([])
    for k, v in README:
        ws.append([k, v])
        ws.cell(ws.max_row, 1).font = Font(bold=True)
    ws.append([])
    ws.append(["Form", "Version in this sheet", "Default version", "Questions", "Note"])
    for c in ws[ws.max_row]:
        c.font = head_font
        c.fill = head_fill
    for f in forms:
        rows = catalog_rows(f["title"], f["schema"])
        n = sum(1 for r in rows if r["item"] in ("Question", "Auto / read only", "Grid"))
        ws.append([f["title"], f"v{f['version']}" if f.get("version") else "—",
                   f"v{f['default_version']}" if f.get("default_version") else "—", n,
                   "Pinned for this project" if f.get("pinned") else ""])
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 110
    for row in ws.iter_rows(min_row=3):
        for c in row:
            c.alignment = wrap

    all_rows: List[Dict[str, Any]] = []
    for f in forms:
        all_rows.extend(catalog_rows(f["title"], f["schema"]))

    def sheet(title: str, rows: List[Dict[str, Any]]) -> None:
        s = wb.create_sheet(title[:31])
        s.append([h for _, h, _ in COLUMNS])
        for c in s[1]:
            c.font = head_font
            c.fill = head_fill
            c.alignment = Alignment(wrap_text=True, vertical="center")
        prev_form = None
        shade = False
        for r in rows:
            if r["form"] != prev_form:
                shade = not shade
                prev_form = r["form"]
            s.append([r[k] for k, _, _ in COLUMNS])
            for c in s[s.max_row]:
                c.alignment = wrap
                c.border = Border(bottom=thin)
                if shade:
                    c.fill = band
        for i, (_, _, width) in enumerate(COLUMNS, start=1):
            s.column_dimensions[get_column_letter(i)].width = width
        s.freeze_panes = "F2"
        s.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{max(1, s.max_row)}"

    sheet("All questions", all_rows)
    for f in forms:
        sheet(f["title"].replace("/", "-"), catalog_rows(f["title"], f["schema"]))
    return wb
