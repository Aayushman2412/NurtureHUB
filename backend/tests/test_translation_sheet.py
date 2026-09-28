"""The language-review sheet must keep telling the truth.

The workbook we send for review lists the sentences the SERVER sends — the
verification email and the notifications — by hand, because they live in
Python rather than in the app's translation files. Nothing stops someone
rewording one of them in the code and leaving the sheet behind, so a reviewer
would be approving a sentence nobody sees any more.

These tests pin both halves: every listed sentence still exists in the file it
claims to come from, and every sentence in the app's English translation files
reaches the sheet — with the context (where, what kind of text, which question)
a reviewer needs to translate it well.

  cd backend && ./venv-win/Scripts/python.exe -m pytest tests/test_translation_sheet.py -v
"""
from __future__ import annotations

import json
import re
from pathlib import Path

import pytest

from scripts.export_translation_sheet import (
    AREAS, BACKEND_STRINGS, HINDI_DRAFTS, LOCALES, MARATHI_DRAFTS, _tokens, collect_rows, load_drafts,
)
from scripts.translation_context import form_rows

BACKEND = Path(__file__).resolve().parents[1]
NOT_FROM_LOCALES = ("Messages & email", "Assessment forms")


def _normalise(text: str) -> str:
    """Compare on words, not typography: the sheet writes plain hyphens, plain
    ampersands and straight quotes where the code uses dashes, HTML entities
    and curly quotes."""
    for a, b in (("—", "-"), ("–", "-"), ("·", "-"), ("•", "-"), ("&amp;", "&"),
                 ("’", "'"), ("“", '"'), ("”", '"')):
        text = text.replace(a, b)
    return re.sub(r"\s+", " ", text).strip()


def _source_text(path: Path) -> str:
    """The file's text with Python's adjacent-literal seams closed up, so a
    sentence written as "half " "and half" reads as one sentence here too."""
    return re.sub(r'"\s*(?:f|r|rf|fr)?"', "", _normalise(path.read_text(encoding="utf-8")))


def _fragments(english: str) -> list[str]:
    """The literal pieces between {placeholders} that must appear in the code."""
    return [f for f in (_normalise(p) for p in re.split(r"\{[^}]*\}", english)) if len(f) >= 15]


@pytest.mark.parametrize("entry", BACKEND_STRINGS, ids=[e[0] for e in BACKEND_STRINGS])
def test_server_sentence_is_still_in_the_code(entry):
    sid, where, english, _note = entry
    match = re.search(r"\(([^)]+\.py)\)", where)
    assert match, f"{sid}: 'where' must name the source file, e.g. (app/utils.py)"
    source = _source_text(BACKEND / match.group(1))

    pieces = _fragments(english)
    if not pieces:                       # very short ones: "Hello,", "Passed"
        pieces = [_normalise(english)]
    missing = [p for p in pieces if p not in source]
    assert not missing, (
        f"{sid}: the sheet says the app sends {missing!r}, but that is no longer in "
        f"{match.group(1)}. Update BACKEND_STRINGS in scripts/export_translation_sheet.py."
    )


def test_every_translation_namespace_has_a_home_in_the_sheet():
    files = {p.stem for p in (LOCALES / "en").glob("*.json")}
    missing = sorted(files - set(AREAS))
    assert not missing, (
        f"New translation file(s) {missing}: add them to AREAS in "
        f"scripts/export_translation_sheet.py so reviewers see which screen they belong to."
    )


@pytest.mark.parametrize("path", [MARATHI_DRAFTS, HINDI_DRAFTS], ids=["marathi", "hindi"])
def test_drafts_keep_every_placeholder(path):
    """A draft that dropped {{name}} would reach reviewers looking finished,
    and a reviewer skimming Marathi will not notice a missing placeholder."""
    drafts = load_drafts(path)
    if not drafts:
        pytest.skip(f"no drafts in {path.name}")
    english = {r["id"]: r["en"] for rows in collect_rows().values() for r in rows}
    broken = [k for k, text in drafts.items() if k in english and _tokens(text) != _tokens(english[k])]
    assert not broken, f"drafts whose placeholders do not match the English: {broken[:10]}"
    # Form rows come from the database, so the export checks those; everything
    # else must still exist.
    unknown = [k for k in drafts if k not in english and not k.startswith("form:")]
    assert not unknown, f"drafts for sentences no longer in the app: {unknown[:10]} — remove or re-key them"


def test_sheet_covers_every_english_string():
    def flatten(d, prefix=""):
        out = {}
        for k, v in d.items():
            key = f"{prefix}.{k}" if prefix else k
            out.update(flatten(v, key)) if isinstance(v, dict) else out.update({key: v})
        return out

    expected = sum(
        len(flatten(json.loads(p.read_text(encoding="utf-8"))))
        for p in (LOCALES / "en").glob("*.json")
    )
    rows = collect_rows()
    in_sheet = sum(len(r) for name, r in rows.items() if name not in NOT_FROM_LOCALES)
    assert in_sheet == expected, "every English string belongs in the review sheet"
    assert len(rows["Messages & email"]) == len(BACKEND_STRINGS)


def test_every_row_says_where_it_is_and_what_it_is():
    blank = [r["id"] for rows in collect_rows().values() for r in rows
             if not str(r["where"]).strip() or not str(r["type"]).strip()]
    assert not blank, f"rows without context: {blank[:10]}"


def test_every_answer_choice_names_its_question():
    """'Normal', 'None', 'Sometimes' cannot be translated well without the
    question. A new option group needs a line in translation_context._OPTION_QUESTIONS."""
    orphans = [r["id"] for rows in collect_rows().values() for r in rows
               if r["type"] == "Answer choice" and not r["part_of"]]
    assert not orphans, f"answer choices with no question in 'Part of': {orphans[:10]}"


def test_form_rows_carry_their_question_and_skip_bare_numbers():
    flow = {
        "startNodeId": "sec",
        "nodes": {
            "sec": {"id": "sec", "kind": "section", "title": "Latching", "next": "tbl", "children": [
                {"id": "q1", "kind": "question", "title": "Baby's cheeks", "helpText": "Look closely.",
                 "options": [{"id": "g", "label": "Full and rounded", "action": {"message": "Well done."}},
                             {"id": "r", "label": "Hollow", "action": {"message": "Relatch the baby."}}]},
            ]},
            "tbl": {"id": "tbl", "kind": "matrix", "title": "Grains", "unitLabel": "", "next": "tbl2",
                    "columns": [{"id": "f", "label": "Frequency", "options": [{"label": "0"}, {"label": "1 day"}]}],
                    "rows": [{"id": "roti", "label": "1 roti", "unit": ""},
                             {"id": "rice", "label": "1 cup rice", "unit": "cup"}]},
            "tbl2": {"id": "tbl2", "kind": "matrix", "title": "Eggs",
                     "columns": [{"id": "f", "label": "Frequency", "options": [{"label": "1 day"}]}],
                     "rows": [{"id": "egg", "label": "1 egg"}]},
        },
    }
    rows = {r["id"]: r for r in form_rows("breastfeeding", flow)}
    q = rows["form:breastfeeding:q1.opt.r.label"]
    assert (q["type"], q["part_of"], q["en"]) == ("Answer choice", "Question: Baby's cheeks", "Hollow")
    assert q["where"] == "Breastfeeding assessment form > Latching"
    coach = rows["form:breastfeeding:q1.opt.r.message"]
    assert coach["type"].startswith("Coaching message") and "answer: Hollow" in coach["part_of"]
    assert rows["form:breastfeeding:sec.title"]["type"] == "Section heading"
    assert rows["form:breastfeeding:tbl.row.rice.unit"]["part_of"] == "Table: Grains - row: 1 cup rice"
    assert not any(r["en"] == "0" for r in rows.values()), "a bare number has nothing to translate"
    assert sum(r["en"] == "1 day" for r in rows.values()) == 1, "a choice repeated across a table is listed once"

    flat = {"fields": [{"id": "place", "label": "Where was it taken?", "placeholder": "Select location",
                        "options": [{"value": "home", "label": "Mother's home"}]}]}
    rows = {r["id"]: r for r in form_rows("growth_monitoring", flat)}
    assert rows["form:growth_monitoring:place.opt.home.label"]["part_of"] == "Question: Where was it taken?"
    assert rows["form:growth_monitoring:place.placeholder"]["type"].startswith("Placeholder")


def test_import_reads_columns_by_header(tmp_path):
    """The importer must follow the headers, so a workbook with the context
    columns — or the older layout with the ID first — reads the same."""
    from openpyxl import Workbook

    from scripts.export_translation_sheet import _write_sheet
    from scripts.import_translation_sheet import read_sheet

    row = {"id": "mother:form.title", "where": "Register a mother form", "type": "Heading", "part_of": "",
           "en": "Register a mother", "hi": "माता का पंजीकरण", "fix": "माता पंजीकरण", "mr": "मातेची नोंदणी",
           "tokens": "", "note": ""}
    form = {"id": "form:antenatal:q.title", "where": "Antenatal", "type": "Question", "part_of": "",
            "en": "Current Weight (kg)", "hi_draft": "वर्तमान वज़न (कि.ग्रा.)", "mr": "सध्याचे वजन (कि.ग्रॅ.)",
            "tokens": "", "note": ""}
    wb = Workbook()
    wb.remove(wb.active)
    _write_sheet(wb, "Learner app", [row])
    _write_sheet(wb, "Assessment forms", [form])
    old = wb.create_sheet("Old layout")
    old.append(["ID - do not change", "Where it appears", "English (in use now)", "Hindi (in use now)",
                "Hindi - correction (only if wrong)", "Marathi (please fill in)"])
    old.append([row["id"], "x", row["en"], row["hi"], "", "मातेची नोंदणी करा"])
    path = tmp_path / "reviewed.xlsx"
    wb.save(path)

    got = {(r["sheet"], r["id"]): r for r in read_sheet(path)}
    ui = got[("Learner app", "mother:form.title")]
    assert (ui["en"], ui["fix"], ui["mr"]) == (row["en"], row["fix"], row["mr"])
    f = got[("Assessment forms", "form:antenatal:q.title")]
    assert (f["hi_draft"], f["mr"]) == (form["hi_draft"], form["mr"])
    assert got[("Old layout", "mother:form.title")]["mr"] == "मातेची नोंदणी करा"
