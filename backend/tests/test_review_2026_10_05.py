"""What the review of 5 Oct 2026 asked for, held to it.

  * one sheet of every form question, with its rules — limits, skip patterns,
    branching, verdicts — and the project-specific lists flagged;
  * the registration forms in the Form Builder carry the instrument the
    registration pages ask (they had kept their placeholder questions);
  * learners recognised by user ID, mobile and a short cadre; found by user
    ID or mobile in a search;
  * the raw test data downloadable: learners × questions on the first attempt,
    scored, with a question reference sheet;
  * batch monitors: tagged learners who answer for a few others, with each
    monitor's learners summarised (on track / need a call).

    cd backend && ./venv-win/Scripts/python.exe -m pytest tests/test_review_2026_10_05.py -v
"""
from __future__ import annotations

import asyncio
import base64
import json
import os
import tempfile
from datetime import date, datetime, timedelta, timezone
from io import BytesIO
from pathlib import Path

import pytest

os.environ.setdefault("APP_ENV", "development")
_TEST_KEY = base64.urlsafe_b64encode(b"0123456789abcdef0123456789abcdef").decode().rstrip("=")
os.environ.setdefault("PHI_ENCRYPTION_KEYS", f"v1:{_TEST_KEY}")
os.environ.setdefault("PHI_ENCRYPTION_ACTIVE_KEY", "v1")
os.environ.setdefault("PHI_INDEX_KEY", "test-index-key-not-for-production")
os.environ.setdefault("AUDIT_HMAC_KEY", "test-audit-key-not-for-production-0123456789")

from fastapi import HTTPException  # noqa: E402
from openpyxl import load_workbook  # noqa: E402
from sqlalchemy import create_engine  # noqa: E402
from sqlalchemy.orm import sessionmaker  # noqa: E402

from app.database import Base  # noqa: E402
from app.form_catalog import build_workbook, catalog_rows  # noqa: E402
from app.masd import engine as E  # noqa: E402
from app.models import (  # noqa: E402
    Designation, Department, ProgramDistrict, Question, QuestionOption, Stage, User,
    Test as Paper, TestAnswer as PaperAnswer, TestAttempt as PaperAttempt,
)
from app.routers.admin import (  # noqa: E402
    download_test_answers, get_results_questions, list_users_for_admin, short_cadre,
)
from app.routers.masd import (  # noqa: E402
    BatchMonitorTagIn, MonitorAssignIn, MonitorLearnersIn,
    put_batch_monitor_tag, put_learner_monitor, put_monitor_learners,
)
from app.seed_forms import build_schema  # noqa: E402
import app.models_live  # noqa: F401,E402
import app.models_security  # noqa: F401,E402

T0 = datetime(2026, 9, 1, 10, 0, tzinfo=timezone.utc)


# ── The form questions sheet ─────────────────────────────────────────────────

def _rows(key: str):
    return catalog_rows(key, build_schema(key))


def test_every_mother_registration_question_is_listed_in_order():
    rows = _rows("mother_registration")
    assert len(rows) == 38, "the EP HST tools instrument (38 with 'other trusted source')"
    assert [r["seq"] for r in rows] == list(range(1, 39))
    age = next(r for r in rows if r["variable"] == "mother_age")
    assert age["limits"] == "10 – 50; whole numbers"
    lmp = next(r for r in rows if r["variable"] == "lmp")
    assert "not in the future" in lmp["date_rules"]


def test_skip_patterns_read_as_sentences():
    field = next(r for r in _rows("mother_registration") if r["variable"] == "education_field")
    assert field["shown_if"].startswith("‘Highest educational qualification’ is ‘Diploma’")
    assert field["required"] == "Yes, when shown"
    ebf = next(r for r in _rows("growth_monitoring") if r["variable"] == "exclusive_breastfeeding")
    assert ebf["shown_if"] == "the child is younger than 183 days"


def test_flow_forms_show_branching_verdicts_and_gates():
    cf = {r["variable"]: r for r in _rows("complementary_feeding")}
    assert cf["cfa_pulses_freq"]["next"].startswith("if ‘0 days’ → ‘Number of days per week milk products")
    assert "Non-Vegetarian" in cf["cfa_chicken_freq"]["shown_if"]
    bf = {r["variable"]: r for r in _rows("breastfeeding")}
    assert "→ As per LAP" in bf["bf_hunger"]["verdicts"] and "→ Needs tutorial" in bf["bf_hunger"]["verdicts"]
    assert bf["bf_hunger"]["coaching"].startswith("If ‘The mother feeds the baby when the baby cries.’")
    pca = _rows("mother_protein_intake")
    assert any(r["item"] == "Grid" for r in pca) and any(r["item"] == "Grid row" for r in pca)


def test_only_place_facility_and_cadre_lists_are_flagged_as_changing():
    flagged = {(k, r["variable"]): r["varies"] for k in ("learner_registration", "mother_registration",
                                                          "child_registration", "growth_monitoring")
               for r in _rows(k) if r["varies"]}
    assert flagged[("mother_registration", "district")] == "Location list"
    assert flagged[("mother_registration", "hwc_name")] == "Facility list", "its help text mentions State, it is still a facility"
    assert flagged[("mother_registration", "ration_card")] == "State-specific list"
    assert flagged[("learner_registration", "designation")] == "Cadre / department list"
    assert ("growth_monitoring", "illness_duration") not in flagged, "'duration' is not a ration card"
    assert ("learner_registration", "years_facility") not in flagged, "a number of years is not a list"
    assert ("child_registration", "ebf_during_stay") not in flagged, "a yes/no about a facility stay is not a facility list"


def test_workbook_has_a_sheet_per_form_and_a_read_me():
    forms = [{"title": t, "form_key": k, "version": 1, "default_version": 1, "pinned": False, "schema": build_schema(k)}
             for k, t in (("mother_registration", "Mother Registration Form"), ("breastfeeding", "Breastfeeding Assessment"))]
    wb = build_workbook(forms, "Demo", "05 Oct 2026")
    assert wb.sheetnames == ["Read me", "All questions", "Mother Registration Form", "Breastfeeding Assessment"]
    header = [c.value for c in wb["All questions"][1]]
    assert header[:5] == ["Form", "No.", "Section / block", "Variable name", "Question"]
    assert "Shown only if (skip pattern)" in header and "Reviewer flag" in header


def test_registration_snapshot_is_the_instrument():
    """The migration publishes this frozen copy; it must be what the
    registration pages ask (the seed builders) at the time it was written."""
    snap = json.loads((Path(__file__).parents[1] / "alembic" / "versions" / "data"
                       / "registration_instrument_2026_10_05.json").read_text(encoding="utf-8"))
    for key, size in (("learner_registration", 30), ("mother_registration", 38), ("child_registration", 19)):
        assert len(snap[key]["fields"]) == size
        assert [f["id"] for f in snap[key]["fields"]] == [f["id"] for f in build_schema(key)["fields"]]


# ── Learners recognised and found ────────────────────────────────────────────

@pytest.mark.parametrize("cadre,short", [
    ("Anganwadi Worker (AWW)", "AWW"), ("Staff Nurse", "SN"), ("ASHA", "ASHA"),
    ("CHO (Community Health Officer)", "CHO"), ("Other (Specify)", "Other"), ("Lady Supervisor", "LS"),
    (None, None),
])
def test_short_cadre(cadre, short):
    assert short_cadre(cadre) == short


@pytest.fixture()
def world():
    fd, path = tempfile.mkstemp(suffix=".sqlite")
    os.close(fd)
    engine = create_engine(f"sqlite:///{path}", connect_args={"check_same_thread": False})
    Base.metadata.create_all(engine)
    db = sessionmaker(bind=engine, expire_on_commit=False)()
    pd = ProgramDistrict(name="Demo", slug="demo")
    db.add(pd)
    db.flush()
    health = Department(code="HFW", name="Health & Family Welfare")
    db.add(health)
    db.flush()
    aww = Designation(department_id=health.id, name="Anganwadi Worker (AWW)")
    db.add(aww)
    db.flush()
    stage = Stage(program_district_id=pd.id, title="Test phase", stage_type="test", order_index=1)
    db.add(stage)
    db.flush()
    test = Paper(stage_id=stage.id, title="Screening", test_type="screening", duration_minutes=10,
                 passing_score_pct=50, max_attempts=2, status="ended")
    db.add(test)
    db.flush()
    opts = {}
    for i in (1, 2, 3):
        q = Question(test_id=test.id, text=f"Question {i}", marks=1, order_index=i)
        db.add(q)
        db.flush()
        a = QuestionOption(question_id=q.id, label="A", text="right", is_correct=True)
        b = QuestionOption(question_id=q.id, label="B", text="wrong", is_correct=False)
        db.add_all([a, b])
        db.flush()
        opts[i] = (q, a, b)

    def learner(name, phone):
        u = User(email=f"{name.lower()}@t.mock", full_name=name, is_verified=True, program_district_id=pd.id,
                 role="Anganwadi Worker (AWW)", designation_id=aww.id, phone=phone)
        db.add(u)
        db.flush()
        return u

    def attempt(user, number, picks):
        att = PaperAttempt(test_id=test.id, user_id=user.id, attempt_number=number,
                           started_at=T0 + timedelta(hours=number), submitted_at=T0 + timedelta(hours=number, minutes=5),
                           score=0.0, is_passed=False)
        db.add(att)
        db.flush()
        for i, pick in picks.items():
            q, a, b = opts[i]
            chosen = {"A": a, "B": b}.get(pick)
            db.add(PaperAnswer(attempt_id=att.id, question_id=q.id, selected_option_id=chosen.id if chosen else None,
                               is_correct=bool(chosen and chosen.is_correct)))

    meena = learner("Meena", "9876543210")
    rani = learner("Rani", "9123456789")
    learner("Kamla", "9000000001")
    attempt(meena, 1, {1: "A", 2: "B", 3: None})
    attempt(meena, 2, {1: "A", 2: "A", 3: "A"})
    attempt(rani, 1, {1: "B", 2: "A", 3: "A"})
    db.commit()
    yield db, test, {"meena": meena, "rani": rani}
    db.close()


def test_search_finds_a_learner_by_user_id_or_mobile(world):
    db, _, who = world
    by_id = list_users_for_admin(q=str(who["rani"].id), project_id=None, limit=10, offset=0, db=db, admin_email="a@t")
    assert [u["id"] for u in by_id["users"]] == [who["rani"].id]
    hashed = list_users_for_admin(q=f"#{who['meena'].id}", project_id=None, limit=10, offset=0, db=db, admin_email="a@t")
    assert [u["id"] for u in hashed["users"]] == [who["meena"].id]
    by_mobile = list_users_for_admin(q="9876543210", project_id=None, limit=10, offset=0, db=db, admin_email="a@t")
    assert [u["id"] for u in by_mobile["users"]] == [who["meena"].id], "mobiles are encrypted; matched in Python"


def test_each_learners_first_attempt_answers_come_with_the_questions(world):
    db, test, who = world
    out = get_results_questions(district="demo", db=db, admin_email="a@t")["tests"][0]
    assert out["answers"][str(who["meena"].id)] == "A|B|", "first attempt, in paper order, blank kept"
    assert out["answers"][str(who["rani"].id)] == "B|A|A"
    assert [q["correct"] for q in out["questions"]] == [1, 1, 1]


def _xlsx(resp):
    async def collect():
        return b"".join([c async for c in resp.body_iterator])
    return load_workbook(BytesIO(asyncio.run(collect())), read_only=True)


def test_raw_test_answers_workbook(world):
    db, test, who = world
    wb = _xlsx(download_test_answers(district="demo", test_id=test.id, db=db, admin_email="a@t"))
    assert wb.sheetnames == ["Read me", "Screening – answers", "Screening – scored", "Screening – questions"]
    answers = list(wb["Screening – answers"].iter_rows(values_only=True))
    head, rows = answers[0], {r[0]: r for r in answers[1:]}
    assert head[:3] == ("User ID", "Learner", "Mobile") and head[-3:] == ("Q1", "Q2", "Q3")
    assert rows[who["meena"].id][2] == "9876543210"
    assert rows[who["meena"].id][-3:] == ("A", "B", None), "first attempt; blank stays blank"
    scored = {r[0]: r for r in list(wb["Screening – scored"].iter_rows(values_only=True))[1:]}
    assert scored[who["meena"].id][-5:] == (1, 0, None, 1, 3), "Q1 right, Q2 wrong, Q3 blank; 1 of 3"
    ref = list(wb["Screening – questions"].iter_rows(values_only=True))
    assert ref[1][:10] == ("Q1", "Question 1", None, None, "A. right", 2, 1, 50, 0, "Medium")


# ── Batch monitors ───────────────────────────────────────────────────────────

def test_only_a_tagged_monitor_can_be_given_and_never_to_themself(world):
    db, _, who = world
    meena, rani = who["meena"], who["rani"]
    with pytest.raises(HTTPException) as e:
        put_learner_monitor(rani.id, MonitorAssignIn(batch_monitor_id=meena.id), "a@t", db)
    assert e.value.status_code == 400, "Meena is not a monitor yet"
    put_batch_monitor_tag(meena.id, BatchMonitorTagIn(is_batch_monitor=True), "a@t", db)
    put_learner_monitor(rani.id, MonitorAssignIn(batch_monitor_id=meena.id), "a@t", db)
    assert db.get(User, rani.id).batch_monitor_id == meena.id
    with pytest.raises(HTTPException):
        put_learner_monitor(meena.id, MonitorAssignIn(batch_monitor_id=meena.id), "a@t", db)
    out = put_monitor_learners(meena.id, MonitorLearnersIn(user_ids=[meena.id, rani.id]), "a@t", db)
    assert out["assigned"] == 1, "the monitor themself is skipped"
    released = put_batch_monitor_tag(meena.id, BatchMonitorTagIn(is_batch_monitor=False), "a@t", db)
    assert released["released"] == 1 and db.get(User, rani.id).batch_monitor_id is None


def test_each_monitor_is_told_who_is_on_track_and_who_needs_a_call():
    L = E.LearnerIn
    learners = [
        L(1, "Monitor One", role="ASHA", block="North", f2f=True, is_monitor=True),
        L(2, "Busy Two", role="ASHA", block="North", f2f=True, monitor_id=1),
        L(3, "Idle Three", role="ASHA", block="North", f2f=True, monitor_id=1),
        L(4, "Loose Four", role="ASHA", block="North", f2f=True),
    ]
    M, C, A = E.MotherIn, E.ChildIn, E.ActivityIn
    mothers = [M(1, 2, date(2026, 6, 10)), M(2, 2, date(2026, 6, 11))]
    children = [C(1, 1, date(2026, 6, 5), date(2026, 6, 10), "Female", 3.0, 49.0),
                C(2, 2, date(2026, 6, 6), date(2026, 6, 11), "Male", 3.1, 50.0)]
    acts = [A("growth_monitoring", 2, date(2026, 9, 25), child_id=1), A("breastfeeding", 2, date(2026, 9, 26), child_id=2)]
    data = E.MasdData(project={"id": 1, "slug": "demo", "name": "Demo", "level": "district"},
                      learners=learners, mothers=mothers, children=children, activities=acts, growth=[],
                      training_date=date(2026, 6, 1), tranche2_start=None, benchmarks=None, calendar_saved=True)
    report = E.compute(data, date(2026, 10, 1))
    bm = report["batch_monitors"]
    assert bm["f2f_without_monitor"] == 1, "Loose Four"
    (mon,) = bm["monitors"]
    assert (mon["id"], mon["learners"]) == (1, 2)
    assert [c["name"] for c in mon["calls"]] == ["Idle Three"] or "Idle Three" in [c["name"] for c in mon["calls"]]
    assert mon["on_track"] + mon["need_call"] == 2
    row = next(r for r in report["learners"] if r["id"] == 3)
    assert row["monitor"] == "Monitor One" and row["monitor_id"] == 1 and not row["is_monitor"]
