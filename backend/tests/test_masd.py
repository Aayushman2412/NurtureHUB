"""The MASD dashboard's numbers.

Managers decide who gets a supervisor's call and how a district is doing from
these figures, and the analysts compare them with their own report, so the
rules have to be the report's rules exactly:

  * the cumulative forms table reproduces every figure quoted when the
    method was agreed (29 Sep 2026): 90 days of 3·3·3 = 6 ANC, 6 protein,
    84 growth, 30 BF, 15 CF; a case followed ~30 days in pregnancy and 32
    after birth = 20;
  * the targets come in tranches, each counted at its own follow-up (refined
    1–3 Oct 2026): from the day it opened (or the learner's training ended,
    if later) plus a buffer — a week for the first, four days after — rounded
    down to 15. Jalna on 1 Oct 2026: 105 · 60 · 15 days. A learner's expected
    activity is, per tranche, the adoptions it asked for × the table at that
    tranche's follow-up, summed;
  * how quickly each tranche was taken up: days from it opening to the first
    adoption for it and to the one that completed it;
  * an adoption is ANC if the mother was taken on before the birth, otherwise
    PNC split at 150 days (5 months);
  * an interim report only charges learners for what was due by then, under
    the targets in force then;
  * pregnancies past their due date, or behind on fortnightly checks, are flagged;
  * the outcome analysis keeps F2F learners' cases minus Staff Nurses', and a
    case's exclusion reason is the FIRST failed data check, in the report's order;
  * the downloaded deck opens in PowerPoint — no fractional offsets, which
    PowerPoint rejects outright.

    cd backend && ./venv-win/Scripts/python.exe -m pytest tests/test_masd.py -v
"""
from __future__ import annotations

import base64
import os
import re
import tempfile
import zipfile
from datetime import date, datetime, timedelta, timezone
from io import BytesIO

import pytest

os.environ.setdefault("APP_ENV", "development")
_TEST_KEY = base64.urlsafe_b64encode(b"0123456789abcdef0123456789abcdef").decode().rstrip("=")
os.environ.setdefault("PHI_ENCRYPTION_KEYS", f"v1:{_TEST_KEY}")
os.environ.setdefault("PHI_ENCRYPTION_ACTIVE_KEY", "v1")
os.environ.setdefault("PHI_INDEX_KEY", "test-index-key-not-for-production")
os.environ.setdefault("AUDIT_HMAC_KEY", "test-audit-key-not-for-production-0123456789")

from app.masd import engine as E  # noqa: E402
from app.masd import insights as I  # noqa: E402
from app.masd import rules as R  # noqa: E402

AS_OF = date(2026, 9, 29)
TRAINING = date(2026, 6, 1)
TRANCHE2 = date(2026, 7, 13)


# ── Rules ──────────────────────────────────────────────────────────────────


def test_expected_table_reproduces_the_agreed_figures():
    t = R.default_expected_forms()
    assert R.validate_expected_forms(t) is None
    gm = {d: R.expected_at(t, R.PNC_LT5, d)["gm"] for d in (15, 30, 90, 105)}
    assert gm == {15: 9, 30: 11, 90: 14, 105: 15}
    assert R.expected_at(t, R.PNC_LT5, 15)["bf"] == 5, "all five BF assessments in the first 15 days"
    assert (R.expected_at(t, R.PNC_GE5, 75)["cf"], R.expected_at(t, R.PNC_GE5, 90)["cf"]) == (4, 5)
    total = {}
    for atype in R.ADOPTION_TYPES:
        for k, v in R.expected_at(t, atype, 90).items():
            total[k] = total.get(k, 0) + 3 * v
    assert total == {"anc": 6, "protein": 6, "gm": 84, "bf": 30, "cf": 15}
    assert R.expected_at(t, R.NURSE, 90) == {"gm": 6, "bf": 5}, "a Staff Nurse's hospital visits only"
    assert R.expected_at(t, R.PNC_LT5, 0) == {"gm": 0, "bf": 0}, "nothing is due before the first step"


def test_each_tranche_has_its_own_follow_up():
    jalna = R.default_targets("jalna", None, None)
    assert [(s["from"], s["buffer"]) for s in jalna] == [("2026-06-11", 7), ("2026-07-27", 4), ("2026-09-07", 4)]
    trs = R.tranches(jalna, date(2026, 10, 1))
    assert [(t["fu_raw"], t["fu_days"]) for t in trs] == [(105, 105), (62, 60), (20, 15)]
    assert all(t["added"] == {"anc": 1, "pnc_lt5": 1, "pnc_ge5": 1, "nurse": 3} for t in trs)
    assert trs[-1]["targets"] == {"anc": 3, "pnc_lt5": 3, "pnc_ge5": 3, "nurse": 9}
    # Inside the buffer: the tranche is open (it counts towards the target) but nothing is due yet.
    assert [t["fu_days"] for t in R.tranches(jalna, date(2026, 9, 10))] == [75, 30, 0]
    assert len(R.tranches(jalna, date(2026, 7, 26))) == 1, "a tranche not yet opened is left out"
    # Trained after tranche 2 opened: both start on the training day, and a
    # later tranche's clock never starts before an earlier one's.
    late = R.tranches(jalna, date(2026, 10, 1), training_end=date(2026, 8, 1))
    assert [(t["opened"], t["clock_from"]) for t in late[:2]] == [("2026-08-01", "2026-08-08"),
                                                                  ("2026-08-01", "2026-08-08")]
    assert R.round_follow_up(1000) == R.FU_MAX_DAYS


def test_expected_activity_adds_up_tranche_by_tranche():
    """Jalna on 1 Oct 2026, one community learner (1·1·1 asked in each
    tranche): 49 + 44 + 33 = 126 forms, where reading all nine adoptions at
    105 days would ask for 147."""
    t = R.default_expected_forms()
    ex = E.expected_per_learner(t, R.tranches(R.default_targets("jalna", None, None), date(2026, 10, 1)))
    assert [tr["community"]["total"] for tr in ex["tranches"]] == [49, 44, 33]
    assert ex["community"]["total"] == 126
    assert ex["community"]["forms"] == {"anc": 6, "protein": 6, "gm": 15 * 2 + 13 * 2 + 9 * 2, "bf": 30, "cf": 5 + 4 + 1}
    assert ex["community"]["by_type"][R.PNC_LT5]["adoptions"] == 3
    assert ex["nurse"]["total"] == 9 * 11, "a Staff Nurse's flat hospital row does not change with follow-up"


def test_case_expected_follows_the_case():
    t = R.default_expected_forms()
    # Adopted in pregnancy, born 28 days later, baby followed for 32 days.
    e = R.case_expected(t, mother_adopted=date(2026, 6, 1), lmp=date(2025, 10, 1), dob=date(2026, 6, 29),
                        baby_adopted=date(2026, 6, 29), end=date(2026, 7, 31))
    assert (e["anc"], e["protein"], e["gm"], e["bf"], e["cf"], e["total"]) == (2, 2, 11, 5, 0, 20)
    # Adopted at 3 months and followed to 195 days: 105 days before 195.
    dob = date(2025, 10, 3)
    e = R.case_expected(t, mother_adopted=dob + timedelta(days=90), lmp=None, dob=dob,
                        baby_adopted=dob + timedelta(days=90), end=dob + timedelta(days=195))
    assert (e["fu_to_195"], e["gm"], e["bf"]) == (105, 15, 5)
    # Adopted at 198 days: no breastfeeding forms, complementary feeding from the first visit.
    dob = date(2025, 10, 1)
    e = R.case_expected(t, mother_adopted=dob + timedelta(days=198), lmp=None, dob=dob,
                        baby_adopted=dob + timedelta(days=198), end=dob + timedelta(days=231))
    assert e["bf"] == 0 and e["cf"] == 1 + 2, "day one, then 210 and 225"
    # Not born yet: fortnightly checks stop at the due date.
    e = R.case_expected(t, mother_adopted=date(2026, 6, 1), lmp=date(2025, 10, 1), dob=None, baby_adopted=None,
                        end=date(2026, 9, 29))
    due = R.edd(date(2025, 10, 1))
    assert e["anc"] == (due - date(2026, 6, 1)).days // 15 + 1 and e["gm"] == 0


def test_targets_rise_in_steps():
    steps = R.default_targets("jalna", date(2026, 6, 7), date(2026, 7, 19))
    assert R.targets_in_force(steps, date(2026, 6, 10)) == ({"anc": 0, "pnc_lt5": 0, "pnc_ge5": 0, "nurse": 0}, None)
    assert R.targets_in_force(steps, date(2026, 6, 11))[0] == {"anc": 1, "pnc_lt5": 1, "pnc_ge5": 1, "nurse": 3}
    assert R.targets_in_force(steps, date(2026, 9, 29)) == ({"anc": 3, "pnc_lt5": 3, "pnc_ge5": 3, "nurse": 9}, 3)
    other = R.default_targets("ujjain", date(2026, 6, 7), date(2026, 7, 19))
    assert [(s["from"], s["buffer"]) for s in other] == [("2026-06-07", 7), ("2026-07-19", 4), ("2026-09-02", 4)]
    assert R.learner_targets({"anc": 3, "pnc_lt5": 5, "pnc_ge5": 3, "nurse": 9}, True) == \
        {R.ANC: 0, R.PNC_LT5: 9, R.PNC_GE5: 0}


def test_edited_table_is_validated():
    t = R.default_expected_forms()
    t["rows"][R.PNC_LT5]["gm"][5] = 1
    assert "cannot fall" in R.validate_expected_forms(t)
    t = R.default_expected_forms()
    t["durations"] = t["durations"][:-1]
    assert R.validate_expected_forms(t)


@pytest.mark.parametrize("actual,expected,band", [
    (0, 50, "none"), (1, 100, "b1_20"), (20, 100, "b1_20"), (21, 100, "b21_40"), (80, 100, "b61_80"),
    (81, 100, "b81_100"), (150, 100, "b81_100"), (5, 0, None),
])
def test_own_bands(actual, expected, band):
    assert R.own_band(actual, expected) == band


@pytest.mark.parametrize("role,group", [
    ("Anganwadi Worker (AWW)", "AWW"), ("Lady Supervisor", "AWSup"), ("ASHA", "ASHA"),
    ("ASHA Facilitator", "ASHA Sup"), ("ANM", "ANM"), ("CHO (Community Health Officer)", "CHO"),
    ("Staff Nurse", "Nursing Staff"), ("Pharmacist", "Other"), (None, "Other"),
])
def test_role_groups(role, group):
    assert R.role_group(role) == group


def test_delivery_groups():
    assert R.delivery_place_group("District Hospital (DH)") == "Government"
    assert R.delivery_place_group("Private Hospital/Nursing Home") == "Private"
    assert R.delivery_place_group("Home") == "Home"
    assert R.delivery_method_group("Caesarean section") == "Caesarean section"


# ── Classification ─────────────────────────────────────────────────────────


def _m(mid=1, adopted=date(2026, 6, 10), lmp=None):
    return E.MotherIn(id=mid, learner_id=1, adoption_date=adopted, lmp=lmp)


def _c(cid=1, mid=1, dob=None, adopted=None, **kw):
    return E.ChildIn(id=cid, mother_id=mid, dob=dob, adoption_date=adopted, **kw)


@pytest.mark.parametrize("age,expected", [(0, R.PNC_LT5), (149, R.PNC_LT5), (150, R.PNC_GE5), (300, R.PNC_GE5)])
def test_pnc_split_at_150_days(age, expected):
    adopted = date(2026, 6, 10)
    child = _c(dob=date.fromordinal(adopted.toordinal() - age), adopted=adopted)
    assert E.classify(_m(adopted=adopted), [child])[0] == expected


def test_anc_when_mother_adopted_before_birth():
    assert E.classify(_m(lmp=date(2026, 3, 1)), [])[0] == R.ANC
    born_later = _c(dob=date(2026, 8, 1), adopted=date(2026, 8, 2))
    assert E.classify(_m(), [born_later])[0] == R.ANC
    assert E.classify(_m(lmp=None), [])[0] == R.UNKNOWN


# ── A small project ────────────────────────────────────────────────────────


def _project() -> E.MasdData:
    """One MT learner with six cases that each hit a different outcome path,
    a Staff Nurse, a learner who did nothing, and a learner outside F2F."""
    L = E.LearnerIn
    learners = [
        L(1, "Asha One", role="ASHA", block="North", f2f=True, trainer_role="master_trainer"),
        L(2, "Nurse Two", role="Staff Nurse", block="South", f2f=True),
        L(3, "Idle Three", role="Anganwadi Worker (AWW)", block="North", f2f=True),
        L(4, "Outside Four", role="ASHA", block="North", f2f=False),
    ]
    M, C, G, A = E.MotherIn, E.ChildIn, E.GrowthIn, E.ActivityIn
    d = date
    mothers = [
        M(1, 1, d(2026, 6, 10)),                         # included PNC<5M
        M(2, 1, d(2026, 6, 15), lmp=d(2026, 3, 1), uid="UID-M2"),  # ANC, not born yet → no_child
        M(3, 1, d(2026, 7, 20)),                         # PNC≥5M, tranche 2 → only_adoption_visit
        M(4, 1, d(2026, 6, 20)),                         # baby adopted before the mother
        M(5, 1, d(2026, 6, 12)),                         # no sex → z_missing
        M(6, 1, d(2026, 6, 14)),                         # 0.85 kg girl → birth_weight_extreme
        M(7, 2, d(2026, 6, 5)),                          # nurse's case
        M(8, 4, d(2026, 6, 12)),                         # outside F2F
    ]
    children = [
        C(1, 1, d(2026, 6, 5), d(2026, 6, 10), "Female", 3.0, 49.0),
        C(3, 3, d(2026, 1, 1), d(2026, 7, 20), "Male", 3.2, 50.0),
        C(4, 4, d(2026, 6, 1), d(2026, 6, 15), "Male", 3.1, 49.5),
        C(5, 5, d(2026, 6, 1), d(2026, 6, 12), None, 3.0, 49.0),
        C(6, 6, d(2026, 6, 1), d(2026, 6, 14), "Female", 0.85, 45.0),
        C(7, 7, d(2026, 6, 5), d(2026, 6, 5), "Male", 3.3, 50.0),
        C(8, 8, d(2026, 6, 1), d(2026, 6, 12), "Female", 3.0, 49.0),
    ]
    growth = [
        G(1, d(2026, 6, 10), 3.1, 49.5), G(1, d(2026, 7, 10), 4.2, 54.0), G(1, d(2026, 9, 1), 5.6, 60.0),
        G(3, d(2026, 7, 20), 7.4, 66.0),
        G(4, d(2026, 6, 15), 3.6, 51.0), G(4, d(2026, 8, 1), 5.5, 58.0),
        G(5, d(2026, 6, 12), 3.4, 50.0), G(5, d(2026, 8, 1), 5.2, 57.0),
        G(6, d(2026, 6, 14), 1.1, 46.0), G(6, d(2026, 8, 1), 2.5, 52.0),
        G(7, d(2026, 6, 5), 3.3, 50.0), G(7, d(2026, 6, 6), 3.2, 50.0),
        G(8, d(2026, 6, 12), 3.5, 51.0), G(8, d(2026, 8, 12), 5.9, 58.0),
    ]
    acts = [A("growth_monitoring", 1, g.on, child_id=g.child_id) for g in growth if g.child_id in (1, 3, 4, 5, 6)]
    acts += [A("breastfeeding", 1, d(2026, 6, 10), child_id=1), A("breastfeeding", 1, d(2026, 7, 10), child_id=1)]
    acts += [A("antenatal", 1, d(2026, 6, 15) + timedelta(days=14 * k), mother_id=2) for k in range(3)]
    acts += [A("mother_protein_intake", 1, d(2026, 6, 16), mother_id=2)]
    acts += [A("growth_monitoring", 2, d(2026, 6, 5), child_id=7)] * 2 + [A("growth_monitoring", 2, d(2026, 6, 6), child_id=7)]
    acts += [A("breastfeeding", 2, d(2026, 6, 5), child_id=7)] * 2
    acts += [A("growth_monitoring", 4, g.on, child_id=8) for g in growth if g.child_id == 8]
    return E.MasdData(
        project={"id": 1, "slug": "demo", "name": "Demo", "level": "district"},
        learners=learners, mothers=mothers, children=children, activities=acts, growth=growth,
        training_date=TRAINING, tranche2_start=TRANCHE2,
        benchmarks=R.DEFAULT_BENCHMARKS["jalna"], calendar_saved=True,
    )


@pytest.fixture()
def report():
    return E.compute(_project(), AS_OF)


def _learner(r, lid):
    return next(x for x in r["learners"] if x["id"] == lid)


def _expected(steps, as_of, training_end, is_nurse=False):
    """Forms expected of one learner: per tranche, what it added × the table
    at its own follow-up."""
    t = R.default_expected_forms()
    total = 0
    for tr in R.tranches(steps, as_of, training_end):
        lt = R.learner_targets(tr["added"], is_nurse)
        total += sum(lt[a] * v for a in R.ADOPTION_TYPES
                     for v in R.expected_at(t, R.row_type(a, is_nurse), tr["fu_days"]).values())
    return total


def test_targets_and_expected_activity(report):
    s = report["summary"]
    assert s["learners"] == 3, "only F2F learners count"
    cal = report["calendar"]
    assert (cal["fu_raw"], cal["fu_days"]) == (113, 105), "1 Jun + a week's buffer → 29 Sep"
    assert [(t["clock_from"], t["fu_days"]) for t in cal["tranches"]] == [("2026-06-08", 105), ("2026-07-17", 60)]
    now = cal["targets"]["now"]
    assert now == {"anc": 2, "pnc_lt5": 2, "pnc_ge5": 2, "nurse": 6} and cal["targets"]["is_default"]
    one = _learner(report, 1)
    assert one["adoptions"] == {"anc": 1, "pnc_lt5": 4, "pnc_ge5": 1, "unknown": 0, "total": 6}
    assert one["target"] == 6 and one["fulfilment_pct"] == 100.0
    assert one["tranche_fu"] == [105, 60]
    assert one["ideal"]["total"] == 49 + 44 == _expected(cal["targets"]["steps"], AS_OF, TRAINING)
    assert one["activities"]["total"] == 16   # 10 growth checks + 2 BF + 3 ANC + 1 protein
    assert one["by_type"][R.ANC] == {"target": 2, "adopted": 1, "adoption_pct": 50.0, "expected": 8,
                                     "actual": 4, "activity_pct": 50.0}
    nurse = _learner(report, 2)
    assert nurse["target"] == 6 and nurse["ideal"]["total"] == 6 * 11 and nurse["activities"]["total"] == 5
    assert nurse["subtype_pct"]["anc"] is None, "a nurse has no antenatal expectation — not 0%"
    assert _learner(report, 3)["nil_days"] is None
    assert _learner(report, 3)["own"]["band"] is None, "no cases, nothing to band"
    assert one["own"]["band"] is not None and one["own"]["expected"]["total"] > 0


def test_batch_end_starts_the_follow_up():
    data = _project()
    data.batches = [{"id": 7, "name": "Batch 2", "end_date": date(2026, 7, 1)}]
    data.learners[0].batch_id, data.learners[0].training_end = 7, date(2026, 7, 1)
    r = E.compute(data, AS_OF)
    one = _learner(r, 1)
    assert (one["batch"], one["fu_raw"], one["fu_days"]) == ("Batch 2", 83, 75), "1 Jul + a week → 29 Sep"
    assert one["tranche_fu"] == [75, 60], "tranche 2 opened after the batch ended: its own clock"
    assert [b["name"] for b in r["calendar"]["batches"]] == ["Batch 2", None], "the rest train on the project date"
    b2 = r["calendar"]["batches"][0]
    assert b2["tranche_fu"] == [75, 60]
    assert b2["expected"]["community"]["total"] == _expected(r["calendar"]["targets"]["steps"], AS_OF, date(2026, 7, 1))


def test_pregnancy_flags():
    r = E.compute(_project(), AS_OF)
    f = r["flags"]
    assert f["summary"]["open_pregnancies"] == 1
    item = f["items"][0]
    assert item["reasons"] == ["anc_behind"] and item["anc_done"] == 3 and item["anc_expected"] == 8
    data = _project()
    data.mothers[1].lmp = date(2025, 11, 1)          # due 8 Aug 2026, no birth entered
    f = E.compute(data, AS_OF)["flags"]
    assert "edd_passed" in f["items"][0]["reasons"] and f["items"][0]["days_past_edd"] == 52


def test_attention_lists_the_idle_learner(report):
    idle = next(a for a in report["attention"] if a["id"] == 3)
    assert "no_activity" in idle["reasons"] and "few_adoptions" in idle["reasons"]
    assert report["attention"][0]["id"] == 3, "no activity at all sorts first"


def test_funnel_and_first_failed_check(report):
    o = report["outcomes"]
    steps = {s["key"]: s for s in o["funnel"]}
    assert (steps["start"]["learners_after"], steps["start"]["cases_after"]) == (3, 8)
    assert (steps["no_f2f"]["learners_removed"], steps["no_f2f"]["cases_removed"]) == (1, 1)
    assert (steps["nursing_staff"]["learners_after"], steps["nursing_staff"]["cases_after"]) == (1, 6)
    reasons = {x["key"]: x["total"] for x in o["exclusions"]["reasons"]}
    for key in ("no_child", "only_adoption_visit", "baby_before_mother", "z_missing", "birth_weight_extreme"):
        assert reasons[key] == 1, key
    assert o["exclusions"]["included"]["total"] == 1
    assert o["exclusions"]["inclusion_pct_mtfl"] == pytest.approx(16.7)
    assert o["prevalence"]["overall"]["n"] == 1
    fixes = o["data_fixes"]
    assert fixes and fixes[0]["cases"] == 3, "the three data-entry slips, not the follow-up gaps"


def test_interim_counts_only_what_was_due_then():
    interim = E.compute(_project(), date(2026, 7, 1))
    cal = interim["calendar"]
    assert cal["tranches_in_force"] == 1 and cal["targets"]["step"] == 1
    one = _learner(interim, 1)
    assert one["target"] == 3 and one["fu_days"] == 15
    assert one["ideal"]["total"] == 33, "15 days of follow-up under 1·1·1: 4 + 14 + 15"
    assert one["adoptions"]["total"] == 5, "the adoption on 20 July is not counted yet"


def test_uptake_of_each_tranche(report):
    """Learner 1 started tranche 1 nine days after it opened and completed it
    with the 20 Jul PNC≥5M adoption; their second PNC<5M came before tranche 2
    opened, but the second ANC never came. The Staff Nurse started tranche 1
    in four days; the idle learner never did."""
    up = {u["step"]: u for u in report["uptake"]}
    t1, t2 = up[1], up[2]
    assert (t1["learners"], t1["started"], t1["completed"]) == (3, 2, 1)
    assert t1["started_by"]["d7"] == 1 and t1["started_by"]["d15"] == 1 and t1["started_by"]["not_yet"] == 1
    assert (t1["median_days_to_start"], t1["median_days_to_complete"]) == (6.5, 49.0)
    assert (t2["started"], t2["completed"], t2["started_by"]["before"]) == (1, 0, 1)
    assert _learner(report, 1)["uptake"] == [{"step": 1, "started_days": 9, "completed_days": 49},
                                             {"step": 2, "started_days": -31, "completed_days": None}]
    assert report["summary"]["tranche_mix"]["t2_pnc_ge5"] == 1, "adopted 20 Jul, after tranche 2 opened"
    texts = [f["text"] for f in I.uptake(report)]
    assert any(t.startswith("Tranche 1 (opened 01 Jun 2026") for t in texts)
    assert any("already adopted for tranche 2 before it opened" in t for t in texts)


def test_comparison_and_findings(report):
    then = E.compute(_project(), date(2026, 7, 12))
    cmp = E.comparison(then, report)
    assert cmp["then"]["target_per_learner"] == 3 and cmp["now"]["target_per_learner"] == 6
    assert {b["key"] for b in cmp["blocks"]} == {"North", "South"}
    report["comparison"] = cmp
    findings = I.all_findings(report, cmp)
    assert all(f["text"] and f["tone"] in ("good", "watch", "info") for fs in findings.values() for f in fs)
    assert any("12 Jul 2026" in f["text"] for f in findings["progress"]), "dates read as dates, not ISO"


def test_deck_and_workbook_build(report):
    from pptx import Presentation

    from app.masd.pptx_export import build_deck
    from app.masd.xlsx_export import build_workbook

    report["comparison"] = E.comparison(E.compute(_project(), date(2026, 7, 12)), report)
    report["insights"] = I.all_findings(report, report["comparison"])
    deck = build_deck(report).read()
    assert len(Presentation(BytesIO(deck)).slides) >= 30
    with zipfile.ZipFile(BytesIO(deck)) as z:
        for name in z.namelist():
            if name.startswith("ppt/slides/slide") and name.endswith(".xml"):
                xml = z.read(name).decode("utf-8")
                assert not re.search(r'<a:(off|ext) [^>]*="\d+\.\d', xml), f"fractional EMU in {name}"
    from openpyxl import load_workbook
    wb = load_workbook(build_workbook(report))
    assert {"MASD learners", "By block", "By cadre", "By adoption type", "Own-case bands", "Batches", "Tranches",
            "Expected forms", "Pregnancy flags", "Malnutrition", "Needs attention"} <= set(wb.sheetnames)
    tranches = [[c.value for c in row] for row in wb["Tranches"].iter_rows(min_row=4)]
    assert [r[0] for r in tranches] == ["Tranche 1", "Tranche 2"] and tranches[0][12] == 49
    flat = " ".join(str(c.value) for ws in wb.worksheets for row in ws.iter_rows() for c in row if c.value)
    assert "UID-M2" not in flat, "the workbook never carries a mother's record ID"


# ── Routes ─────────────────────────────────────────────────────────────────


@pytest.fixture()
def world():
    from sqlalchemy import create_engine
    from sqlalchemy.orm import sessionmaker

    import app.models_live  # noqa: F401
    import app.models_security  # noqa: F401
    from app import models
    from app.database import Base

    fd, path = tempfile.mkstemp(suffix=".sqlite")
    os.close(fd)
    engine = create_engine(f"sqlite:///{path}", connect_args={"check_same_thread": False})
    Base.metadata.create_all(engine)
    db = sessionmaker(bind=engine, expire_on_commit=False)()
    pd = models.ProgramDistrict(name="Demo", slug="demo")
    db.add(pd)
    db.flush()
    trained = models.User(email="t@t.mock", full_name="Trained", is_verified=True, program_district_id=pd.id, role="ASHA")
    other = models.User(email="o@t.mock", full_name="Other", is_verified=True, program_district_id=pd.id, role="ASHA")
    db.add_all([trained, other])
    db.flush()
    db.add(models.FaceToFaceSelection(user_id=trained.id, program_district_id=pd.id))
    mother = models.Mother(mother_uid="T-M1", registered_by_user_id=trained.id, mother_name="M",
                           adoption_date=date(2026, 6, 10))
    db.add(mother)
    db.flush()
    child = models.Child(child_uid="T-C1", mother_id=mother.id, child_name="C", dob=date(2026, 6, 5),
                         adoption_date=date(2026, 6, 10), gender="Female", birth_weight=3.0, birth_length=49.0)
    db.add(child)
    db.flush()
    for on, w, ln in ((date(2026, 6, 10), "3.1", "49.5"), (date(2026, 8, 10), "5.0", "57.0")):
        db.add(models.FormResponse(
            form_key="growth_monitoring", child_id=child.id, submitted_by_user_id=trained.id,
            assessment_date=on, status="submitted",
            answers_json=[{"nodeId": "baby_weight", "question": "Baby's weight (kg)", "questionType": "number", "value": w},
                          {"nodeId": "baby_length", "question": "Baby's length (cm)", "questionType": "number", "value": ln}],
            created_at=datetime(2026, 8, 10, tzinfo=timezone.utc),
        ))
    db.commit()
    yield db, trained, other
    db.close()


def test_report_route_reads_the_database(world):
    from app.routers.masd import masd_projects, masd_report_json

    db, trained, _ = world
    projects = {p["slug"]: p for p in masd_projects(db=db)}
    assert projects["demo"]["f2f_learners"] == 1
    r = masd_report_json(project="demo", as_of=AS_OF, compare=None, db=db)
    assert r["summary"]["learners"] == 1 and r["summary"]["adoptions"] == 1
    assert r["summary"]["subtype_actual"]["gm"] == 2
    assert r["outcomes"]["exclusions"]["included"]["total"] == 1, "two measured visits with dates that agree"
    assert r["insights"]["overview"]


def test_trainer_role_only_for_f2f_learners(world):
    from fastapi import HTTPException

    from app.routers.masd import TrainerRoleIn, put_trainer_role

    db, trained, other = world
    with pytest.raises(HTTPException) as refused:
        put_trainer_role(other.id, TrainerRoleIn(trainer_role="facilitator"), admin_email="a@t", db=db)
    assert refused.value.status_code == 404
    out = put_trainer_role(trained.id, TrainerRoleIn(trainer_role="facilitator"), admin_email="a@t", db=db)
    assert out["trainer_role"] == "facilitator"


def test_batches_targets_and_table_settings(world):
    from fastapi import HTTPException

    from app.routers.masd import (BatchAssignIn, ExpectedFormsIn, SettingsIn, get_settings, masd_report_json,
                                  put_expected_forms, put_learner_batch, put_settings, reset_expected_forms)

    db, trained, other = world
    out = put_settings(SettingsIn(
        training_date=date(2026, 6, 1),
        targets=[{"from": "2026-06-01", "anc": 1, "pnc_lt5": 1, "pnc_ge5": 1, "nurse": 3},
                 {"from": "2026-08-01", "anc": 3, "pnc_lt5": 5, "pnc_ge5": 3, "nurse": 9, "buffer": 10}],
        batches=[{"name": "Batch 1", "end_date": "2026-06-01"}, {"name": "Batch 2", "end_date": "2026-06-20"}],
    ), project="demo", admin_email="a@t", db=db)
    assert not out["targets_are_default"] and [b["name"] for b in out["batches"]] == ["Batch 1", "Batch 2"]
    assert [s["buffer"] for s in out["targets"]] == [None, 10], "unset = the default buffer"
    b2 = out["batches"][1]["id"]
    with pytest.raises(HTTPException):
        put_learner_batch(other.id, BatchAssignIn(batch_id=b2), admin_email="a@t", db=db)   # not F2F
    put_learner_batch(trained.id, BatchAssignIn(batch_id=b2), admin_email="a@t", db=db)
    r = masd_report_json(project="demo", as_of=AS_OF, compare=None, db=db)
    row = next(x for x in r["learners"] if x["id"] == trained.id)
    assert (row["batch"], row["fu_raw"], row["target"]) == ("Batch 2", 94, 11), "20 Jun + a week → 29 Sep"
    assert row["tranche_fu"] == [90, 45], "tranche 2: 1 Aug + its own 10-day buffer → 49 days"
    assert get_settings(project="demo", db=db)["batches"][1]["learners"] == 1

    table = R.default_expected_forms()
    table["rows"][R.PNC_LT5]["gm"] = [v + 1 for v in table["rows"][R.PNC_LT5]["gm"]]
    saved = put_expected_forms(ExpectedFormsIn(**table), admin_email="a@t", db=db)
    assert not saved["is_default"]
    r2 = masd_report_json(project="demo", as_of=AS_OF, compare=None, db=db)
    assert next(x for x in r2["learners"] if x["id"] == trained.id)["ideal"]["gm"] > row["ideal"]["gm"]
    bad = R.default_expected_forms()
    bad["rows"][R.ANC]["anc"][3] = 0
    with pytest.raises(HTTPException) as refused:
        put_expected_forms(ExpectedFormsIn(**bad), admin_email="a@t", db=db)
    assert refused.value.status_code == 422
    assert reset_expected_forms(admin_email="a@t", db=db)["is_default"]

    # Deleting a batch puts its learners back on the project's training date.
    put_settings(SettingsIn(training_date=date(2026, 6, 1), batches=[{"id": out["batches"][0]["id"],
                                                                        "name": "Batch 1", "end_date": "2026-06-01"}]),
                 project="demo", admin_email="a@t", db=db)
    r3 = masd_report_json(project="demo", as_of=AS_OF, compare=None, db=db)
    assert next(x for x in r3["learners"] if x["id"] == trained.id)["batch"] is None


def test_settings_reject_tranche_before_training():
    from pydantic import ValidationError

    from app.routers.masd import SettingsIn

    with pytest.raises(ValidationError):
        SettingsIn(training_date=date(2026, 7, 1), tranche2_start=date(2026, 6, 1))
    with pytest.raises(ValidationError):
        SettingsIn(benchmarks={"label": "x", "age_bands": {"lt6": {"stunting": 140}}})
    with pytest.raises(ValidationError):
        SettingsIn(targets=[{"from": "2026-06-01", "anc": 1}, {"from": "2026-06-01", "anc": 2}])
    with pytest.raises(ValidationError):
        SettingsIn(batches=[{"name": "A", "end_date": "2026-06-01"}, {"name": "a ", "end_date": "2026-06-02"}])


# ── Review of 3 Oct 2026 ───────────────────────────────────────────────────


def test_severe_forms_are_counted_below_minus_three(report):
    prev = report["outcomes"]["prevalence"]["overall"]
    for ind in R.SEVERE_INDICATORS:
        assert set(prev[ind]) >= {"bv", "av", "lv", "abs_change", "rel_change"}
        base = R.SEVERE_OF[ind]
        for p in ("bv", "av", "lv"):
            if prev[base][p] is not None and prev[ind][p] is not None:
                assert prev[ind][p] <= prev[base][p], "below −3 SD is a subset of below −2 SD"


def test_nfhs_first_year_is_weighted_by_sample_size():
    assert R.combined_prevalence(28.8, 1000, 24.8, 1000) == 26.8
    assert R.combined_prevalence(30.0, 300, 20.0, 100) == 27.5
    assert R.combined_prevalence(30.0, None, 20.0, 100) is None, "no sample size, no combined figure"
    data = _project()
    data.benchmarks = {"label": "NFHS-5", "age_bands": {
        "lt6": {"stunting": 30.0, "sam": 4.0, "n": 300}, "m6_11": {"stunting": 20.0, "sam": 2.0, "n": 100}}}
    lt12 = E.compute(data, AS_OF)["outcomes"]["benchmarks_lt12"]
    assert (lt12["stunting"], lt12["sam"], lt12["n"]) == (27.5, 3.5, 400)
    assert lt12["wasting"] is None
    assert E.compute(_project(), AS_OF)["outcomes"]["benchmarks_lt12"] is None, "the default has no sample sizes"


def test_settings_take_sample_sizes_severe_and_under5():
    from pydantic import ValidationError

    from app.routers.masd import SettingsIn

    s = SettingsIn(benchmarks={"label": "NFHS-5", "age_bands": {"lt6": {"stunting": 28.8, "sam": 3.1, "n": 1200}},
                               "under5": {"stunting": 38.0, "severe_stunting": 14.0}})
    dumped = s.benchmarks.model_dump()
    assert dumped["age_bands"]["lt6"]["n"] == 1200 and dumped["under5"]["severe_stunting"] == 14.0
    with pytest.raises(ValidationError):
        SettingsIn(benchmarks={"label": "x", "age_bands": {"lt6": {"n": -5}}})


def test_pnc_split_is_one_setting():
    from app.routers.metadata import programme_constants

    out = programme_constants()
    assert out == {"pnc_split_days": R.FIVE_MONTHS_DAYS, "pnc_split_months": R.PNC_SPLIT_MONTHS}
    assert R.ADOPTION_TYPE_LABELS[R.PNC_LT5] == f"PNC <{R.PNC_SPLIT_MONTHS}M"
