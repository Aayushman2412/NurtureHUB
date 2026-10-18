"""The supervisors' review of 18 Oct 2026.

* The learner's tests are fed back topic by topic, with the videos to watch
  again for the weak topics.
* The dashboard opens on Training until the learner is selected for
  face-to-face training (or already adopting), then on Adoptions.
* A case with no adoption date entered takes the day the mother was
  registered — as the MASD dashboard does — so its expected counts are not nil.
* Faltering cases: the moderate/severe grades of the latest visit and a fall
  in weight; filters; an Excel sheet per batch monitor.

    cd backend && ./venv-win/Scripts/python.exe -m pytest tests/test_review_2026_10_18.py -v
"""
from __future__ import annotations

import base64
import os
import tempfile
from datetime import date, datetime, timedelta, timezone

import pytest

os.environ.setdefault("APP_ENV", "development")
_TEST_KEY = base64.urlsafe_b64encode(b"0123456789abcdef0123456789abcdef").decode().rstrip("=")
os.environ.setdefault("PHI_ENCRYPTION_KEYS", f"v1:{_TEST_KEY}")
os.environ.setdefault("PHI_ENCRYPTION_ACTIVE_KEY", "v1")
os.environ.setdefault("PHI_INDEX_KEY", "test-index-key-not-for-production")
os.environ.setdefault("AUDIT_HMAC_KEY", "test-audit-key-not-for-production-0123456789")

from sqlalchemy import create_engine  # noqa: E402
from sqlalchemy.orm import sessionmaker  # noqa: E402

from app import faltering, models  # noqa: E402
from app.database import Base  # noqa: E402
from app.learner_tests import topic_report  # noqa: E402
from app.routers import growth  # noqa: E402
import app.models_live  # noqa: F401,E402
import app.models_security  # noqa: F401,E402


@pytest.fixture()
def db():
    fd, path = tempfile.mkstemp(suffix=".sqlite")
    os.close(fd)
    engine = create_engine(f"sqlite:///{path}", connect_args={"check_same_thread": False})
    Base.metadata.create_all(engine)
    session = sessionmaker(bind=engine)()
    yield session
    session.close()
    engine.dispose()
    os.unlink(path)


def _learner(db, project=None, email="a@example.org"):
    u = models.User(email=email, full_name="Asha Learner", password_hash="x", is_verified=True,
                    role="AWW", program_district_id=project.id if project else None)
    db.add(u)
    db.flush()
    return u


# ── Tests, topic by topic ──────────────────────────────────────────────────

def test_topic_report_scores_each_topic_and_suggests_videos(db):
    project = models.ProgramDistrict(name="Demo", slug="demo")
    db.add(project)
    db.flush()
    videos = models.Stage(program_district_id=project.id, title="Phase 1", order_index=0)
    exam = models.Stage(program_district_id=project.id, title="Phase 2", order_index=1, stage_type="test")
    db.add_all([videos, exam])
    db.flush()
    db.add_all([
        models.Tutorial(stage_id=videos.id, title="Counseling Mothers on Breastfeeding", order_index=0),
        models.Tutorial(stage_id=videos.id, title="Nutrition & Growth Standards", order_index=1),
    ])
    test = models.Test(stage_id=exam.id, title="Formative", test_type="formative", passing_score_pct=70,
                       total_questions=5, duration_minutes=10, max_attempts=3, status="ended")
    db.add(test)
    db.flush()
    qs = [models.Question(test_id=test.id, text=f"Q{i}", order_index=i, topic=topic)
          for i, topic in enumerate(["Breastfeeding", "Breastfeeding", "Growth monitoring",
                                     "Growth monitoring", None])]
    db.add_all(qs)
    learner = _learner(db, project)
    attempt = models.TestAttempt(user_id=learner.id, test_id=test.id, attempt_number=1, score=60.0,
                                 submitted_at=datetime.now(timezone.utc))
    db.add(attempt)
    db.flush()
    # Breastfeeding 0/2, growth 2/2, untagged 1/1
    for q, ok in zip(qs, [False, False, True, True, True]):
        db.add(models.TestAnswer(attempt_id=attempt.id, question_id=q.id, is_correct=ok))
    db.commit()

    [report] = topic_report(db, learner)
    by_topic = {r["topic"]: r for r in report["topics"]}
    assert by_topic["Breastfeeding"]["pct"] == 0.0
    assert by_topic["Growth monitoring"]["pct"] == 100.0
    assert by_topic[None]["total"] == 1                     # untagged questions still count
    assert report["topics"][0]["topic"] == "Breastfeeding"   # weakest first
    assert report["weak"] == ["Breastfeeding"]
    assert report["strong"] == ["Growth monitoring"]
    assert (report["correct"], report["total"]) == (3, 5)
    [advice] = report["videos"]
    assert [v["title"] for v in advice["videos"]] == ["Counseling Mothers on Breastfeeding"]


def test_topic_report_skips_tests_not_taken(db):
    project = models.ProgramDistrict(name="Demo", slug="demo")
    db.add(project)
    db.flush()
    stage = models.Stage(program_district_id=project.id, title="Phase 2", order_index=1, stage_type="test")
    db.add(stage)
    db.flush()
    db.add(models.Test(stage_id=stage.id, title="Screening", test_type="screening", passing_score_pct=70,
                       total_questions=1, duration_minutes=10, max_attempts=3, status="active"))
    learner = _learner(db, project)
    db.commit()
    assert topic_report(db, learner) == []


# ── Dashboard modes ────────────────────────────────────────────────────────

def test_dashboard_opens_on_adoptions_once_selected(db):
    from app.routers.dashboard import get_dashboard_data

    learner = _learner(db)
    db.commit()
    data = get_dashboard_data(current_user=learner, db=db)
    assert (data.selected_for_f2f, data.adoptions_open) == (False, False)

    db.add(models.FaceToFaceSelection(user_id=learner.id))
    db.commit()
    data = get_dashboard_data(current_user=learner, db=db)
    assert (data.selected_for_f2f, data.adoptions_open) == (True, True)
    assert data.awaiting_results is False


# ── Expected counts: registration day when no adoption date ───────────────

def test_case_without_adoption_date_takes_the_registration_day():
    today = date.today()
    registered = today - timedelta(days=60)
    case = {
        "child": {"id": 1, "dob": (today - timedelta(days=90)).isoformat(), "adoption_date": None},
        "mother": {"id": 1, "adoption_date": None, "lmp": None, "created_on": registered.isoformat()},
        "learner": {"id": 1, "role": "AWW"},
    }
    adoption = growth._adoption(case, None, today)
    assert adoption["type"] is not None              # was None: "why is expected count nil?"
    assert adoption["anchor"] == registered
    assert adoption["expected"]["cg"] is not None

    case["mother"]["created_on"] = None              # nothing to anchor on: still unknown
    assert growth._adoption(case, None, today)["type"] is None


# ── Faltering cases ────────────────────────────────────────────────────────

def _row(cats=None):
    names = growth.Z_CATEGORY_NAMES
    return {"categories": {m: {"lv": (cats or {}).get(m, names[m][0])} for m in names}}


def test_faltering_flags_grades_and_weight_loss():
    names = growth.Z_CATEGORY_NAMES
    visits = [{"date": "2026-09-01", "weight": 6.2}, {"date": "2026-09-20", "weight": 6.05}]
    flags, change = faltering._flags(_row({"wfaz": names["wfaz"][3], "hfaz": names["hfaz"][2]}), {"visits": visits})
    assert flags == ["suw", "mst", "faltering"]
    assert change == -150

    gained = [{"date": "2026-09-01", "weight": 6.0}, {"date": "2026-09-20", "weight": 6.3}]
    assert faltering._flags(_row({"wfhz": names["wfhz"][1]}), {"visits": gained}) == ([], 300)  # mild: not listed


def _falter(case_id, flags, monitor=None, batch=None):
    return {"case_id": case_id, "flags": flags, "batch": batch,
            "monitor": {"id": monitor, "name": f"Monitor {monitor}"} if monitor else None,
            "learner": {"id": case_id, "name": f"L{case_id}", "cadre": "AWW", "block": "Warora", "mobile": None},
            "mother": {"name": "M", "uid": "U", "village": "V"},
            "child": {"name": "B", "uid": "C", "sex": "girls", "dob": None, "age_months": 4.0},
            "last_visit": "2026-10-01", "days_since_visit": 17, "weight": 5.1, "length": 60.0,
            "weight_change_g": -100, "z": {"wfaz": -3.2, "hfaz": -1.0, "wfhz": -2.1},
            "categories": {"wfaz": "Severe underweight (SUW)", "hfaz": "Mild stunting", "wfhz": "Moderate wasting (MAM)"}}


def test_faltering_filters_and_a_sheet_per_batch_monitor():
    from openpyxl import load_workbook

    rows = [_falter(1, ["suw"], monitor=7, batch="Batch 1"), _falter(2, ["mam"], monitor=8),
            _falter(3, ["faltering"]), _falter(4, ["muw", "faltering"], monitor=7)]
    assert [r["case_id"] for r in faltering.apply_filters(rows, conditions=["faltering"])] == [3, 4]
    assert [r["case_id"] for r in faltering.apply_filters(rows, monitor_id=7)] == [1, 4]
    assert [r["case_id"] for r in faltering.apply_filters(rows, no_monitor=True)] == [3]
    assert [r["case_id"] for r in faltering.apply_filters(rows, batch="Batch 1")] == [1]

    wb = load_workbook(faltering.xlsx(rows, project="Demo", conditions=faltering.CONDITION_KEYS))
    assert wb.sheetnames == ["All faltering cases", "Monitor - Monitor 7", "Monitor - Monitor 8", "No batch monitor"]
    assert wb["Monitor - Monitor 7"].max_row == 4 + 2        # title, note, blank, header + 2 cases
    assert wb["All faltering cases"]["A5"].value == "Severe underweight (SUW)"


def test_demo_topics_from_question_text():
    from scripts.demo_accounts import topic_for

    assert topic_for("At what age should complementary feeding start alongside continued breastfeeding?") == "Complementary feeding"
    assert topic_for("What is colostrum?") == "Breastfeeding"
    assert topic_for("What is the MUAC threshold for SAM?") == "Malnutrition (SAM / MAM)"
    assert topic_for("Something else entirely") == "General"
